"""跨工作簿引用：解析与校验放行，本地试算能真正打开外部工作簿读取。"""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from excel_formula.evaluator import UnsupportedFormula, evaluate_formula, format_value
from excel_formula.excel_reader import WorkbookView
from excel_formula.external import ExternalBookLoader
from excel_formula.formula_parser import parse_formula
from excel_formula.validator import validate_formula


def test_unquoted_external_ref_passes_validation():
    result, ast = validate_formula("=SUM([book2]Sheet1!A1:A5)", sheet_names=["Sheet1"])
    assert result.ok
    assert ast is not None
    assert any("跨工作簿" in w for w in result.warnings)


def test_quoted_external_ref_passes_validation():
    result, _ = validate_formula("=SUM('[book2]Sheet1'!A1:A5)", sheet_names=["Sheet1"])
    assert result.ok
    assert any("跨工作簿" in w for w in result.warnings)


def test_path_style_external_ref_parses():
    result, ast = validate_formula(
        "=SUM('C:\\数据\\[book2.xlsx]Sheet1'!A1:A5)", sheet_names=["Sheet1"]
    )
    assert result.ok and ast is not None
    assert any("跨工作簿" in w for w in result.warnings)


def test_external_ref_warning_is_listed_once():
    result, _ = validate_formula(
        "=SUM([book2]Sheet1!A1:A5)+SUM([book2]Sheet1!B1:B5)", sheet_names=["Sheet1"]
    )
    hits = [w for w in result.warnings if "跨工作簿" in w]
    assert len(hits) == 1


def test_external_ref_marks_unverified_in_evaluator(sample_xlsx):
    with WorkbookView(sample_xlsx) as view:
        sheets = {name: view.wb_values[name] for name in view.sheet_names}
    with pytest.raises(UnsupportedFormula, match="跨工作簿"):
        evaluate_formula(parse_formula("=[book2]Sheet1!A1"), sheets, "Sheet1")


# ------------------------------------------------------------------ 本地试算（真打开外部工作簿）
@pytest.fixture
def linked_workbooks(tmp_path):
    """主工作簿 main.xlsx 与同目录 data.xlsx；data 含缓存值、无缓存公式与不支持函数。"""
    data = tmp_path / "data.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for offset, value in enumerate([100, 200, 300, 400], start=2):
        sheet.cell(row=offset, column=2, value=value)  # B2:B5
    sheet["C2"] = "=B2*2"  # 无缓存公式（单层）
    sheet["E2"] = "=C2*2"  # 无缓存公式且引用另一条无缓存公式（两层递归）
    sheet["D2"] = "=SUBTOTAL(9,B2:B5)"  # 本地不支持的函数
    workbook.save(data)

    main = tmp_path / "main.xlsx"
    main_book = Workbook()
    main_book.active.title = "Sheet1"
    main_book.save(main)
    return main, data


def _eval_linked(main, formula):
    with WorkbookView(main) as view:
        sheets = {name: view.wb_values[name] for name in view.sheet_names}
        loader = ExternalBookLoader()
        try:
            return evaluate_formula(
                parse_formula(formula),
                sheets,
                "Sheet1",
                load_external=loader,
                book_path=view.path,
            )
        finally:
            loader.close()


@pytest.mark.parametrize(
    "formula,expected",
    [
        ("=SUM([data.xlsx]Sheet1!B2:B5)", "1000"),
        ("=[data]Sheet1!B2", "100"),  # 省略扩展名按同目录补全
        ("='[data.xlsx]Sheet1'!B2", "100"),  # 带引号写法
        ("=[data.xlsx]SHEET1!B2", "100"),  # 工作表名不区分大小写
        ("=[data.xlsx]Sheet1!C2", "200"),  # 外部公式无缓存 → 递归本地求值
        ("=[data.xlsx]Sheet1!E2", "400"),  # 两层递归
        ('=SUMIF([data.xlsx]Sheet1!B2:B5,">150")', "900"),
    ],
)
def test_cross_workbook_local_evaluation(linked_workbooks, formula, expected):
    main, _ = linked_workbooks
    assert format_value(_eval_linked(main, formula)) == expected


def test_path_style_external_reference_evaluates(linked_workbooks):
    """带完整目录的引用（Excel 保存远端文件时的常见形态）直接用该目录解析。"""
    main, data = linked_workbooks
    formula = f"='{data.parent}\\[data.xlsx]Sheet1'!B2"
    assert format_value(_eval_linked(main, formula)) == "100"


def test_unsupported_function_inside_external_book_stays_unverified(linked_workbooks):
    main, _ = linked_workbooks
    with pytest.raises(UnsupportedFormula, match="SUBTOTAL"):
        _eval_linked(main, "=[data.xlsx]Sheet1!D2")


@pytest.mark.parametrize(
    "formula,match",
    [
        ("=[missing.xlsx]Sheet1!A1", "找不到外部工作簿"),
        ("=[data.xlsx]NoSheet!A1", "找不到工作表"),
    ],
)
def test_unavailable_external_book_stays_unverified(linked_workbooks, formula, match):
    """文件或工作表不存在时降级为"未验证（原因）"，绝不报错阻断。"""
    main, _ = linked_workbooks
    with pytest.raises(UnsupportedFormula, match=match):
        _eval_linked(main, formula)


def test_circular_external_references_are_detected(tmp_path):
    """两个工作簿互引且都无缓存时，被循环检测拦住而不是无限递归。"""
    first = tmp_path / "first.xlsx"
    book = Workbook()
    book.active.title = "Sheet1"
    book.active["A1"] = "=[second.xlsx]Sheet1!A1"
    book.save(first)

    second = tmp_path / "second.xlsx"
    book = Workbook()
    book.active.title = "Sheet1"
    book.active["A1"] = "=[first.xlsx]Sheet1!A1"
    book.save(second)

    with pytest.raises(UnsupportedFormula, match="循环"):
        _eval_linked(first, "=[first.xlsx]Sheet1!A1")


def test_corrupt_external_file_stays_unverified(tmp_path):
    main = tmp_path / "main.xlsx"
    book = Workbook()
    book.active.title = "Sheet1"
    book.save(main)
    (tmp_path / "broken.xlsx").write_bytes(b"not an excel file")

    with pytest.raises(UnsupportedFormula, match="无法打开"):
        _eval_linked(main, "=[broken.xlsx]Sheet1!A1")


def test_same_folder_reference_requires_book_path(linked_workbooks):
    """未提供当前工作簿路径（不知道同目录在哪）时降级为未验证而非瞎猜。"""
    main, _ = linked_workbooks
    with WorkbookView(main) as view:
        sheets = {name: view.wb_values[name] for name in view.sheet_names}
        loader = ExternalBookLoader()
        try:
            with pytest.raises(UnsupportedFormula, match="未提供当前工作簿目录"):
                evaluate_formula(
                    parse_formula("=[data.xlsx]Sheet1!B2"),
                    sheets,
                    "Sheet1",
                    load_external=loader,
                )
        finally:
            loader.close()
