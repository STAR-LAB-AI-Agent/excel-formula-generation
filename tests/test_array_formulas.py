"""数组语义公式：写入侧 CSE 化与读取侧归一，保证本地结果与 Excel 打开后一致。

区域参与一元/二元运算的公式（=SUM(B2:B6*C2:C6)、条件数组写法等）本地按数组
语义求值；若以裸公式写入 .xlsx，Excel/WPS 打开时会按传统"隐式交叉"只取交叉
单值，与本地结果及网页展示不一致。写入侧改用数组公式（CSE）形态，另一侧
（读取器/摘要/网格/覆盖提醒）把 openpyxl 的 ArrayFormula 对象归一为公式文本。
"""
from __future__ import annotations

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

import webapp
from excel_formula.evaluator import UnsupportedFormula, evaluate_formula, needs_array_formula
from excel_formula.excel_reader import WorkbookView
from excel_formula.formula_parser import parse_formula
from excel_formula.writer import CellWrite, apply_writes

ARRAY_FORMULA = "=SUM(B2:B6*C2:C6)"  # 同形区域逐元素相乘：85*78+88*85+80*82+90*87+78*80
# “班里第 k 名的姓名”写法：IF 数组 + LARGE + MATCH + ROW()，k 随公式行号变化
RANK_FORMULA = (
    '=IFERROR(INDEX(成绩单!$B$2:$B$13,MATCH(LARGE(IF(成绩单!$C$2:$C$13="高一1班",'
    '成绩单!$E$2:$E$13),ROW()-1),IF(成绩单!$C$2:$C$13="高一1班",成绩单!$E$2:$E$13),0)),"")'
)


# ------------------------------------------------------------------ 判定器
@pytest.mark.parametrize(
    "formula,expected",
    [
        ("=SUM(B2:B6*C2:C6)", True),                      # 同形区域逐元素运算
        ("=SUM(B2:B6*B2:B4)", True),                      # 维度不兼容：本地给 #N/A，仍需 CSE
        ('=SUM((A2:A6="Math")*B2:B6)', True),             # 条件布尔数组乘回取值区
        ("=-B2:F2", True),                                # 区域一元运算
        ("=IF(A2:A6>85,B2:B6,0)", True),                  # IF 数组形式
        ('=SUMPRODUCT((A2:A6="Math")*B2:B6)', False),     # 参数天然数组，不需要 CSE
        ("=B2*C2+D2", False),                             # 纯标量公式
        ("=SUM(B2:F2)", False),                           # 区域作为汇总函数参数无需展开
        ('=IF(A2>2,"高","低")', False),                   # IF 标量形式
    ],
)
def test_needs_array_formula(formula, expected):
    assert needs_array_formula(parse_formula(formula)) is expected


def test_needs_array_formula_for_whole_column():
    """整列引用参与比较也算数组语义（公式里很常见的 =SUM((D:D="华北")*1) 写法）。"""
    assert needs_array_formula(parse_formula('=SUM((D:D="华北")*1)')) is True


def test_needs_array_formula_for_if_array_constant():
    """IF({1,0},…) 反向查找写法：数组常量参与，需要 CSE 才能正确求值。"""
    assert needs_array_formula(parse_formula("=VLOOKUP(A2,IF({1,0},B2:B9,A2:A9),2,FALSE)")) is True


# ------------------------------------------------------------------ 写入侧：CSE 化
def test_array_semantics_formula_written_as_cse(sample_xlsx, settings):
    """数组语义公式以数组公式形态写入：读回是 ArrayFormula，文本与 ref 不变。"""
    apply_writes(
        sample_xlsx,
        [CellWrite(sheet="Sheet1", cell="H2", formula=ARRAY_FORMULA)],
        settings,
    )
    workbook = load_workbook(sample_xlsx)
    value = workbook["Sheet1"]["H2"].value
    workbook.close()
    assert isinstance(value, ArrayFormula)
    assert value.text == ARRAY_FORMULA
    assert value.ref == "H2"


def test_plain_formulas_stay_plain(sample_xlsx, settings):
    """标量公式与 SUMPRODUCT（参数天然按数组处理）保持普通公式形态。"""
    apply_writes(
        sample_xlsx,
        [
            CellWrite(sheet="Sheet1", cell="H2", formula="=B2*2"),
            CellWrite(sheet="Sheet1", cell="H3", formula='=SUMPRODUCT((A2:A6="Math")*B2:B6)'),
        ],
        settings,
    )
    workbook = load_workbook(sample_xlsx)
    assert workbook["Sheet1"]["H2"].value == "=B2*2"
    assert workbook["Sheet1"]["H3"].value == '=SUMPRODUCT((A2:A6="Math")*B2:B6)'
    workbook.close()


def test_cse_survives_second_write(sample_xlsx, settings):
    """二次写入其他格子后，已写入的数组公式仍以 CSE 形态保留（真实流程会反复写）。"""
    apply_writes(
        sample_xlsx,
        [CellWrite(sheet="Sheet1", cell="H2", formula=ARRAY_FORMULA)],
        settings,
    )
    apply_writes(sample_xlsx, [CellWrite(sheet="Sheet1", cell="H4", formula="=B2*3")], settings)
    workbook = load_workbook(sample_xlsx)
    value = workbook["Sheet1"]["H2"].value
    workbook.close()
    assert isinstance(value, ArrayFormula)
    assert value.text == ARRAY_FORMULA


# ------------------------------------------------------------------ 读取侧：归一为公式文本
def test_reader_normalizes_array_formula(sample_xlsx, settings):
    """cell_formula 与摘要都把 ArrayFormula 归一为公式文本，不出现对象 repr。"""
    apply_writes(
        sample_xlsx,
        [CellWrite(sheet="Sheet1", cell="H2", formula=ARRAY_FORMULA)],
        settings,
    )
    with WorkbookView(sample_xlsx) as view:
        assert view.cell_formula("Sheet1", "H2") == ARRAY_FORMULA
        text = view.digest().to_prompt()
    assert ARRAY_FORMULA in text
    assert "ArrayFormula" not in text


# ------------------------------------------------------------------ 网页网格
def test_grid_shows_array_result_and_keeps_formula(sample_xlsx, settings):
    """网格对数组公式格显示本地数组语义结果；引用它的链式公式也能递归求值。"""
    apply_writes(
        sample_xlsx,
        [
            CellWrite(sheet="Sheet1", cell="H2", formula=ARRAY_FORMULA),
            CellWrite(sheet="Sheet1", cell="H3", formula="=H2+1"),
        ],
        settings,
    )
    with WorkbookView(sample_xlsx) as view:
        payload = webapp.build_display_grid(view, view.digest())
    rows = {row["row"]: row for row in payload["rows"]}
    assert rows[2]["values"][7] == "34740"  # 85*78+88*85+80*82+90*87+78*80
    assert rows[2]["formulas"][7] == ARRAY_FORMULA
    assert rows[3]["values"][7] == "34741"  # 引用了无缓存的 CSE 格：递归本地求值
    assert rows[3]["formulas"][7] == "=H2+1"


# ------------------------------------------------------------------ ROW/COLUMN 位置函数
def test_row_column_need_cell_position():
    """无参 ROW()/COLUMN() 依赖公式所在位置：缺位置标未验证，给位置返回行列号。"""
    sheets: dict = {}
    with pytest.raises(UnsupportedFormula):
        evaluate_formula(parse_formula("=ROW()"), sheets, "Sheet1")
    assert evaluate_formula(
        parse_formula("=ROW()"), sheets, "Sheet1", current_cell=(16, 10)
    ) == 16.0
    assert evaluate_formula(
        parse_formula("=COLUMN()"), sheets, "Sheet1", current_cell=(16, 10)
    ) == 10.0


def test_row_column_with_reference():
    """带引用参数：返回引用首行/首列；区域形式按数组参与计算（SUM 求和整段行号）。"""
    sheets: dict = {}
    assert evaluate_formula(parse_formula("=ROW($C$5)"), sheets, "Sheet1") == 5.0
    assert evaluate_formula(parse_formula("=COLUMN(C5)"), sheets, "Sheet1") == 3.0
    assert evaluate_formula(parse_formula("=SUM(ROW(A1:A5))"), sheets, "Sheet1") == 15.0
    assert evaluate_formula(parse_formula("=SUM(COLUMN(A3:D3))"), sheets, "Sheet1") == 10.0


def test_needs_array_formula_for_row_range():
    """ROW(A2:A6) 区域形式在数组上下文返回数组，需要 CSE；无参 ROW() 是标量不需要。"""
    assert needs_array_formula(parse_formula("=SUM(ROW(A2:A6))")) is True
    assert needs_array_formula(parse_formula("=ROW()-1")) is False


def test_row_based_rank_formula_gets_result(class_scores_xlsx, settings):
    """条件排名公式（IF 数组 + LARGE + MATCH + ROW）从“未验证”变为算出结果。

    公式贴在班级信息 E2（ROW()-1=1），取高一1班数学第一名：92 分对应的学生1。
    """
    apply_writes(
        class_scores_xlsx,
        [CellWrite(sheet="班级信息", cell="E2", formula=RANK_FORMULA)],
        settings,
    )
    with WorkbookView(class_scores_xlsx) as view:
        payload = webapp.build_display_grid(view, view.digest("班级信息"))
    rows = {row["row"]: row for row in payload["rows"]}
    assert rows[2]["values"][4] == "学生1"
    assert rows[2]["formulas"][4] == RANK_FORMULA
