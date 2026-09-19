"""求值器与读取器测试：整表无损文本化、本地公式计算（可选功能"公式自动验证"的基础）。"""
from __future__ import annotations

import pytest

from excel_formula.evaluator import ExcelError, UnsupportedFormula, evaluate_formula, format_value
from excel_formula.excel_reader import WorkbookView
from excel_formula.formula_parser import parse_formula


def _eval(path, formula):
    with WorkbookView(path) as view:
        sheets = {name: view.wb_values[name] for name in view.sheet_names}
        return evaluate_formula(parse_formula(formula), sheets, "Sheet1")


# ------------------------------------------------------------------ 读取器
def test_digest_dumps_every_cell(sample_xlsx):
    """读取器只搬运不判断：每一行都要在里面，行号连续。"""
    with WorkbookView(sample_xlsx) as view:
        digest = view.digest()
        assert digest.name == "Sheet1"  # Sheet2 为空表，应自动跳过
        assert digest.data_range == "A1:F6"
        assert [row for row, _ in digest.rows] == [1, 2, 3, 4, 5, 6]
        assert digest.rows[0][1] == [
            "Subject", "Student1", "Student2", "Student3", "Student4", "Student5"
        ]
        assert digest.rows[1][1] == ["Math", "85", "78", "92", "88", "95"]
        assert digest.truncated_rows == 0


def test_digest_prompt_is_lossless_tsv(sample_xlsx):
    """提示文本里必须带坐标系，且末行数据不能丢。"""
    with WorkbookView(sample_xlsx) as view:
        text = view.digest().to_prompt()
    assert "\tA\tB\tC\tD\tE\tF" in text  # 列字母表头
    assert "1\tSubject\tStudent1" in text  # 行号 + 原始值
    assert "6\tBiology\t78\t80\t86\t83\t81" in text  # 最后一行也在
    assert "已省略" not in text  # 小表不应触发截断


def test_digest_truncates_large_table_at_both_ends(tmp_path):
    """超过行数闸门时保留头尾两段，并显式告知已截断。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["序号", "重量"])
    for index in range(1, 500):
        sheet.append([index, index * 2])
    path = tmp_path / "big.xlsx"
    workbook.save(path)
    workbook.close()

    with WorkbookView(path) as view:
        digest = view.digest()
        text = digest.to_prompt()
    assert digest.max_row == 500
    assert digest.truncated_rows == 500 - 60  # 头 50 行 + 尾 10 行
    assert [row for row, _ in digest.rows][:3] == [1, 2, 3]
    assert [row for row, _ in digest.rows][-1] == 500
    assert "已省略" in text


# ------------------------------------------------------------------ 求值器
@pytest.mark.parametrize(
    "formula,expected",
    [
        ("=SUM(B2:F2)", "438"),
        ("=AVERAGE(B2:F2)", "87.6"),
        ("=MAX(B2:F6)", "95"),
        ("=MIN(B2:F6)", "75"),
        ("=COUNT(B2:F2)", "5"),
        ('=COUNTIF(B2:F2,">85")', "3"),
        ('=SUMIF(A2:A6,"Math",B2:B6)', "85"),
        ('=IF(AVERAGE(B2:F2)>=85,"优秀","一般")', "优秀"),
        ("=ROUND(AVERAGE(B2:F2),1)", "87.6"),
        ("=B2/SUM(B2:B6)", "0.2019"),  # 85/421
        ("=B2&\"分\"", "85分"),
        ('=VLOOKUP("Physics",A2:F6,2,FALSE)', "80"),
    ],
)
def test_evaluate_common_formulas(sample_xlsx, formula, expected):
    assert format_value(_eval(sample_xlsx, formula)) == expected


def test_division_by_zero_returns_excel_error(sample_xlsx):
    assert format_value(_eval(sample_xlsx, "=B2/(B2-B2)")) == "#DIV/0!"


@pytest.mark.parametrize("data_type,is_error", [("e", True), ("s", False)])
def test_cell_error_type_is_distinct_from_text(sample_xlsx, data_type, is_error):
    with WorkbookView(sample_xlsx) as view:
        ws = view.value_sheet("Sheet1")
        ws["G2"] = "#DIV/0!"
        ws["G2"].data_type = data_type
        value = evaluate_formula(parse_formula("=G2"), {"Sheet1": ws}, "Sheet1")
    assert isinstance(value, ExcelError) is is_error
    assert format_value(value) == "#DIV/0!"


def test_formula_view_distinguishes_missing_cache_from_blank(sample_xlsx):
    with WorkbookView(sample_xlsx) as view:
        formulas = view.formula_sheet("Sheet1")
        values = view.value_sheet("Sheet1")
        formulas["G2"] = "=2"
        with pytest.raises(UnsupportedFormula, match="尚未计算"):
            evaluate_formula(parse_formula("=1/G2"), {"Sheet1": values}, "Sheet1",
                             formula_sheets={"Sheet1": formulas})
        assert format_value(evaluate_formula(
            parse_formula("=1/H2"), {"Sheet1": values}, "Sheet1",
            formula_sheets={"Sheet1": formulas},
        )) == "#DIV/0!"


def test_unsupported_function_raises(sample_xlsx):
    with pytest.raises(UnsupportedFormula):
        _eval(sample_xlsx, "=SUBTOTAL(9,B2:B6)")


def test_reverse_lookup_with_array_constant(tmp_path):
    """IF({1,0},姓名列,编号列) 用数组常量翻转列顺序，让 VLOOKUP 能向左返回。"""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in [
        ["编号", "姓名"],
        ["A01048", "叶知"],
        ["A02267", "陈佩亮"],
    ]:
        sheet.append(row)
    path = tmp_path / "emp.xlsx"
    workbook.save(path)
    workbook.close()

    # 姓名在右侧，仍可用 VLOOKUP 按姓名反查左侧编号
    assert format_value(
        _eval(path, '=VLOOKUP("陈佩亮",IF({1,0},B2:B3,A2:A3),2,0)')
    ) == "A02267"
    # 查不到时走精确匹配的 #N/A，而不是抛异常
    assert format_value(
        _eval(path, '=VLOOKUP("查无此人",IF({1,0},B2:B3,A2:A3),2,0)')
    ) == "#N/A"


def test_array_constant_literal_shapes(sample_xlsx):
    """数组常量本身可以参与求和，分号分行、逗号分列。"""
    assert format_value(_eval(sample_xlsx, "=SUM({1,2,3})")) == "6"
    assert format_value(_eval(sample_xlsx, "=COUNT({1,2;3,4})")) == "4"


# ------------------------------------------------------------------ 数组运算
@pytest.mark.parametrize(
    "formula,expected",
    [
        # 区域与标量比较得到布尔数组，再乘回区域汇总（SUMPRODUCT 条件写法的基础）
        ("=SUM((B2:F2>85)*B2:F2)", "275"),  # 92+88+95
        ('=SUMPRODUCT((A2:A6="Math")*B2:B6)', "85"),
        # 布尔数组直接求和：TRUE 计 1
        ("=SUM((B2:F2>90)*1)", "2"),
        # 同形区域逐元素相减后再汇总
        ("=SUM(B2:F2-B3:F3)", "2"),  # -3-7+1+6+5
        # 行区域 × 列区域按 Excel 广播成外积：总和 = 行合计 × 列合计
        ("=SUM(B2:F2*B3:B6)", "147168"),  # 438 × 336
        # 区域一元运算
        ("=SUM(-B2:F2)", "-438"),
        ("=SUM(B2:F2%)", "4.38"),
    ],
)
def test_array_operands_broadcast_elementwise(sample_xlsx, formula, expected):
    assert format_value(_eval(sample_xlsx, formula)) == expected


def test_array_shape_mismatch_returns_na(sample_xlsx):
    """维度不兼容（5 行比 3 行）无从广播：给 #N/A，不静默钳制出错误结果。"""
    assert format_value(_eval(sample_xlsx, "=SUM(B2:B6*B2:B4)")) == "#N/A"


def test_array_op_keeps_element_errors(sample_xlsx):
    """区域里的错误值按位置保留为错误元素：不抛异常，也不吞成静默值。"""
    with WorkbookView(sample_xlsx) as view:
        ws = view.value_sheet("Sheet1")
        ws["G2"] = "#DIV/0!"
        ws["G2"].data_type = "e"
        result = evaluate_formula(parse_formula("=B2:F2*G2:G2"), {"Sheet1": ws}, "Sheet1")
    assert format_value(result) == "#DIV/0!"  # 顶层取左上角元素（隐式交叉近似）


# ---------------------------------------------------------------- 展示级递归求值
def test_expand_uncached_resolves_formula_chain(sample_xlsx):
    """展示场景（expand_uncached）：主簿里没有缓存的公式也递归求值，链式依赖同样能算出来。"""
    with WorkbookView(sample_xlsx) as view:
        formulas = view.formula_sheet("Sheet1")
        values = view.value_sheet("Sheet1")
        formulas["G2"] = "=2"
        formulas["G3"] = "=G2*10"  # 依赖上一层公式，两层都没有缓存
        result = evaluate_formula(
            parse_formula("=G3+1"), {"Sheet1": values}, "Sheet1",
            formula_sheets={"Sheet1": formulas}, expand_uncached=True,
        )
    assert format_value(result) == "21"


def test_expand_uncached_detects_cycle(sample_xlsx):
    with WorkbookView(sample_xlsx) as view:
        formulas = view.formula_sheet("Sheet1")
        values = view.value_sheet("Sheet1")
        formulas["G2"] = "=G3+1"
        formulas["G3"] = "=G2+1"
        with pytest.raises(UnsupportedFormula, match="循环引用"):
            evaluate_formula(
                parse_formula("=G2+1"), {"Sheet1": values}, "Sheet1",
                formula_sheets={"Sheet1": formulas}, expand_uncached=True,
            )


def test_expand_uncached_reuses_dependency_result(sample_xlsx, monkeypatch):
    """同一格被引用多次：递归求值只做一次，后续命中缓存。"""
    import excel_formula.evaluator as evaluator_module

    parsed_texts: list[str] = []
    original = evaluator_module.parse_formula

    def counting_parse(text):
        parsed_texts.append(text)
        return original(text)

    monkeypatch.setattr(evaluator_module, "parse_formula", counting_parse)
    with WorkbookView(sample_xlsx) as view:
        formulas = view.formula_sheet("Sheet1")
        values = view.value_sheet("Sheet1")
        formulas["G2"] = "=2"
        result = evaluate_formula(
            parse_formula("=G2+G2+G2"), {"Sheet1": values}, "Sheet1",
            formula_sheets={"Sheet1": formulas}, expand_uncached=True,
        )
    assert format_value(result) == "6"
    assert parsed_texts.count("=2") == 1
