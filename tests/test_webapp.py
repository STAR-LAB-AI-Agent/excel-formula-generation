"""webapp 接口测试：真实启动本机 HTTP 服务，模型调用由 FakeClient 顶替。

覆盖：bootstrap / 表结构网格 / 四类意图分发（校验与概览必须 0 Token）/
SSE 打字机事件序列 / 生成→确认→写入全链路 / 备份 / 越界安全拒绝 / 思考档位切换。
"""
from __future__ import annotations

import json
import threading
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import webapp
from excel_formula.llm_client import ChatResult, Usage
from excel_formula.pipeline import FormulaService


class FakeClient:
    """与真实客户端签名兼容：回放两段 delta（思考 + 正文）后返回预置 JSON。"""

    def __init__(self, replies: list | None = None):
        self.replies = list(replies or [])
        self.calls: list[list[dict]] = []

    def chat(self, messages, *, json_mode=True, max_tokens=800, on_delta=None, thinking=None):
        self.calls.append(messages)
        reply = self.replies.pop(0)
        content = reply if isinstance(reply, str) else json.dumps(reply, ensure_ascii=False)
        if on_delta is not None:
            on_delta("先看看表结构再动手。", "reasoning")
            on_delta('{"formulas"', "content")
        return ChatResult(
            content=content, usage=Usage(prompt_tokens=120, completion_tokens=40, calls=1)
        )


class WebHarness:
    """对外只暴露 get / post / sse 三个动作，屏蔽端口与线程细节。"""

    def __init__(self, server, client: FakeClient):
        self.server = server
        self.client = client
        self.base = f"http://127.0.0.1:{server.server_address[1]}"

    def _request(self, path: str, payload: dict | None = None):
        data = None if payload is None else json.dumps(payload).encode("utf-8")
        request = urllib.request.Request(self.base + path, data=data)
        if data is not None:
            request.add_header("Content-Type", "application/json")
        try:
            with urllib.request.urlopen(request, timeout=15) as response:
                return response.status, response.read().decode("utf-8")
        except urllib.error.HTTPError as exc:
            return exc.code, exc.read().decode("utf-8")

    def get(self, path: str) -> tuple[int, dict]:
        status, raw = self._request(path)
        return status, json.loads(raw)

    def get_text(self, path: str) -> tuple[int, str]:
        request = urllib.request.Request(self.base + path)
        with urllib.request.urlopen(request, timeout=15) as response:
            return response.status, response.read().decode("utf-8")

    def post(self, path: str, payload: dict) -> tuple[int, dict]:
        status, raw = self._request(path, payload)
        return status, json.loads(raw)

    def post_raw(self, path: str, blob: bytes) -> tuple[int, dict]:
        """上传接口：请求体是原始文件字节（非 JSON）。"""
        request = urllib.request.Request(self.base + path, data=blob)
        request.add_header("Content-Type", "application/octet-stream")
        try:
            with urllib.request.urlopen(request, timeout=15) as response:
                return response.status, json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            return exc.code, json.loads(exc.read().decode("utf-8"))

    def sse(self, payload: dict) -> list[dict]:
        """读完整条 SSE 流并解析为事件列表（chunked 由 urllib 自动解码）。"""
        _, raw = self._request("/api/chat", payload)
        events: list[dict] = []
        for block in raw.split("\n\n"):
            for line in block.split("\n"):
                if line.startswith("data:"):
                    events.append(json.loads(line[5:].strip()))
        return events


@pytest.fixture()
def web(settings):
    client = FakeClient()
    service = FormulaService(settings, logger=None, client=client)
    server = webapp.make_server(service, port=0)  # 0 = 随机可用端口
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield WebHarness(server, client)
    finally:
        server.shutdown()
        server.server_close()


def _event_types(events: list[dict]) -> list[str]:
    return [e["type"] for e in events]


# ------------------------------------------------------------------ 基础接口
def test_index_page_served(web):
    status, html = web.get_text("/")
    assert status == 200
    assert "ExcelCR" in html and "api/chat" in html


def test_bootstrap_lists_files_and_samples(sample_xlsx, web):
    status, data = web.get("/api/bootstrap")
    assert status == 200
    assert data["files"] == ["scores.xlsx"]
    assert data["has_key"] is True
    assert data["samples"][0]["text"]


def test_sheet_returns_grid(sample_xlsx, web):
    status, data = web.get("/api/sheet?file=scores.xlsx")
    assert status == 200
    assert data["current"] == "Sheet1"
    assert data["grid"]["max_row"] == 6
    assert data["grid"]["rows"][0]["values"][0] == "Subject"
    assert [s["sheet"] for s in data["sheets"]] == ["Sheet1", "Sheet2"]


def test_sheet_rejects_path_outside_whitelist(web):
    status, data = web.get("/api/sheet?file=..%2F..%2Fsecret.xlsx")
    assert status == 400
    assert "错误" in data["error"] or "超出" in data["error"]


@pytest.fixture()
def formula_xlsx(settings):
    """含公式的表格：openpyxl 写出的公式没有缓存值，走本地试算分支。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append([1, 2])
    sheet["C2"] = "=A2+B2"          # 可本地试算
    sheet["D2"] = "=C2*10"          # 依赖没有缓存的公式格：应递归求值
    sheet["C3"] = "=BrokenRef!A1"   # 引用不存在的工作表：应退回公式原文
    path = settings.allowed_roots[0] / "calc.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def test_sheet_displays_formula_result_with_formula_metadata(formula_xlsx, web):
    """网格公式格直接显示计算结果值，原公式放在平行 formulas 表里供悬停查看。"""
    status, data = web.get("/api/sheet?file=calc.xlsx")
    assert status == 200
    rows = {row["row"]: row for row in data["grid"]["rows"]}
    assert rows[2]["values"][2] == "3"           # C2 = A2+B2 的本地试算结果
    assert rows[2]["formulas"][2] == "=A2+B2"    # 公式保留
    assert rows[2]["values"][3] == "30"          # D2 引用的 C2 无缓存：递归求值链
    assert rows[2]["formulas"][3] == "=C2*10"
    assert rows[1]["formulas"][0] == ""          # 普通单元格没有公式
    assert rows[3]["values"][2] == "=BrokenRef!A1"  # 试算失败退回公式原文
    assert rows[3]["formulas"][2] == "=BrokenRef!A1"


def test_start_in_background_serves_and_returns_url(settings, sample_xlsx):
    """main.py 复用的后台启动入口：真起服务、返回地址、可正常收请求。"""
    service = FormulaService(settings, logger=None, client=FakeClient())
    server, url = webapp.start_in_background(service, port=0, open_browser=False)
    try:
        assert server is not None and url is not None
        with urllib.request.urlopen(url + "api/bootstrap", timeout=15) as response:
            data = json.loads(response.read().decode("utf-8"))
        assert data["files"] == ["scores.xlsx"]
    finally:
        server.shutdown()
        server.server_close()


def test_bind_server_skips_occupied_port():
    """端口被占时必须顺延（Windows 需独占绑定，否则第二个实例会重复绑定同一端口）。"""
    first = webapp.make_server(object(), port=0)
    try:
        port = first.server_address[1]
        second = webapp._bind_server(object(), "127.0.0.1", port, attempts=3)
        try:
            assert second is not None
            assert second.server_address[1] != port
        finally:
            second.server_close()
    finally:
        first.server_close()


# ------------------------------------------------------------------ 意图分发
def test_describe_is_local_with_zero_model_calls(sample_xlsx, web):
    events = web.sse({"request": "这个表有哪些列？", "file": "scores.xlsx"})
    assert web.client.calls == []  # 概览不花 Token
    assert _event_types(events) == ["stage", "result"]
    result = events[-1]
    assert result["kind"] == "describe"
    assert result["payload"]["sheets"][0]["sheet"] == "Sheet1"
    assert "Sheet1" in result["payload"]["digest_text"]


def test_validate_is_local_with_zero_model_calls(sample_xlsx, web):
    events = web.sse({"request": '=SUMIF(B2:B6,">85") 这个公式对不对？', "file": "scores.xlsx"})
    assert web.client.calls == []
    result = events[-1]
    assert result["kind"] == "validate"
    assert result["payload"]["validation"]["ok"] is True


def test_frame_chat_is_local_and_applies_zero_write_borders(sample_xlsx, web):
    """“将范围框起来”：0 Token 本地出预览，确认后只套边框、不写任何单元格。"""
    events = web.sse({"request": "将a14到B18的表格范围框起来", "file": "scores.xlsx"})
    assert web.client.calls == []  # 纯格式需求不花 Token

    stages = [e for e in events if e["type"] == "stage"]
    assert any("不调用模型" in e["text"] for e in stages)

    result = events[-1]
    assert result["type"] == "result" and result["kind"] == "generate"
    payload = result["payload"]
    assert payload["ok"] is True
    assert payload["cells"] == []
    assert payload["table"]["table_range"] == "A14:B18"
    assert payload["table"]["label"] == "Sheet1!A14:B18（范围加细边框）"

    status, data = web.post("/api/apply", {"proposal_id": payload["proposal_id"]})
    assert status == 200
    applied = data["applied"]
    assert applied["count"] == 0
    assert applied["tables"] == ["Sheet1!A14:B18（范围加细边框）"]
    assert applied["backup_name"]

    workbook = load_workbook(sample_xlsx)
    sheet = workbook["Sheet1"]
    assert sheet["A14"].border.left.style == "thin"
    assert sheet["B18"].border.bottom.style == "thin"
    assert sheet["A14"].value is None
    workbook.close()


def test_frame_chat_without_range_asks_clarification(sample_xlsx, web):
    """没说范围时也走本地：返回追问卡片（0 Token），不送模型空跑。"""
    events = web.sse({"request": "给这个表格加边框", "file": "scores.xlsx"})
    assert web.client.calls == []

    payload = events[-1]["payload"]
    assert payload["ok"] is False
    assert "范围" in (payload["clarification"] or "")


# ------------------------------------------------------------------ 生成全链路
def test_generate_streams_deltas_and_returns_proposal(sample_xlsx, web):
    web.client.replies.append({
        "formulas": [{
            "target": "G2",
            "formula": "=SUM(B2:F2)",
            "explanation": "求B2到F2的总分",
            "fill_to": "G6",
        }]
    })
    events = web.sse({"request": "帮我在G2算每个科目的总分，填充到G6", "file": "scores.xlsx"})

    kinds = _event_types(events)
    assert "delta" in kinds  # 思考链与正文都走打字机
    deltas = [e for e in events if e["type"] == "delta"]
    assert {d["kind"] for d in deltas} == {"reasoning", "content"}

    result = events[-1]
    assert result["type"] == "result" and result["kind"] == "generate"
    payload = result["payload"]
    assert payload["ok"] is True
    assert payload["formula"] == "=SUM(B2:F2)"
    assert len(payload["cells"]) == 5  # G2 → G6
    assert payload["predicted_value"] == "438"  # 85+78+92+88+95
    assert payload["proposal_id"]


def test_generate_repairs_invalid_formula_and_reports_attempts(sample_xlsx, web):
    """首轮公式自引用被本地校验拦下，回传修复后第二轮通过——前端能看到两轮记录。"""
    web.client.replies.extend([
        {"formulas": [{"target": "G2", "formula": "=SUM(B2:G2)", "explanation": "含自身"}]},
        {"formulas": [{"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "总分"}]},
    ])
    events = web.sse({"request": "算总分", "file": "scores.xlsx"})
    payload = events[-1]["payload"]
    assert payload["ok"] is True
    assert len(payload["attempts"]) == 2
    assert payload["attempts"][0]["ok"] is False
    assert payload["attempts"][1]["ok"] is True
    assert len(web.client.calls) == 2


def test_chat_reports_error_events_for_missing_file(web):
    events = web.sse({"request": "算总分", "file": None})
    assert events[-1]["type"] == "error"
    assert "Excel 文件" in events[-1]["message"]


# ------------------------------------------------------------------ 写入闭环
def test_apply_writes_cells_and_creates_backup(sample_xlsx, web):
    web.client.replies.append({
        "formulas": [{"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "总分", "fill_to": "G6"}]
    })
    events = web.sse({"request": "算总分，填充到G6", "file": "scores.xlsx"})
    proposal_id = events[-1]["payload"]["proposal_id"]

    status, data = web.post("/api/apply", {"proposal_id": proposal_id})
    assert status == 200
    applied = data["applied"]
    assert applied["count"] == 5
    assert applied["backup_name"]

    # 文件真的被写入，且原文件已备份
    workbook = load_workbook(sample_xlsx)
    assert workbook["Sheet1"]["G2"].value == "=SUM(B2:F2)"
    assert workbook["Sheet1"]["G6"].value == "=SUM(B6:F6)"
    workbook.close()
    assert list((sample_xlsx.parent / "backups").glob("scores_*.xlsx"))

    # 令牌一次性：再写一次会被拒绝
    status, data = web.post("/api/apply", {"proposal_id": proposal_id})
    assert status == 400
    assert "失效" in data["error"]


def test_apply_rejects_unknown_token(web):
    status, data = web.post("/api/apply", {"proposal_id": "deadbeef0000"})
    assert status == 400
    assert "失效" in data["error"]


# ------------------------------------------------------------------ 思考档位
def test_think_switch_takes_effect_immediately(sample_xlsx, web):
    status, data = web.post("/api/think", {"value": "1"})
    assert status == 200
    assert data["thinking"] == "enabled"

    _, bootstrap = web.get("/api/bootstrap")
    assert bootstrap["thinking"] == "enabled"

    status, data = web.post("/api/think", {"value": "auto"})
    assert status == 200
    assert data["thinking"] == "auto"


def test_think_switch_rejects_invalid_value(web):
    status, data = web.post("/api/think", {"value": "sometimes"})
    assert status == 400
    assert "思考开关" in data["error"]


# ------------------------------------------------------------------ 打开原文件 / 提交表格
def test_open_file_uses_system_default_app(sample_xlsx, web, monkeypatch):
    opened: list[Path] = []
    monkeypatch.setattr(webapp, "_open_externally", opened.append)

    status, data = web.post("/api/open", {"file": "scores.xlsx"})
    assert status == 200
    assert data["file"] == "scores.xlsx"
    assert [p.name for p in opened] == ["scores.xlsx"]
    assert opened[0].is_file()

    # 文件不存在 / 越界路径 / 空选择：都拒绝，且不会再触发打开
    assert web.post("/api/open", {"file": "missing.xlsx"})[0] == 400
    assert web.post("/api/open", {"file": "../evil.xlsx"})[0] == 400
    status, data = web.post("/api/open", {"file": ""})
    assert status == 400 and "请先选择" in data["error"]
    assert len(opened) == 1


def test_upload_saves_file_and_lists_it(sample_xlsx, web):
    blob = sample_xlsx.read_bytes()
    status, data = web.post_raw("/api/upload?name=" + urllib.parse.quote("我的成绩.xlsx"), blob)
    assert status == 200
    assert data["file"] == "我的成绩.xlsx"
    assert data["renamed"] is False
    assert "我的成绩.xlsx" in data["files"]
    assert (sample_xlsx.parent / "我的成绩.xlsx").read_bytes() == blob

    # 上传的文件立即可用：直接走 /api/sheet 打开
    status, sheet = web.get("/api/sheet?file=" + urllib.parse.quote("我的成绩.xlsx"))
    assert status == 200
    assert sheet["current"] == "Sheet1"


def test_upload_renames_when_name_taken(sample_xlsx, web):
    blob = sample_xlsx.read_bytes()
    status, first = web.post_raw("/api/upload?name=mine.xlsx", blob)
    assert (status, first["file"], first["renamed"]) == (200, "mine.xlsx", False)

    status, second = web.post_raw("/api/upload?name=mine.xlsx", blob)
    assert status == 200
    assert second["file"] == "mine_1.xlsx"
    assert second["renamed"] is True


def test_upload_rejects_invalid_requests(sample_xlsx, web):
    blob = sample_xlsx.read_bytes()

    status, data = web.post_raw("/api/upload?name=notes.txt", blob)
    assert status == 400 and "只支持" in data["error"]

    status, data = web.post_raw("/api/upload?name=" + urllib.parse.quote("../evil.xlsx"), blob)
    assert status == 400 and "文件名无效" in data["error"]
    assert not (sample_xlsx.parent / "evil.xlsx").exists()

    status, data = web.post_raw("/api/upload?name=fake.xlsx", b"definitely not a workbook")
    assert status == 400 and "不是有效的 Excel 文件" in data["error"]
    assert not (sample_xlsx.parent / "fake.xlsx").exists()

    status, data = web.post_raw("/api/upload?name=empty.xlsx", b"")
    assert status == 400 and "为空" in data["error"]

    status, data = web.post_raw("/api/upload", blob)
    assert status == 400 and "文件名无效" in data["error"]


# ------------------------------------------------------------------ 新建表格格式
def test_generate_new_table_streams_table_format_and_applies_it(sample_xlsx, web):
    """需求识别为新建表格：阶段提示标明、预览带表格格式，写入回执列出已套用格式。"""
    web.client.replies.append({
        "formulas": [
            {"target": "H2", "value": "科目", "explanation": "表头"},
            {"target": "I2", "value": "总分", "explanation": "表头"},
            {"target": "H3", "value": "Math", "explanation": "科目名"},
            {"target": "I3", "formula": "=SUM(B2:F2)", "explanation": "各科总分", "fill_to": "I5"},
        ],
        "table": {"header": "H2:I2", "range": "H2:I5"},
        "assumptions": [],
        "clarification": None,
    })
    events = web.sse({"request": "新建一个表格统计各科总分", "file": "scores.xlsx"})

    stages = [e for e in events if e["type"] == "stage"]
    assert any("识别为新建表格" in e["text"] for e in stages)

    payload = events[-1]["payload"]
    assert payload["ok"] is True
    label = "Sheet1!H2:I5（表头 H2:I2 浅蓝底、范围加细边框）"
    assert payload["table"]["label"] == label

    status, data = web.post("/api/apply", {"proposal_id": payload["proposal_id"]})
    assert status == 200
    assert data["applied"]["tables"] == [label]


# ------------------------------------------------------------------ 下拉列表

def _create_dropdown(web, file="scores.xlsx", request="把 H2:H3 设置为下拉列表，选项为：001、002"):
    events = web.sse({"request": request, "file": file})
    proposal = events[-1]["payload"]
    assert proposal["ok"], proposal
    assert proposal["operation"] == "dropdown"
    assert proposal["usage"]["calls"] == 0
    status, applied = web.post("/api/apply", {"proposal_id": proposal["proposal_id"]})
    assert status == 200, applied
    return applied


def _dropdown_selection(data, cell="H2"):
    return {"file": data["file"], "sheet": data["current"], "cell": cell,
            "rule_id": data["grid"]["dropdowns"][0]["id"], "option_index": 1,
            "file_version": data["file_version"]}


def test_dropdown_preview_apply_and_select(sample_xlsx, web):
    before = sample_xlsx.read_bytes()
    events = web.sse({"request": "把 H2 设置为下拉列表，选项为：001、002", "file": "scores.xlsx"})
    proposal = events[-1]["payload"]
    assert proposal["ok"] and proposal["dropdown"]["options"] == ["001", "002"]
    assert sample_xlsx.read_bytes() == before
    status, applied = web.post("/api/apply", {"proposal_id": proposal["proposal_id"]})
    assert status == 200 and applied["applied"]["dropdown"]["range"] == "H2"
    assert applied["applied"]["backup_name"]
    assert web.post("/api/apply", {"proposal_id": proposal["proposal_id"]})[0] == 400
    _, data = web.get("/api/sheet?file=scores.xlsx")
    rule = data["grid"]["dropdowns"][0]
    assert rule["cells"]["H2"]["selected"] is None
    status, result = web.post("/api/dropdown/select", _dropdown_selection(data))
    assert status == 200, result
    assert result["applied"]["backup_name"] != applied["applied"]["backup_name"]
    _, fresh = web.get("/api/sheet?file=scores.xlsx")
    assert fresh["grid"]["dropdowns"][0]["cells"]["H2"]["selected"] == 1
    assert fresh["file_version"] != data["file_version"]
    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"]["H2"].value == "002"
    wb.close()
    assert web.client.calls == []


@pytest.mark.parametrize("change", [
    {"option_index": -1}, {"option_index": 100}, {"option_index": True}, {"option_index": "1"},
    {"rule_id": "伪造"}, {"cell": "A1"}, {"cell": "H2:H3"}, {"file_version": "过期"},
    {"file": "../../outside.xlsx"}, {"sheet": "不存在"},
])
def test_dropdown_selection_rejects_tampering(sample_xlsx, web, change):
    _create_dropdown(web)
    _, data = web.get("/api/sheet?file=scores.xlsx")
    before = sample_xlsx.read_bytes()
    status, result = web.post("/api/dropdown/select", {**_dropdown_selection(data), **change})
    assert status == 400 and result["error"]
    assert sample_xlsx.read_bytes() == before


def test_dropdown_source_formula_updates_after_selection(sample_xlsx, web):
    wb = load_workbook(sample_xlsx)
    wb["Sheet1"]["I2"] = "=H2*10"
    wb.save(sample_xlsx)
    wb.close()
    _create_dropdown(web, request="把 H2 设置为下拉列表，选项来自 Sheet1!B2:B3")
    _, data = web.get("/api/sheet?file=scores.xlsx")
    status, result = web.post("/api/dropdown/select", _dropdown_selection(data))
    assert status == 200, result
    _, fresh = web.get("/api/sheet?file=scores.xlsx")
    row = next(r for r in fresh["grid"]["rows"] if r["row"] == 2)
    assert row["values"][8] == "880"
    assert row["formulas"][8] == "=H2*10"


def test_dropdown_blank_area_long_options_and_unsupported_rule(sample_xlsx, web):
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = load_workbook(sample_xlsx)
    text = "这是一个必须完整保留不能被截断的下拉选项" * 2
    rule = DataValidation(type="list", formula1='"' + text + ',短选项"')
    rule.add("H50")
    wb["Sheet1"].add_data_validation(rule)
    unsupported = DataValidation(type="list", formula1='INDIRECT("A1:A2")')
    unsupported.add("H51")
    wb["Sheet1"].add_data_validation(unsupported)
    wb.save(sample_xlsx)
    wb.close()
    _, data = web.get("/api/sheet?file=scores.xlsx")
    rules = data["grid"]["dropdowns"]
    assert rules[0]["options"][0]["label"] == text
    assert "H50" in rules[0]["cells"]
    assert rules[1]["error"] and rules[1]["cells"]["H51"]["error"]


def test_dropdown_literal_equals_not_a_formula(sample_xlsx, web):
    _create_dropdown(web, request="把 H2 设置为下拉列表，选项为：普通、=1+1")
    _, data = web.get("/api/sheet?file=scores.xlsx")
    assert web.post("/api/dropdown/select", _dropdown_selection(data))[0] == 200
    _, fresh = web.get("/api/sheet?file=scores.xlsx")
    row = next(r for r in fresh["grid"]["rows"] if r["row"] == 2)
    assert row["values"][7] == "=1+1" and row["formulas"][7] == ""


# ------------------------------------------------------------------ 模拟运算表兼容
def test_display_value_falls_back_for_datatable():
    """模拟运算表单元格的值是 DataTableFormula 对象：摘要应退回缓存结果而不是对象文本。"""
    from openpyxl.worksheet.formula import DataTableFormula

    from excel_formula.excel_reader import _cell_text, _display_value

    wb_f = Workbook()
    ws_f = wb_f.active
    ws_f["D4"] = 0
    cell = ws_f["D4"]
    cell._value = DataTableFormula(ref="D4")  # 模拟 openpyxl 读入的对象型公式
    cell.data_type = "f"

    wb_v = Workbook()
    ws_v = wb_v.active
    ws_v["D4"] = 12904.55
    assert _display_value(ws_f, ws_v, 4, 4) == 12904.55
    assert "DataTableFormula" not in _cell_text(_display_value(ws_f, ws_v, 4, 4))

    wb_empty = Workbook()  # 无缓存（例如被 openpyxl 重新保存过）：退回空文本
    assert _display_value(ws_f, wb_empty.active, 4, 4) is None
    assert _cell_text(_display_value(ws_f, wb_empty.active, 4, 4)) == ""


def test_grid_tolerates_datatable_formula(tmp_path):
    """对象型公式不得进入网格 JSON（曾让 /api/sheet 整个 500），显示退回缓存结果。"""
    import types

    from openpyxl.worksheet.formula import DataTableFormula

    from excel_formula.excel_reader import SheetDigest

    path = tmp_path / "datatable.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "占位"
    workbook.save(path)
    workbook.close()

    ws_f = Workbook().active
    ws_f.title = "模拟运算表"
    ws_f["A1"] = "本金"
    ws_f["D4"] = 0
    cell = ws_f["D4"]
    cell._value = DataTableFormula(ref="D4")
    cell.data_type = "f"

    ws_v = Workbook().active
    ws_v.title = "模拟运算表"
    ws_v["D4"] = 12904.55

    # build_display_grid / dropdown_metadata 需要的最小视图接口
    view = types.SimpleNamespace(
        sheet_names=["模拟运算表"],
        wb_formulas={"模拟运算表": ws_f},
        wb_values={"模拟运算表": ws_v},
        path=path,
    )

    digest = SheetDigest(
        name="模拟运算表", max_row=4, max_column=4,
        rows=[(1, ["本金", "", "", ""]), (4, ["", "", "", "12904.55"])],
    )
    payload = webapp.build_display_grid(view, digest)
    json.dumps(payload, ensure_ascii=False)  # 修复前此处抛 TypeError

    row = next(r for r in payload["rows"] if r["row"] == 4)
    assert row["values"][3] == "12904.55"
    assert row["formulas"][3] == ""  # 对象型公式没有可展示的原文
