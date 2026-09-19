"""列表验证全本地回归：规则、类型、边界、确认版本与原子保存。"""
from __future__ import annotations

import datetime as dt
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from excel_formula import data_validation as dv
from excel_formula import writer
from excel_formula.excel_reader import WorkbookView
from excel_formula.intent import INTENT_DROPDOWN, INTENT_GENERATE, INTENT_VALIDATE, classify
from excel_formula.llm_client import ChatResult, Usage
from excel_formula.pipeline import FormulaService


class NoModel:
    """无兜底回复的模型替身：兜底调用会抛错，验证失败路径能静默回退。"""

    def chat(self, *args, **kwargs):
        raise AssertionError("下拉测试未预置模型回复")


class FallbackClient:
    """本地解析失败后的兜底替身：回放预置 JSON，并记录每次调用。"""

    def __init__(self, reply: dict | str):
        self.reply = reply if isinstance(reply, str) else json.dumps(reply, ensure_ascii=False)
        self.calls: list[list[dict]] = []

    def chat(self, messages, *, json_mode=True, max_tokens=800, on_delta=None, thinking=None):
        self.calls.append(messages)
        return ChatResult(
            content=self.reply, usage=Usage(prompt_tokens=480, completion_tokens=60, calls=1)
        )


@pytest.fixture()
def service(settings):
    return FormulaService(settings, client=NoModel())


@pytest.fixture()
def dropdown_book(tmp_path):
    wb = Workbook()
    query = wb.active
    query.title = "订单查询"
    query["B3"] = "SO-1"
    query["B4"] = '=IFERROR(VLOOKUP(B3,销售订单!A2:B3,2,0),"")'
    query["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    source = wb.create_sheet("销售订单")
    source.append(["订单号", "金额"])
    source.append(["SO-1", 100])
    source.append(["SO-2", 200])
    path = tmp_path / "dropdown.xlsx"
    wb.save(path)
    wb.close()
    return path


def propose(service, path, text="把订单查询!B3 设置为下拉列表，选项来自销售订单!A2:A3"):
    return service.propose_dropdown(path, text, sheet="订单查询")


def select_args(path, cell="B3"):
    with WorkbookView(path) as view:
        rule = dv.dropdown_metadata(view, "订单查询")[0]
    return {"sheet": "订单查询", "cell": cell, "rule_id": rule["id"],
            "option_index": 1, "file_version": dv.file_version(path)}


@pytest.mark.parametrize("text", [
    "把 B3 设置为下拉列表，选项为：甲、乙",
    "把 B3 设成下拉框，选项为：甲，乙",
    "给 B3 添加数据验证，选项为：甲,乙",
    "把 B3 设置为下拉列表，内容为：甲、乙",
    "把 B3 设置为下拉列表，值为：甲、乙",
])
def test_intent_and_inline(service, dropdown_book, text):
    assert classify(text).kind == INTENT_DROPDOWN
    before = dv.file_version(dropdown_book)
    p = propose(service, dropdown_book, text)
    assert p.ok, p.error or p.clarification
    assert p.dropdown.options == ["甲", "乙"]
    assert p.usage.calls == 0
    assert dv.file_version(dropdown_book) == before
    assert "不在选项中" in p.render()
    result = service.apply(p)
    assert Path(result["backup"]).exists()
    wb = load_workbook(dropdown_book)
    ws = wb["订单查询"]
    rule = ws.data_validations.dataValidation[0]
    assert rule.type == "list" and rule.showDropDown is False
    assert rule.showErrorMessage and rule.errorStyle == "stop" and rule.allowBlank
    assert rule.formula1 == '"甲,乙"' and str(rule.sqref) == "B3"
    assert ws["B3"].value == "SO-1" and ws["B3"].fill.fgColor.rgb == "00FFF2CC"
    assert ws["B4"].data_type == "f"
    wb.close()


def test_formula_intents_unchanged():
    assert classify("在B3生成公式并下拉到B9").kind == INTENT_GENERATE
    assert classify("=SUM(A1:A3)这个公式对不对").kind == INTENT_VALIDATE
    assert classify('检查 =IF(A1=1,"下拉列表","数据验证")').kind == INTENT_VALIDATE
    assert classify('在B3写入公式 =IF(A1=1,"下拉列表","")').kind == INTENT_GENERATE


@pytest.mark.parametrize("text", [
    "不要把B3设置为下拉列表，选项为：甲、乙",
    "删除B3的下拉列表",
    "解释B3的数据验证",
    "如何设置B3的下拉列表，选项为：甲、乙",
    "把B3设为下拉列表并计算公式，选项为：甲、乙",
    "把B3和C3设置为下拉列表，选项为：甲、乙",
    "把B3设为下拉列表，选项为：甲、乙，并删除C3",
    "把B3设为下拉列表，选项为：甲、乙，同时计算C3",
    "把B3设为下拉列表，选项为：甲、乙，然后把C3加边框",
    "把B3设为下拉列表，选项为：甲、乙，写入C3",
])
def test_non_creation_never_writes(service, dropdown_book, text):
    before = dv.file_version(dropdown_book)
    p = propose(service, dropdown_book, text)
    assert not p.ok
    with pytest.raises(ValueError):
        service.apply(p)
    assert dv.file_version(dropdown_book) == before


def test_option_words_do_not_become_commands_or_sources(service, dropdown_book):
    p = propose(service, dropdown_book, "把B3设为下拉列表，选项为：来自北京、取消、删除")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.options == ["来自北京", "取消", "删除"]


@pytest.mark.parametrize("text", [
    "把订单查询!B3 设置为下拉列表，选项来自销售订单!A2:A3",
    "把订单查询!B3 设置为下拉列表，选项是来自销售订单!A2:A3",
    "把订单查询!B3 设置为下拉列表，选项是销售订单!A2:A3",
    "把订单查询!B3设置为下拉列表，选项来自销售订单表!A2:A3",
    "给订单查询!B3加下拉列表，选项来自销售订单的A2到A3",
    "订单查询!B3加下拉，来源销售订单!$A$2:$A$3",
    "帮我给订单查询!B3做一个下拉列表，引用销售订单!A2:A3",
    "订单查询B3下拉列表，选项来自销售订单!A2:A3",
    "给订单查询!B3加数据有效性，列表来源销售订单!A2:A3",
    "把订单查询!B3 设置为下拉列表，选项来自销售订单!A2:A3，以便B4查询",
    "把订单查询!B3 设置为下拉列表，选项是来自销售订单!A2:A3，以便B4查询",
    "把订单查询!B3 设置为下拉列表，选项来自 销售订单 的 A2 到 A3",
])
def test_colloquial_source_variants(service, dropdown_book, text):
    """口语来源写法（选项是来自/表名带表字/无感叹号/到字区间/尾注）都要解析到位。"""
    p = propose(service, dropdown_book, text)
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "订单查询" and p.dropdown.cell_range == "B3"
    assert p.dropdown.formula1 == "'销售订单'!$A$2:$A$3"
    assert p.dropdown.options == ["SO-1", "SO-2"]


def test_colloquial_variant_applies(service, dropdown_book):
    """口语写法走完整链路：apply 后静态名称与规则都落在源表上。"""
    p = propose(service, dropdown_book, "把订单查询!B3 设置为下拉列表，选项是来自销售订单!A2:A3")
    assert p.ok, p.error
    result = service.apply(p)
    assert Path(result["backup"]).exists()
    wb = load_workbook(dropdown_book)
    rule = wb["订单查询"].data_validations.dataValidation[0]
    assert rule.formula1.startswith("=_ExcelCR_DV_")
    assert wb.defined_names[rule.formula1[1:]].attr_text == "'销售订单'!$A$2:$A$3"
    wb.close()


def test_source_lead_is_not_swallowed_by_options(service, dropdown_book):
    """“选项是来自北京”是字面选项，不能因为“来自”把北京当成来源。"""
    p = propose(service, dropdown_book, "把B3设为下拉列表，选项是来自北京")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.options == ["来自北京"]
    p = propose(service, dropdown_book, "把B3设为下拉列表，选项是来自北京、上海")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.options == ["来自北京", "上海"]


def test_inline_area_reference(service, dropdown_book):
    """“选项是C2:C3”指向本表区域，而不是逐字选项。"""
    wb = load_workbook(dropdown_book)
    wb["订单查询"]["C2"] = "甲"
    wb["订单查询"]["C3"] = "乙"
    wb.save(dropdown_book)
    wb.close()
    p = propose(service, dropdown_book, "把B3设为下拉列表，选项是C2:C3")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.formula1 == "'订单查询'!$C$2:$C$3"
    assert p.dropdown.options == ["甲", "乙"]


@pytest.mark.parametrize("source", ["销售订单!A列", "销售订单!A:A", "1:1"])
def test_whole_column_source_hint(service, dropdown_book, source):
    """整列/整行引用给专门的提示，而不是笼统的“不支持”。"""
    p = propose(service, dropdown_book, f"把B3设为下拉列表，选项来自{source}")
    assert not p.ok and "整列" in (p.error or "")


@pytest.mark.parametrize("text", [
    "将b3设置为下拉列表，数据来自于销售订单的订单号",
    "把订单查询!B3 设置为下拉列表，数据来自于销售订单的订单号",
    "把订单查询!B3 设置为下拉列表，数据来源是销售订单的订单号",
    "把订单查询!B3 设置为下拉列表，数据取自销售订单的订单号",
    "把订单查询!B3 设置为下拉列表，选项是来自销售订单的订单号",
    "把订单查询!B3 设置为下拉列表，选项是销售订单的订单号",
    "把订单查询!B3 设置为下拉列表，来自销售订单的订单号列",
    "订单查询!B3加下拉，来源销售订单表!订单号",
    "把订单查询!B3 设置为下拉列表，数据来自于销售订单的订单编号",
    "把b3改为下拉列表，内容为销售订单中的订单编号",
    "把订单查询!B3 设置为下拉列表，内容是销售订单的订单号",
    "把订单查询!B3 设置为下拉列表，值为销售订单的订单号",
])
def test_header_column_source(service, dropdown_book, text):
    """来源写成列标题（含“订单编号→订单号”这类口语近似）时按标题列数据区解析。"""
    p = propose(service, dropdown_book, text)
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "订单查询" and p.dropdown.cell_range == "B3"
    assert p.dropdown.formula1 == "'销售订单'!$A$2:$A$3"
    assert p.dropdown.options == ["SO-1", "SO-2"]


def test_header_column_without_data_keeps_hint(service, dropdown_book):
    """列标题存在但下方没有数据时仍走提示，而不是给出空来源。"""
    wb = load_workbook(dropdown_book)
    wb["销售订单"]["C1"] = "备注"
    wb.save(dropdown_book)
    wb.close()
    p = propose(service, dropdown_book, "把B3设为下拉列表，来源销售订单的备注")
    assert not p.ok and "无法识别下拉来源" in (p.error or "")


def test_source_mention_does_not_move_target_sheet(service, dropdown_book):
    """句里顺带提到的来源表名不能被当成目标表：未指定目标表时改的永远是当前表。"""
    p = propose(service, dropdown_book, "用销售订单的订单编号给B3设置下拉，选项为：甲、乙")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "订单查询" and p.dropdown.cell_range == "B3"
    # 表名与目标格绑定（“销售订单的B3”）时才算明确指定
    p = propose(service, dropdown_book, "把销售订单的B3设置为下拉列表，选项为：甲、乙")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "销售订单"
    p = propose(service, dropdown_book, "在订单查询表里把B3设置为下拉列表，选项为：甲、乙")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "订单查询"


# ---------------------------------------------------- 模型兜底（本地未命中时一次调用）
def test_fallback_understands_free_phrasing(settings, dropdown_book):
    """本地认不出的口语来源（“选项用…”）由模型兜底理解，输出仍走严格校验。"""
    assert classify("把B3改为下拉列表，选项用销售订单表里的订单号").kind == INTENT_DROPDOWN
    client = FallbackClient({
        "intent": "create", "sheet": "订单查询", "target": "B3",
        "source": {"sheet": "销售订单", "header": "订单号"},
    })
    service = FormulaService(settings, client=client)
    stages: list[str] = []
    p = service.propose_dropdown(
        dropdown_book, "把B3改为下拉列表，选项用销售订单表里的订单号",
        sheet="订单查询", on_stage=stages.append,
    )
    assert p.ok, p.error or p.clarification
    assert p.dropdown.formula1 == "'销售订单'!$A$2:$A$3"
    assert p.dropdown.options == ["SO-1", "SO-2"]
    assert p.usage.calls == 1 and len(client.calls) == 1
    assert any("1 次调用" in note for note in stages)
    user_text = client.calls[0][1]["content"]
    assert "订单查询" in user_text and "销售订单" in user_text  # 上下文带上了两张表的内容


def test_fallback_accepts_region_source(settings, dropdown_book):
    """句尾“…做选项”触发缺选项追问时，模型给出区域来源同样能兜底成功。"""
    client = FallbackClient({
        "intent": "create", "sheet": "订单查询", "target": "B3",
        "source": {"sheet": "销售订单", "range": "A2:A3"},
    })
    service = FormulaService(settings, client=client)
    stages: list[str] = []
    p = service.propose_dropdown(
        dropdown_book, "把B3改为下拉列表，用销售订单表里的订单号做选项", on_stage=stages.append,
    )
    assert p.ok, p.error or p.clarification
    assert p.dropdown.options == ["SO-1", "SO-2"]
    assert p.usage.calls == 1
    assert "1 次调用" in p.render()  # 预览如实标注这次消耗，不再写“0 Token”
    assert any("1 次调用" in note for note in stages)
    assert not any("0 Token" in note for note in stages)  # 兜底发生时不得再宣传 0 Token


def test_fallback_rejects_hallucinated_source(settings, dropdown_book):
    """模型编造不存在的来源表时必须退回本地提示，不产生方案。"""
    client = FallbackClient({
        "intent": "create", "sheet": "订单查询", "target": "B3",
        "source": {"sheet": "不存在的表", "range": "A2:A9"},
    })
    service = FormulaService(settings, client=client)
    p = service.propose_dropdown(dropdown_book, "把B3改为下拉列表，选项用订单数据里的编号")
    assert not p.ok and "无法识别下拉来源" in (p.error or "")
    assert p.dropdown is None
    assert p.usage.calls == 1  # 白跑的一次也要如实计入


def test_fallback_rejects_non_create_intent(settings, dropdown_book):
    """模型判定为删除/询问等非创建意图时，兜底放弃、保留原提示。"""
    client = FallbackClient({"intent": "remove", "clarification": "本系统只支持创建下拉列表"})
    service = FormulaService(settings, client=client)
    p = service.propose_dropdown(dropdown_book, "把B3改为下拉列表，选项用订单数据里的编号")
    assert not p.ok and p.error and p.dropdown is None
    assert len(client.calls) == 1


def test_fallback_survives_model_failure(service, dropdown_book):
    """模型不可用（如未配密钥）时静默回退，用户看到的仍是本地提示。"""
    p = propose(service, dropdown_book, "把B3改为下拉列表，选项用订单数据里的编号")
    assert not p.ok and p.error and p.dropdown is None
    assert p.usage.calls == 0


def test_local_hit_never_calls_model(settings, dropdown_book):
    """本地规则命中时一次模型调用都不能发生（0 Token 卖点）。"""
    client = FallbackClient({})
    service = FormulaService(settings, client=client)
    stages: list[str] = []
    p = service.propose_dropdown(
        dropdown_book, "把B3设置为下拉列表，选项为：甲、乙", sheet="订单查询", on_stage=stages.append,
    )
    assert p.ok and p.usage.calls == 0
    assert client.calls == []
    assert any("0 Token" in note for note in stages)


def test_fallback_keeps_current_sheet_when_model_confuses_source(dropdown_book, settings):
    """兜底时模型把来源表填进 sheet 字段：没有绑定到目标格就拉回当前表。"""
    client = FallbackClient({
        "intent": "create", "sheet": "销售订单", "target": "B3",
        "source": {"sheet": "销售订单", "header": "订单号"},
    })
    service = FormulaService(settings, client=client)
    p = service.propose_dropdown(
        dropdown_book, "把B3改为下拉列表，选项用销售订单表里的订单号", sheet="订单查询",
    )
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "订单查询"
    assert p.dropdown.formula1 == "'销售订单'!$A$2:$A$3"
    assert p.usage.calls == 1


def test_fallback_target_sheet_rules(dropdown_book, settings):
    """兜底目标表规则：句里与目标格绑定的表名照办；凭空出现的表名不算数。"""
    client = FallbackClient({
        "intent": "create", "target": "B3",
        "source": {"sheet": "销售订单", "header": "订单号"},
    })
    service = FormulaService(settings, client=client)
    p = service.propose_dropdown(
        dropdown_book, "在销售订单的B3做一个下拉列表，数据从订单号来", sheet="订单查询",
    )
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "销售订单"  # “销售订单的B3”是明确指定

    wb = load_workbook(dropdown_book)
    wb.create_sheet("台账")
    wb.save(dropdown_book)
    wb.close()
    client = FallbackClient({
        "intent": "create", "sheet": "台账", "target": "B3",
        "source": {"sheet": "销售订单", "header": "订单号"},
    })
    service = FormulaService(settings, client=client)
    p = service.propose_dropdown(
        dropdown_book, "把B3改为下拉列表，选项用销售订单表里的订单号", sheet="订单查询",
    )
    assert p.ok, p.error or p.clarification
    assert p.dropdown.sheet == "订单查询"  # 句中没提“台账”：不写到第三张表上


def test_clarification_and_rectangular_range(service, dropdown_book):
    p = propose(service, dropdown_book, "设置下拉列表，选项为：001、002")
    assert p.clarification
    p = propose(service, dropdown_book, "设置下拉列表，选项为：001、002（补充：B3到C5）")
    assert p.ok, p.error or p.clarification
    assert p.dropdown.cell_range == "B3:C5"
    assert p.dropdown.options == ["001", "002"]


def test_cross_sheet_selection_and_preservation(service, dropdown_book):
    p = propose(service, dropdown_book)
    assert p.ok, p.error
    assert p.sheet == "订单查询"
    result = service.apply(p)
    args = select_args(dropdown_book)
    result2 = service.select_dropdown(dropdown_book, **args)
    assert result["backup"] != result2["backup"]
    wb = load_workbook(dropdown_book)
    ws = wb["订单查询"]
    rule = ws.data_validations.dataValidation[0]
    assert rule.formula1.startswith("=_ExcelCR_DV_")
    assert wb.defined_names[rule.formula1[1:]].attr_text == "'销售订单'!$A$2:$A$3"
    assert ws["B3"].value == "SO-2"
    assert ws["B4"].data_type == "f"
    wb.close()
    p2 = propose(service, dropdown_book)
    assert p2.dropdown.unchanged
    before = dv.file_version(dropdown_book)
    assert service.apply(p2)["unchanged"]
    assert dv.file_version(dropdown_book) == before


@pytest.mark.parametrize("values", [[1, 2.5], [False, True], ["001", "002"],
    [dt.datetime(2026, 1, 1), dt.datetime(2026, 2, 2)], ["普通", "=1+1"], ["普通", "#N/A"]])
def test_selection_keeps_source_types(service, dropdown_book, values):
    wb = load_workbook(dropdown_book)
    for row, value in enumerate(values, 2):
        cell = wb["销售订单"].cell(row, 1, value)
        if isinstance(value, str):
            cell.data_type = "s"
    wb.save(dropdown_book)
    wb.close()
    service.apply(propose(service, dropdown_book))
    service.select_dropdown(dropdown_book, **select_args(dropdown_book))
    wb = load_workbook(dropdown_book)
    cell = wb["订单查询"]["B3"]
    assert type(cell.value) is type(values[1]) and cell.value == values[1]
    if isinstance(values[1], str):
        assert cell.data_type == "s"
    wb.close()


@pytest.mark.parametrize("source,expected", [
    ("销售订单!A2:B3", "单行或单列"), ("销售订单!A1:A1001", "1000"),
    ("销售订单!Z2:Z3", "为空"), ("不存在!A1:A2", "不存在"),
    ('INDIRECT("A1:A2")', "动态"), ("[other.xlsx]Sheet1!A1:A2", "外部"),
    ("销售订单!XFE1", "超出"),
])
def test_invalid_sources(service, dropdown_book, source, expected):
    p = propose(service, dropdown_book, f"把B3设置为下拉列表，选项来自{source}")
    assert not p.ok and expected in (p.error or "")


def test_target_limit_and_inline_limit(service, dropdown_book):
    assert not propose(service, dropdown_book, "把B1:B201设为下拉列表，选项为：甲、乙").ok
    assert not propose(service, dropdown_book, "把B3设为下拉列表，选项为：" + "甲" * 256).ok


def test_protected_merged_and_overlapping_targets(service, dropdown_book):
    wb = load_workbook(dropdown_book)
    ws = wb["订单查询"]
    ws.merge_cells("B3:C3")
    wb.save(dropdown_book)
    wb.close()
    assert "合并" in propose(service, dropdown_book).error
    wb = load_workbook(dropdown_book)
    ws = wb["订单查询"]
    ws.unmerge_cells("B3:C3")
    ws.protection.sheet = True
    wb.save(dropdown_book)
    wb.close()
    assert "保护" in propose(service, dropdown_book).error
    wb = load_workbook(dropdown_book)
    ws = wb["订单查询"]
    ws.protection.sheet = False
    rule = DataValidation(type="list", formula1='"甲,乙"')
    rule.add("B3:B4")
    ws.add_data_validation(rule)
    wb.save(dropdown_book)
    wb.close()
    assert "冲突" in propose(service, dropdown_book).error


def test_quoted_source_name_and_local_defined_name(service, dropdown_book):
    wb = load_workbook(dropdown_book)
    wb["销售订单"].title = "O'Brien 数据"
    wb["订单查询"].defined_names.add(DefinedName("订单源", attr_text="'O''Brien 数据'!$A$2:$A$3"))
    wb.save(dropdown_book)
    wb.close()
    p = propose(service, dropdown_book, "把订单查询!B3设为下拉列表，选项来自订单源")
    assert p.ok, p.error
    service.apply(p)
    with WorkbookView(dropdown_book) as view:
        assert dv.dropdown_metadata(view, "订单查询")[0]["options"][1]["label"] == "SO-2"


def test_source_formula_without_cache(service, dropdown_book):
    wb = load_workbook(dropdown_book)
    wb["销售订单"]["A2"] = "=1+2"
    wb["销售订单"]["A3"] = "=A2*2"
    wb.save(dropdown_book)
    wb.close()
    p = propose(service, dropdown_book)
    assert p.ok and p.dropdown.options == [3, 6]


def test_stale_creation_and_selection_are_rejected(service, dropdown_book):
    p = propose(service, dropdown_book)
    wb = load_workbook(dropdown_book)
    wb["订单查询"]["C1"] = "外部修改"
    wb.save(dropdown_book)
    wb.close()
    with pytest.raises(ValueError, match="版本|改变"):
        service.apply(p)
    service.apply(propose(service, dropdown_book))
    args = select_args(dropdown_book)
    service.select_dropdown(dropdown_book, **args)
    with pytest.raises(ValueError, match="版本|改变"):
        service.select_dropdown(dropdown_book, **args)


def test_concurrent_selection_only_one_version_succeeds(service, dropdown_book):
    service.apply(propose(service, dropdown_book))
    args = select_args(dropdown_book)
    def run():
        try:
            service.select_dropdown(dropdown_book, **args)
            return True
        except ValueError:
            return False
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda _: run(), range(2)))
    assert sorted(results) == [False, True]


def test_save_failure_keeps_original(service, dropdown_book, monkeypatch):
    p = propose(service, dropdown_book)
    before = dropdown_book.read_bytes()
    def fail(*args, **kwargs):
        raise PermissionError("文件占用")
    monkeypatch.setattr(writer.os, "replace", fail)
    with pytest.raises(PermissionError):
        service.apply(p)
    assert dropdown_book.read_bytes() == before
    assert not list(dropdown_book.parent.glob(".excelcr-*"))


def test_x14_blocks_creation_without_changing_file(service, dropdown_book):
    with zipfile.ZipFile(dropdown_book) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
        b"</worksheet>", b'<extLst><ext uri="test"><x:dataValidations xmlns:x="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"/></ext></extLst></worksheet>')
    with zipfile.ZipFile(dropdown_book, "w") as output:
        for name, data in entries.items():
            output.writestr(name, data)
    before = dropdown_book.read_bytes()
    assert "x14" in propose(service, dropdown_book).error
    assert dropdown_book.read_bytes() == before


def test_xml_dtd_rejected(tmp_path):
    path = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(path, "w") as output:
        output.writestr("xl/worksheets/sheet1.xml", '<!DOCTYPE a [<!ENTITY b "boom">]><a>&b;</a>')
    with pytest.raises(ValueError, match="DTD"):
        dv.extension_warning(path)


def test_cli_dropdown_preview_and_apply(service, dropdown_book, capsys):
    import json
    from excel_formula import cli
    args = cli.build_parser().parse_args([
        "--json", "nl", "把B3设为下拉列表，选项为：001、002",
        "--file", str(dropdown_book), "--sheet", "订单查询",
    ])
    before = dropdown_book.read_bytes()
    assert cli._cmd_nl(service, args) == cli.EXIT_OK
    assert json.loads(capsys.readouterr().out)["dropdown"]["range"] == "B3"
    assert dropdown_book.read_bytes() == before
    args.apply = args.yes = True
    assert cli._cmd_nl(service, args) == cli.EXIT_OK
    payload = json.loads(capsys.readouterr().out)
    assert payload["applied"]["dropdown"] and payload["applied"]["backup"]


def test_console_dropdown_clarification_and_cancel(service, dropdown_book, monkeypatch, capsys):
    from main import Console
    console = Console.__new__(Console)
    console.service = service
    before = dropdown_book.read_bytes()
    answers = iter(["选项为：甲、乙", "n"])
    monkeypatch.setattr("builtins.input", lambda prompt: next(answers))
    console.do_dropdown(dropdown_book, "把B3设为下拉列表", "订单查询")
    assert "已取消" in capsys.readouterr().out
    assert dropdown_book.read_bytes() == before


def test_current_web_script_syntax():
    """有 Node.js 时直接检查当前网页内联脚本，不依赖日志里的旧副本。"""
    import re
    import shutil
    import subprocess
    node = shutil.which("node")
    if not node:
        pytest.skip("当前 PATH 未找到 Node.js，跳过独立脚本语法检查")
    html = (Path(__file__).resolve().parents[1] / "web" / "index.html").read_text(encoding="utf-8")
    scripts = re.findall(r"<script[^>]*>(.*?)</script>", html, re.S | re.I)
    assert scripts
    for script in scripts:
        result = subprocess.run([node, "--check", "-"], input=script, text=True,
                                encoding="utf-8", capture_output=True, timeout=15)
        assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("mode", ["cancel", "failure", "confirm", "saving", "loading", "switched", "switch_during_save"])
def test_web_dropdown_event_handler(mode):
    """隔离执行真实事件处理代码；模拟确认结果，不替代浏览器原生对话框验收。"""
    import json
    import shutil
    import subprocess
    node = shutil.which("node")
    if not node:
        pytest.skip("当前 PATH 未找到 Node.js，跳过前端事件单元测试")
    html = (Path(__file__).resolve().parents[1] / "web" / "index.html").read_text(encoding="utf-8")
    fragment = html[html.index("let lastGridData = null;"):html.index("function rowHtml(")]
    setup = r'''
const assert = require('node:assert/strict');
const handlers = {}, messages = [], posts = [], loads = [];
const wrap = { addEventListener: (name, fn) => { handlers[name] = fn; } };
const side = {};
const $ = (selector) => selector === '#grid-wrap' ? wrap : side;
const esc = (value) => String(value ?? '');
const state = { file: 'test.xlsx', sheet: 'query', saving: mode === 'saving', sheetLoading: mode === 'loading' };
const window = { confirm: () => mode !== 'cancel' };
const toast = (message) => messages.push(message);
const cellMarks = () => {};
const loadSheet = async (sheet) => { loads.push(sheet); };
const postJSON = async (url, body) => {
  posts.push({ url, body });
  if (mode === 'failure') throw new Error('保存失败');
  if (mode === 'switch_during_save') state.sheet = 'other';
  return { applied: { backup_name: 'backup.xlsx' } };
};
const Option = class {
  constructor(text, value, defaultSelected, selected) {
    Object.assign(this, { text, value, defaultSelected, selected });
  }
};
'''
    assertions = r'''
(async () => {
  const entry = { rule: { id: 'rule-id', options: [{ label: '001' }, { label: '002' }] },
                  current: { selected: null, current: '' } };
  dropdownCells.set('B3', entry);
  assert.match(dropdownHtml('B3'), /value="" selected disabled/);
  const select = { dataset: { cell: 'B3', original: '' }, value: '', disabled: false,
                   options: [], add(option) { this.options.push(option); }, closest() { return this; } };
  handlers.focusin({ target: select });
  assert.equal(select.value, '');
  assert.equal(select.options[0].disabled, true);
  assert.deepEqual(select.options.slice(1).map(o => o.text), ['001', '002']);
  assert.equal(posts.length, 0);
  entry.current = { selected: 0, current: '001' };
  select.dataset.original = '0';
  select.value = '1';
  lastGridData = { file: state.file, sheet: state.sheet, file_version: 'original-version' };
  if (mode === 'switched') state.sheet = 'other';
  await handlers.change({ target: select });
  if (['cancel', 'saving', 'loading', 'switched'].includes(mode)) {
    assert.equal(posts.length, 0);
    assert.equal(select.value, '0');
    assert.equal(loads.length, 0);
  } else {
    assert.equal(posts.length, 1);
    assert.equal(posts[0].url, '/api/dropdown/select');
    assert.deepEqual(posts[0].body, { file: 'test.xlsx', sheet: 'query', cell: 'B3',
      rule_id: 'rule-id', option_index: 1, file_version: 'original-version' });
    assert.equal(state.saving, false);
    assert.equal(select.disabled, false);
    if (mode === 'failure') {
      assert.equal(select.value, '0');
      assert.equal(loads.length, 0);
      assert.ok(messages.some(message => message.includes('保存失败')));
    } else {
      assert.match(side.textContent, /backup.xlsx/);
      assert.deepEqual(loads, mode === 'switch_during_save' ? [] : ['query']);
    }
  }
})().catch(error => { console.error(error); process.exitCode = 1; });
'''
    result = subprocess.run([node, "-"], input="const mode = " + json.dumps(mode) + ";\n" + setup + fragment + assertions,
                            text=True, encoding="utf-8", capture_output=True, timeout=15)
    assert result.returncode == 0, result.stderr


def run_browser_demo():
    """人工浏览器验收入口：只操作缓存目录的新建工作簿，不接触用户业务表。"""
    import tempfile
    import webapp
    from excel_formula.config import Settings
    root = Path(__file__).resolve().parents[1]
    directory = Path(tempfile.mkdtemp(prefix="dropdown-browser-", dir=root / ".pytest_cache"))
    wb = Workbook()
    ws = wb.active
    ws.title = "订单查询"
    ws.append(["订单编号", "客户", "金额"])
    ws["A3"] = "订单号"
    ws["B3"] = "SO-1"
    ws["A4"] = "客户"
    ws["B4"] = '=IFERROR(VLOOKUP(B3,销售订单!A2:C3,2,0),"")'
    ws["A5"] = "金额"
    ws["B5"] = '=IFERROR(VLOOKUP(B3,销售订单!A2:C3,3,0),"")'
    source = wb.create_sheet("销售订单")
    source.append(["订单号", "客户", "金额"])
    source.append(["SO-1", "甲客户", 100])
    source.append(["SO-2", "乙客户", 200])
    wb.save(directory / "下拉验收.xlsx")
    wb.close()
    service = FormulaService(Settings(allowed_roots=[directory]), client=NoModel())
    server = webapp.make_server(service, port=8769)
    print(f"浏览器验收：http://127.0.0.1:{server.server_address[1]} 文件目录：{directory}", flush=True)
    try:
        server.serve_forever()
    finally:
        server.server_close()
