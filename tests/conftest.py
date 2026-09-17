"""pytest 公共夹具：构造与 测试数据.xlsx / Excel大作业.xlsm 结构一致的临时工作簿。"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from excel_formula.config import Settings  # noqa: E402

ROWS = [
    ["Subject", "Student1", "Student2", "Student3", "Student4", "Student5"],
    ["Math", 85, 78, 92, 88, 95],
    ["English", 88, 85, 91, 82, 90],
    ["Physics", 80, 82, 89, 75, 88],
    ["Chemistry", 90, 87, 84, 91, 86],
    ["Biology", 78, 80, 86, 83, 81],
]

# 订单表取自 Excel大作业.xlsm 的前 8 条真实记录（订单日期为真正的日期类型）
SALES_HEADER = [
    "订单编号", "订单日期", "客户名称", "所属大区", "产品类别", "产品名称",
    "销售负责人", "销售数量", "单价(元)", "总销售额(元)", "总成本(元)", "利润(元)",
]
SALES_ROWS = [
    ["ORD-001", datetime(2025, 1, 5), "王明", "华北", "电子产品", "笔记本电脑", "张伟", 5, 5500, 27500, 20000, 7500],
    ["ORD-002", datetime(2025, 1, 7), "李红", "华东", "家具", "办公桌", "王芳", 10, 1200, 12000, 9000, 3000],
    ["ORD-003", datetime(2025, 1, 10), "赵丽", "华南", "服装", "T恤", "李娜", 50, 80, 4000, 3000, 1000],
    ["ORD-004", datetime(2025, 1, 12), "陈明", "西南", "办公用品", "打印机", "刘洋", 3, 1500, 4500, 3200, 1300],
    ["ORD-005", datetime(2025, 1, 15), "刘刚", "西北", "食品", "坚果礼盒", "陈静", 20, 200, 4000, 2800, 1200],
    ["ORD-006", datetime(2025, 1, 18), "张丽", "华北", "电子产品", "智能手机", "张伟", 8, 3000, 24000, 18000, 6000],
    ["ORD-007", datetime(2025, 1, 20), "王强", "华东", "家具", "椅子", "王芳", 15, 500, 7500, 5500, 2000],
    ["ORD-008", datetime(2025, 2, 1), "李梅", "华南", "服装", "连衣裙", "李娜", 30, 150, 4500, 3200, 1300],
]


@pytest.fixture()
def sample_xlsx(tmp_path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in ROWS:
        sheet.append(row)
    workbook.create_sheet("Sheet2")
    path = tmp_path / "scores.xlsx"
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture()
def sales_xlsx(tmp_path: Path) -> Path:
    """复刻 Excel大作业.xlsm 的核心形态：中文表名、日期列、跨表查询区。"""
    workbook = Workbook()
    raw = workbook.active
    raw.title = "原始数据表"
    raw.append(SALES_HEADER)
    for row in SALES_ROWS:
        raw.append(row)

    query = workbook.create_sheet("查询表")
    query.append(["销售订单查询系统", None, None, "负责销售人", "总销售额"])
    query.append([None, None, None, "张伟", None])
    query.append(["请选择订单编号：", "ORD-003"])
    query.append([None, None, None, "李娜", None])
    query.append(["客户名称", None])
    query.append(["销售负责人", None])
    query.append(["总销售额（元）", None])

    path = tmp_path / "sales.xlsx"
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture()
def settings(tmp_path: Path) -> Settings:
    return Settings(
        api_key="sk-test-key-not-real",
        allowed_roots=[tmp_path.resolve()],
        log_dir=tmp_path / "logs",
        max_repair_rounds=2,
    )
