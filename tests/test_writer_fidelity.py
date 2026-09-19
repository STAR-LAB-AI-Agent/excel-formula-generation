"""写入保真：图片/图表/透视表部件在保存前后的检测、保留与丢失警告。"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XlImage

from excel_formula import writer
from excel_formula.writer import CellWrite, apply_writes, inspect_rich_objects


@pytest.fixture()
def chart_xlsx(tmp_path: Path) -> Path:
    """含一个柱状图的工作簿。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Name", "Value"])
    sheet.append(["A", 1])
    sheet.append(["B", 2])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    sheet.add_chart(chart, "D2")
    path = tmp_path / "chart.xlsx"
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture()
def rich_xlsx(tmp_path: Path) -> Path:
    """含图片与图表的工作簿；图片生成依赖 Pillow，缺失时跳过。"""
    pil_image = pytest.importorskip("PIL.Image")
    png = tmp_path / "logo.png"
    pil_image.new("RGB", (60, 60), (200, 30, 30)).save(png)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Name", "Value"])
    sheet.append(["A", 1])
    sheet.add_image(XlImage(str(png)), "D2")
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "G2")
    path = tmp_path / "rich.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def _write(settings, path: Path, cell: str = "C2"):
    return apply_writes(
        path, [CellWrite(sheet="Sheet1", cell=cell, formula="=B2*2")], settings
    )


# ---------------------------------------------------------------- 部件扫描
def test_plain_workbook_has_no_rich_objects(sample_xlsx):
    assert inspect_rich_objects(sample_xlsx) == {
        "charts": 0, "images": 0, "pivots": 0, "drawings": 0, "wmf_images": 0,
    }


def test_inspect_counts_pivot_and_wmf_parts(tmp_path: Path):
    """部件计数只看 zip 结构，不要求文件是合法工作簿；WMF 单独计数。"""
    path = tmp_path / "parts.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("xl/pivotTables/pivotTable1.xml", "<x/>")
        archive.writestr("xl/pivotTables/pivotTable2.xml", "<x/>")
        archive.writestr("xl/media/image1.png", b"x")
        archive.writestr("xl/media/clip.wmf", b"x")
    counts = inspect_rich_objects(path)
    assert counts["pivots"] == 2
    assert counts["images"] == 1
    assert counts["wmf_images"] == 1


# ---------------------------------------------------------------- 保留
def test_write_preserves_chart(chart_xlsx, settings):
    result = _write(settings, chart_xlsx)
    fidelity = result["fidelity"]
    assert fidelity["after"]["charts"] == 1
    assert fidelity["warnings"] == []
    assert "图表 1 个" in fidelity["preserved"]
    workbook = load_workbook(chart_xlsx)
    assert workbook["Sheet1"]["C2"].value == "=B2*2"
    workbook.close()


def test_write_preserves_image_and_chart(rich_xlsx, settings):
    result = _write(settings, rich_xlsx)
    fidelity = result["fidelity"]
    assert fidelity["after"]["images"] == 1
    assert fidelity["after"]["charts"] == 1
    assert any(item.startswith("图片") for item in fidelity["preserved"])
    assert fidelity["warnings"] == []


# ---------------------------------------------------------------- 报警与拦截
def test_lost_objects_are_reported(chart_xlsx, settings, monkeypatch):
    """写后部件减少必须给出警告，而不是静默丢对象。"""
    real = writer.inspect_rich_objects
    calls = {"n": 0}

    def fake(path):
        calls["n"] += 1
        counts = real(path)
        if calls["n"] > 1:  # 第二次是写后校验：模拟 openpyxl 丢弃图表
            counts["charts"] = 0
            counts["drawings"] = 0
        return counts

    monkeypatch.setattr(writer, "inspect_rich_objects", fake)
    result = _write(settings, chart_xlsx)
    assert result["fidelity"]["after"]["charts"] == 0
    assert any("对象丢失" in w for w in result["fidelity"]["warnings"])


def test_images_without_pillow_are_rejected(rich_xlsx, settings, monkeypatch):
    """环境没有 Pillow 且文件含图片时拒绝写入，文件保持原样。"""
    monkeypatch.setattr(writer, "PILImage", None)
    with pytest.raises(ValueError, match="Pillow"):
        _write(settings, rich_xlsx)
    workbook = load_workbook(rich_xlsx)
    assert workbook["Sheet1"]["C2"].value is None
    workbook.close()


def test_wmf_images_get_warning(sample_xlsx, settings):
    """WMF 图片不受支持：写入前给出即可感知的警告。"""
    with zipfile.ZipFile(sample_xlsx) as source:
        entries = {name: source.read(name) for name in source.namelist()}
    entries["xl/media/clip.wmf"] = b"wmf"
    with zipfile.ZipFile(sample_xlsx, "w") as target:
        for name, blob in entries.items():
            target.writestr(name, blob)
    result = _write(settings, sample_xlsx, cell="H1")
    assert any("WMF" in w for w in result["fidelity"]["warnings"])
