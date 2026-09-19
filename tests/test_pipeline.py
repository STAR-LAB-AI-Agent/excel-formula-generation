"""端到端流水线测试：用假客户端替代 DeepSeek，覆盖成功、重试修复、追问、安全边界。"""
from __future__ import annotations

import json
from dataclasses import replace

import pytest
from openpyxl import Workbook, load_workbook

from excel_formula.config import SecurityError
from excel_formula.excel_reader import WorkbookView
from excel_formula.llm_client import ChatResult, DeepSeekClient, LLMError, Usage
from excel_formula.pipeline import FormulaService
from excel_formula.writer import CellWrite, TableFormat, apply_writes


class FakeClient:
    """按顺序返回预置回复，并记录收到的消息，便于断言"重试只回传错误"。"""

    def __init__(self, replies: list[dict | str]):
        self.replies = list(replies)
        self.calls: list[list[dict]] = []

    def chat(self, messages, *, json_mode=True, max_tokens=800):
        self.calls.append(messages)
        reply = self.replies.pop(0)
        content = reply if isinstance(reply, str) else json.dumps(reply, ensure_ascii=False)
        return ChatResult(content=content, usage=Usage(prompt_tokens=120, completion_tokens=40, calls=1))


def _service(settings, replies):
    return FormulaService(settings, logger=None, client=FakeClient(replies))


# ------------------------------------------------------------------ 多目标
def test_keeps_every_target_when_model_returns_several(sample_xlsx, settings):
    """“在 H2 H3 H4 分别求…”这类需求：模型给几个就要全部保留，不能只取第一个。"""
    service = _service(
        settings,
        [{
            "formulas": [
                {"target": "H2", "formula": "=MIN(B2:F2)", "explanation": "Math 最小值"},
                {"target": "H3", "formula": "=MIN(B3:F3)", "explanation": "English 最小值"},
                {"target": "H4", "formula": "=MIN(B4:F4)", "explanation": "Physics 最小值"},
            ],
            "assumptions": [],
            "clarification": None,
        }],
    )
    proposal = service.propose(sample_xlsx, "在H2 H3 H4分别求前三科的最低分")

    assert proposal.ok
    assert [w.cell for w in proposal.writes] == ["H2", "H3", "H4"]
    assert [w.formula for w in proposal.writes] == [
        "=MIN(B2:F2)", "=MIN(B3:F3)", "=MIN(B4:F4)"
    ]
    # 每个单元格都应有自己的本地试算结果
    assert [w.predicted for w in proposal.writes] == ["78", "82", "75"]
    rendered = proposal.render()
    assert "H3 = =MIN(B3:F3)" in rendered
    assert "H4 = =MIN(B4:F4)" in rendered

    applied = service.apply(proposal)
    assert applied["count"] == 3


def test_rejects_batch_when_one_formula_is_invalid(sample_xlsx, settings):
    """一批里有一个不合法就整批不写，并把出错单元格标出来。"""
    bad_batch = {
        "formulas": [
            {"target": "H2", "formula": "=MIN(B2:F2)", "explanation": "正常"},
            {"target": "H3", "formula": '=INDIRECT("B3")', "explanation": "高风险"},
        ],
        "assumptions": [],
        "clarification": None,
    }
    service = _service(settings, [bad_batch, bad_batch, bad_batch])
    proposal = service.propose(sample_xlsx, "在H2 H3分别求最低分")

    assert not proposal.ok
    assert not proposal.writes
    assert "H3" in proposal.error  # 错误带上了出错的单元格
    assert "INDIRECT" in proposal.error


# ------------------------------------------------------------------ 常量值写入
def test_value_writes_alongside_formulas(sample_xlsx, settings):
    """新建汇总表场景：分类名用 value 直接写入，公式引用这些名称。"""
    service = _service(
        settings,
        [{
            "formulas": [
                {"target": "A8", "value": "甲班", "explanation": "分组名称"},
                {"target": "B8", "value": 3.5},
                {"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "总分", "fill_to": "G6"},
            ],
            "assumptions": [],
            "clarification": None,
        }],
    )
    proposal = service.propose(sample_xlsx, "建一张分组表")

    assert proposal.ok
    assert [(w.cell, w.value) for w in proposal.writes if not w.formula] == [
        ("A8", "甲班"), ("B8", 3.5)
    ]
    rendered = proposal.render()
    assert "A8 = 甲班" in rendered
    assert "B8 = 3.5" in rendered
    assert "填充: G3 → G6（共 4 个单元格）" in rendered

    applied = service.apply(proposal)
    assert applied["count"] == 7
    workbook = load_workbook(sample_xlsx)
    assert workbook["Sheet1"]["A8"].value == "甲班"
    assert workbook["Sheet1"]["B8"].value == 3.5
    assert workbook["Sheet1"]["G2"].value == "=SUM(B2:F2)"
    workbook.close()


def test_value_starting_with_equals_is_rejected(sample_xlsx, settings):
    """常量值不能夹带公式：以 = 开头的 value 直接拒绝，不给写入通道。"""
    bad = {
        "formulas": [{"target": "A8", "value": "=SUM(A1:A5)", "explanation": "伪装成值"}],
        "assumptions": [],
        "clarification": None,
    }
    service = _service(settings, [bad, bad, bad])
    proposal = service.propose(sample_xlsx, "在A8写点东西")

    assert not proposal.ok
    assert not proposal.writes
    assert "formula 字段" in proposal.error


# ------------------------------------------------------------------ 正常链路
def test_generate_and_apply(sample_xlsx, settings):
    service = _service(
        settings,
        [{"formulas": [{"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "B2到F2求和",
                        "fill_to": "G6"}], "assumptions": [], "clarification": None}],
    )
    proposal = service.propose(sample_xlsx, "帮我算每个科目的总分")

    assert proposal.ok
    assert proposal.formula == "=SUM(B2:F2)"
    assert [w.cell for w in proposal.writes] == ["G2", "G3", "G4", "G5", "G6"]
    assert proposal.writes[-1].formula == "=SUM(B6:F6)"  # 相对引用已平移
    assert proposal.predicted == "438"  # 写入前的本地预期结果
    assert proposal.usage.total_tokens == 160

    applied = service.apply(proposal)
    assert applied["count"] == 5
    assert applied["backup"] is not None

    workbook = load_workbook(sample_xlsx)
    assert workbook["Sheet1"]["G2"].value == "=SUM(B2:F2)"
    assert workbook["Sheet1"]["G6"].value == "=SUM(B6:F6)"
    workbook.close()


def test_preview_does_not_touch_file(sample_xlsx, settings):
    before = sample_xlsx.read_bytes()
    service = _service(
        settings,
        [{"formulas": [{"target": "G2", "formula": "=AVERAGE(B2:F2)", "explanation": "平均分"}]}],
    )
    proposal = service.propose(sample_xlsx, "算平均分")
    assert proposal.ok
    assert sample_xlsx.read_bytes() == before  # 预览阶段绝不写盘


# ------------------------------------------------------------------ 大作业场景（跨表）
def test_cross_sheet_sumifs_end_to_end(sales_xlsx, settings):
    """复刻大作业跨表统计：查询表按负责人汇总销售额，预览即给出本地独立试算值。"""
    cross_sheet = {
        "formulas": [
            {"target": "E2",
             "formula": "=SUMIFS(原始数据表!$J$2:$J$9, 原始数据表!$G$2:$G$9, D2)",
             "explanation": "按 D2 的负责人汇总销售额"},
        ],
        "assumptions": [],
        "clarification": None,
    }
    service = _service(settings, [cross_sheet])
    proposal = service.propose(sales_xlsx, "在查询表E2按负责人汇总销售额", sheet="查询表")

    assert proposal.ok
    assert proposal.predicted == "51500"  # 张伟 27500 + 24000，一次通过不需重试
    assert len(service.client.calls) == 1

    applied = service.apply(proposal)
    assert applied["count"] == 1
    workbook = load_workbook(sales_xlsx)
    assert workbook["查询表"]["E2"].value.startswith("=SUMIFS(")
    workbook.close()


def test_cross_sheet_vlookup_end_to_end(sales_xlsx, settings):
    """查找区域含日期列时不再拖垮整式：预览给出客户名称而不是"未验证"。"""
    lookup = {
        "formulas": [
            {"target": "B5",
             "formula": '=IFERROR(VLOOKUP($B$3, 原始数据表!$A$2:$L$9, 3, FALSE), "未找到订单")',
             "explanation": "按订单编号查客户名称"},
        ],
        "assumptions": [],
        "clarification": None,
    }
    service = _service(settings, [lookup])
    proposal = service.propose(sales_xlsx, "在查询表B5按订单编号查客户名称", sheet="查询表")

    assert proposal.ok
    assert proposal.predicted == "赵丽"


def test_cross_workbook_sum_end_to_end(tmp_path, settings):
    """跨工作簿引用：预览阶段自动按同目录打开外部工作簿，给出真正的试算值。"""
    orders = tmp_path / "orders.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "明细"
    for offset, value in enumerate([120, 340, 210], start=2):
        sheet.cell(row=offset, column=1, value=value)  # A2:A4
    book.save(orders)

    report = tmp_path / "报表.xlsx"
    workbook = Workbook()
    workbook.active.title = "汇总"
    workbook.save(report)

    service = _service(
        settings,
        [{
            "formulas": [
                {"target": "B2",
                 "formula": "=SUM([orders.xlsx]明细!A2:A4)",
                 "explanation": "跨工作簿汇总订单金额"},
            ],
            "assumptions": [],
            "clarification": None,
        }],
    )
    proposal = service.propose(report, "在汇总表B2跨工作簿汇总订单金额", sheet="汇总")

    assert proposal.ok
    assert proposal.predicted == "670"  # 120 + 340 + 210
    assert any("跨工作簿" in w for w in proposal.validation.warnings)


# ------------------------------------------------------------------ 工作表名对齐
def test_sheet_name_alignment_recovers_mention(class_xlsx, settings):
    """“班级信息表教室列…”被提取成句子片段时，要能按原文找回真实表名，而不是报“工作表不存在”。"""
    service = _service(
        settings,
        [{
            "formulas": [
                {"target": "D2", "formula": "=SUM(成绩单!C2:C5)", "explanation": "全部成绩合计"},
            ],
            "assumptions": [],
            "clarification": None,
        }],
    )
    proposal = service.propose(
        class_xlsx,
        "在班级信息表教室列后面增加一列班级学生平均成绩，成绩来自表一成绩单",
        sheet="教室列后面增加一列班级学",  # 模拟意图提取出的错误表名
        target="D2",
    )

    assert proposal.ok
    assert proposal.sheet == "班级信息"
    assert proposal.predicted == "300"  # 90 + 80 + 70 + 60


def test_sheet_name_alignment_fills_default_from_text(class_xlsx, settings):
    """不传 sheet 时（提取器没给出表名），原文提到的目标表应优于“第一张有数据的表”。"""
    service = _service(
        settings,
        [{
            "formulas": [
                {"target": "D2", "formula": "=SUM(成绩单!C2:C5)", "explanation": "全部成绩合计"},
            ],
            "assumptions": [],
            "clarification": None,
        }],
    )
    proposal = service.propose(
        class_xlsx,
        "在班级信息表教室列后面增加一列班级学生平均成绩，成绩来自表一成绩单",
        target="D2",
    )

    assert proposal.ok
    assert proposal.sheet == "班级信息"


# ------------------------------------------------------------------ 跨表上下文与试算错误
_CLASS_REQUEST = "在班级信息表教室列后面增加一列班级学生平均成绩，成绩来自表一成绩单"
_CLASS_AVERAGE = (
    "=(SUMIF(成绩单!$C$2:$C$13,A2,成绩单!$D$2:$D$13)"
    "+SUMIF(成绩单!$C$2:$C$13,A2,成绩单!$E$2:$E$13)"
    "+SUMIF(成绩单!$C$2:$C$13,A2,成绩单!$F$2:$F$13))"
    "/(COUNTIF(成绩单!$C$2:$C$13,A2)*3)"
)


def _class_reply(formula):
    return {"formulas": [
        {"target": "D1", "value": "班级学生平均成绩"},
        {"target": "D2", "formula": formula, "fill_to": "D4"},
    ]}


def test_cross_sheet_context_and_division_repair(class_scores_xlsx, settings):
    before = class_scores_xlsx.read_bytes()
    wrong = "=AVERAGEIF(成绩单!$A$2:$A$100,A2,成绩单!$B$2:$B$100)"
    service = _service(settings, [_class_reply(wrong), _class_reply(_CLASS_AVERAGE)])
    proposal = service.propose(class_scores_xlsx, _CLASS_REQUEST)
    assert proposal.ok
    assert proposal.sheet == "班级信息"
    assert len(proposal.attempts) == 2
    assert "#DIV/0!" in proposal.attempts[0].errors[0]
    assert any("D4" in error for error in proposal.attempts[0].errors)
    context = service.client.calls[0][1]["content"]
    assert "目标工作表" in context
    assert "来源工作表（仅供引用）: 成绩单" in context
    assert "1\t学号\t姓名\t班级\t语文\t数学\t英语" in context
    assert "统计区" in context
    assert context.count("工作表: 班级信息 |") == 1
    repair = service.client.calls[1][-1]["content"]
    assert "#DIV/0!" in repair and "IFERROR" in repair
    assert "学号\t姓名" not in repair
    assert class_scores_xlsx.read_bytes() == before
    with WorkbookView(class_scores_xlsx) as view:
        results = [service._predict(view, "班级信息", w.formula)[0]
                   for w in proposal.writes if w.formula]
    assert [float(value) for value in results] == pytest.approx(
        [77.833333, 80.416667, 76.416667], abs=0.0001
    )
    assert service.apply(proposal)["count"] == 4
    with WorkbookView(class_scores_xlsx) as view:
        assert view.cell_formula("班级信息", "D2") == _CLASS_AVERAGE
        assert view.value_sheet("成绩单")["C2"].value == "高一1班"
        assert view.value_sheet("成绩单")["H2"].value is None


def test_source_context_included_without_explicit_source_name(class_scores_xlsx, settings):
    service = _service(settings, [_class_reply(_CLASS_AVERAGE)])
    proposal = service.propose(class_scores_xlsx, "从另一张表计算班级三科平均分", sheet="班级信息")
    assert proposal.ok
    assert "1\t学号\t姓名\t班级" in service.client.calls[0][1]["content"]


def test_source_budget_explicitly_marks_omissions(class_scores_xlsx, settings, monkeypatch):
    monkeypatch.setattr("excel_formula.pipeline._MAX_SOURCE_CONTEXT_CHARS", 0)
    service = _service(settings, [{"formulas": [], "clarification": "请提供成绩单数据"}])
    proposal = service.propose(class_scores_xlsx, _CLASS_REQUEST)
    assert proposal.clarification
    context = service.client.calls[0][1]["content"]
    assert "因上下文预算未提供内容的工作表: 成绩单" in context
    assert "不要猜测列号" in context
    assert "1\t学号\t姓名" not in context


def test_source_budget_prioritizes_named_sheet(class_scores_xlsx, settings, monkeypatch):
    with WorkbookView(class_scores_xlsx) as view:
        extra = view.wb_formulas.create_sheet("无关表", 0)
        extra["A1"] = "无关内容"
        view.wb_values.create_sheet("无关表", 0)
        budget = len(view.digest("成绩单").to_prompt())
        monkeypatch.setattr("excel_formula.pipeline._MAX_SOURCE_CONTEXT_CHARS", budget)
        sources, omitted = _service(settings, [])._source_context(view, "班级信息", _CLASS_REQUEST)
    assert list(sources) == ["成绩单"]
    assert omitted == ["无关表"]


@pytest.mark.parametrize("formula,code", [
    ("=1/0", "#DIV/0!"),
    ('=VALUE("abc")', "#VALUE!"),
    ("=SQRT(-1)", "#NUM!"),
    ("=VLOOKUP(B2,IF({1,0},E2:E6,D2:D6),2,0)", "#N/A"),
])
def test_excel_errors_exhaust_retries_and_block_apply(sample_xlsx, settings, formula, code):
    before = sample_xlsx.read_bytes()
    reply = {"formulas": [{"target": "G2", "formula": formula}]}
    service = _service(settings, [reply] * 3)
    proposal = service.propose(sample_xlsx, "计算结果")
    assert not proposal.ok and not proposal.writes
    assert not proposal.validation.ok
    assert code in proposal.error
    assert len(service.client.calls) == 3
    with pytest.raises(ValueError, match="拒绝写入"):
        service.apply(proposal)
    assert sample_xlsx.read_bytes() == before


def test_middle_fill_error_is_not_missed(sample_xlsx, settings):
    reply = {"formulas": [{"target": "G2", "formula": "=1/(B2-88)", "fill_to": "G4"}]}
    service = _service(settings, [reply] * 3)
    proposal = service.propose(sample_xlsx, "计算并填充")
    assert not proposal.ok
    assert "Sheet1!G3" in proposal.error
    assert "#DIV/0!" in proposal.error


@pytest.mark.parametrize("formula,expected", [
    ('="#DIV/0!"', "#DIV/0!"),
    ('=IFERROR(1/0,"无数据")', "无数据"),
])
def test_error_text_or_handled_error_is_not_rejected(sample_xlsx, settings, formula, expected):
    service = _service(settings, [{"formulas": [{"target": "G2", "formula": formula}]}])
    proposal = service.propose(sample_xlsx, "生成提示")
    assert proposal.ok and proposal.predicted == expected
    assert len(service.client.calls) == 1


def test_unsupported_evaluation_is_still_unverified(sample_xlsx, settings):
    service = _service(settings, [{"formulas": [{"target": "G2", "formula": "=SUBTOTAL(9,B2:F2)"}]}])
    proposal = service.propose(sample_xlsx, "计算小计")
    assert proposal.ok and proposal.predicted is None
    assert proposal.predicted_note
    assert len(service.client.calls) == 1


def test_predictions_use_pending_constants_independent_of_order(sample_xlsx, settings):
    before = sample_xlsx.read_bytes()
    service = _service(settings, [{"formulas": [
        {"target": "G2", "formula": "=100/G3"},
        {"target": "G3", "value": 2},
    ]}])
    proposal = service.propose(sample_xlsx, "填写常量与计算结果")
    assert proposal.ok and proposal.predicted == "50"
    assert sample_xlsx.read_bytes() == before


def test_pending_formula_dependency_is_unverified(sample_xlsx, settings):
    service = _service(settings, [{"formulas": [
        {"target": "G2", "formula": "=100/G3"},
        {"target": "G3", "formula": "=2"},
    ]}])
    proposal = service.propose(sample_xlsx, "填写相互依赖的公式")
    assert proposal.ok and proposal.predicted is None
    assert "尚未计算" in proposal.predicted_note


def test_missing_formula_cache_is_not_treated_as_zero(sample_xlsx, settings):
    workbook = load_workbook(sample_xlsx)
    workbook["Sheet1"]["G3"] = "=2"
    workbook.save(sample_xlsx)
    workbook.close()
    service = _service(settings, [{"formulas": [{"target": "G2", "formula": "=100/G3"}]}])
    proposal = service.propose(sample_xlsx, "读取尚未重算的公式")
    assert proposal.ok and proposal.predicted is None
    assert "尚未计算" in proposal.predicted_note


def test_overlapping_fill_is_rejected(sample_xlsx, settings):
    reply = {"formulas": [
        {"target": "G2", "formula": "=SUM(B2:F2)", "fill_to": "G4"},
        {"target": "G3", "value": 100},
    ]}
    service = _service(settings, [reply] * 3)
    proposal = service.propose(sample_xlsx, "填充与单格写入冲突")
    assert not proposal.ok and "重叠" in proposal.error


# ------------------------------------------------------------------ 校验重试
def test_repair_loop_fixes_invalid_formula(sample_xlsx, settings):
    service = _service(
        settings,
        [
            {"formulas": [{"target": "G2", "formula": "=SUM(B2:G2)", "explanation": "错误：含自身"}]},
            {"formulas": [{"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "修正版"}]},
        ],
    )
    proposal = service.propose(sample_xlsx, "算总分写到G2")

    assert proposal.ok
    assert len(proposal.attempts) == 2
    assert not proposal.attempts[0].ok
    assert "循环引用" in proposal.attempts[0].errors[0]
    # 修复请求只带公式与错误，不重复发送表格内容
    repair_message = service.client.calls[1][-1]["content"]
    assert "循环引用" in repair_message
    assert "Student1" not in repair_message


def test_gives_up_after_max_rounds(sample_xlsx, settings):
    bad = {"formulas": [{"target": "G2", "formula": '=INDIRECT("B2")', "explanation": "高风险"}]}
    service = _service(settings, [bad, bad, bad])
    proposal = service.propose(sample_xlsx, "随便算点东西")

    assert not proposal.ok
    assert len(proposal.attempts) == settings.max_repair_rounds + 1
    assert "高风险函数" in proposal.error
    with pytest.raises(ValueError):
        service.apply(proposal)  # 未通过校验的方案禁止写入


def test_reverse_vlookup_with_array_constant_works(sample_xlsx, settings):
    """反向 VLOOKUP（IF({1,0},…)）现在能通过校验，并被本地独立算出预期值。"""
    reverse = {
        "formulas": [
            {"target": "B5",
             "formula": "=VLOOKUP(E2,IF({1,0},E2:E9,D2:D9),2,0)",
             "explanation": "用数组常量翻转列顺序做反向查找"}
        ]
    }
    service = _service(settings, [reverse])
    proposal = service.propose(sample_xlsx, "在B5用VLOOKUP根据B2的姓名查员工号")

    assert proposal.ok
    assert len(service.client.calls) == 1  # 一次通过，不需要重试
    # E2 是 88，翻转查找列后返回 D2 的 92；真正的 #N/A 另测拒绝写入。
    assert proposal.predicted == "92"


def test_truncated_call_usage_is_counted(sample_xlsx, settings):
    """第二轮输出被截断时，那次调用的用量不能从总账中丢失。"""

    class TruncatingClient:
        def __init__(self):
            self.calls = 0

        def chat(self, messages, *, json_mode=True, max_tokens=800):
            self.calls += 1
            if self.calls == 1:
                # 第一轮给个含自身的错误公式，触发一次修复
                return ChatResult(
                    content=json.dumps({"formulas": [
                        {"target": "G2", "formula": "=SUM(B2:G2)", "explanation": "含自身"}]}),
                    usage=Usage(prompt_tokens=700, completion_tokens=1500, calls=1),
                )
            # 第二轮模拟被 max_tokens 截断：报错但携上用量
            raise LLMError(
                "模型输出被 max_tokens 截断",
                usage=Usage(prompt_tokens=1000, completion_tokens=2000, calls=1),
            )

    service = FormulaService(settings, logger=None, client=TruncatingClient())
    proposal = service.propose(sample_xlsx, "算总分写到G2")

    assert not proposal.ok
    assert "截断" in proposal.error
    # 两次调用的用量都要计入：1500 + 2000
    assert proposal.usage.calls == 2
    assert proposal.usage.completion_tokens == 3500
    assert proposal.usage.prompt_tokens == 1700


def test_clarification_is_returned(sample_xlsx, settings):
    service = _service(
        settings,
        [{"formulas": [], "clarification": "你想统计哪一列的分数？", "assumptions": []}],
    )
    proposal = service.propose(sample_xlsx, "帮我统计一下")
    assert not proposal.ok
    assert proposal.clarification == "你想统计哪一列的分数？"


def test_missing_target_triggers_question(sample_xlsx, settings):
    service = _service(
        settings, [{"formulas": [{"target": "", "formula": "=SUM(B2:F2)", "explanation": "总分"}]}]
    )
    proposal = service.propose(sample_xlsx, "算总分")
    assert proposal.clarification is not None


# ------------------------------------------------------------------ 其他意图
def test_validate_reports_predicted_value(sample_xlsx, settings):
    service = _service(settings, [])
    result = service.validate(sample_xlsx, '=COUNTIF(B2:F2,">85")', target="H2")
    assert result["validation"]["ok"]
    assert result["predicted_value"] == "3"


def test_write_formula_rejects_invalid(sample_xlsx, settings):
    service = _service(settings, [])
    with pytest.raises(ValueError, match="未通过校验"):
        service.write_formula(sample_xlsx, "G2", "=SUM(B2:G2")


def test_overwrite_is_reported_in_preview(sample_xlsx, settings):
    service = _service(
        settings, [{"formulas": [{"target": "A2", "formula": "=UPPER(A3)", "explanation": "覆盖测试"}]}]
    )
    proposal = service.propose(sample_xlsx, "把A2改成大写的A3")
    assert proposal.ok
    assert proposal.writes[0].old_value == "Math"
    assert "将覆盖已有内容" in proposal.render()


def test_array_formula_preview_and_overwrite(sample_xlsx, settings):
    """数组语义公式：预览给出本地数组结果；覆盖已写入的数组公式时提醒公式原文。"""
    apply_writes(
        sample_xlsx,
        [CellWrite(sheet="Sheet1", cell="H2", formula="=SUM(B2:B6*C2:C6)")],
        settings,
    )
    service = _service(
        settings,
        [{"formulas": [{"target": "H2", "formula": "=MAX(B2:F2)", "explanation": "覆盖数组公式"}]}],
    )
    proposal = service.propose(sample_xlsx, "把H2改成数学最高分")
    assert proposal.ok
    assert proposal.writes[0].old_formula == "=SUM(B2:B6*C2:C6)"
    assert "将覆盖已有内容" in proposal.render()


def test_array_formula_predicted_value_in_preview(sample_xlsx, settings):
    """含区域运算的公式在预览里给出本地数组语义试算值，而不是“未验证”。"""
    service = _service(
        settings,
        [{"formulas": [{"target": "H2", "formula": "=SUM(B2:B6*C2:C6)",
                        "explanation": "逐元素相乘求和"}]}],
    )
    proposal = service.propose(sample_xlsx, "在H2计算两列成绩乘积之和")
    assert proposal.ok
    assert proposal.writes[0].predicted == "34740"
    assert proposal.writes[0].predicted_note is None


RANK_FORMULA = (
    '=IFERROR(INDEX(成绩单!$B$2:$B$13,MATCH(LARGE(IF(成绩单!$C$2:$C$13="高一1班",'
    '成绩单!$E$2:$E$13),ROW()-1),IF(成绩单!$C$2:$C$13="高一1班",成绩单!$E$2:$E$13),0)),"")'
)


def test_row_based_rank_formula_predicted_in_preview(class_scores_xlsx, settings):
    """含 ROW() 的数组公式在写入预览里给出本地试算值（位置来自目标单元格）。"""
    service = _service(
        settings,
        [{"formulas": [{
            "target": "E2",
            "formula": RANK_FORMULA,
            "explanation": "高一1班数学第一名姓名",
        }]}],
    )
    proposal = service.propose(
        class_scores_xlsx, "在班级信息表E2列出高一1班数学第一名", sheet="班级信息"
    )
    assert proposal.ok
    assert proposal.writes[0].predicted == "学生1"
    assert proposal.writes[0].predicted_note is None


def test_validate_reports_predicted_for_row_formula(class_scores_xlsx, settings):
    """校验意图：目标位置参与 ROW() 求值，预测值不再是“未验证（原因）”。"""
    service = _service(settings, [])
    result = service.validate(class_scores_xlsx, RANK_FORMULA, sheet="班级信息", target="E2")
    assert result["validation"]["ok"]
    assert result["predicted_value"] == "学生1"


# ------------------------------------------------------------------ 安全边界
def test_path_outside_workspace_rejected(settings, tmp_path):
    service = _service(settings, [])
    outside = tmp_path.parent / "outside.xlsx"
    with pytest.raises(SecurityError):
        service.describe(outside)


def test_non_excel_suffix_rejected(settings, tmp_path):
    service = _service(settings, [])
    text_file = tmp_path / "data.csv"
    text_file.write_text("a,b", encoding="utf-8")
    with pytest.raises(SecurityError):
        service.describe(text_file)


def test_fill_range_cell_limit(sample_xlsx, settings):
    service = _service(
        settings,
        [{"formulas": [{"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "过大范围",
                        "fill_to": "G5000"}]}] * 3,
    )
    proposal = service.propose(sample_xlsx, "整列都填上总分")
    assert not proposal.ok
    assert "最多写入" in proposal.error


# ------------------------------------------------------------------ 真实客户端接缝
def test_pipeline_with_real_client_and_fake_socket(sample_xlsx, settings):
    """用真的 DeepSeekClient（仅替换 HTTP 层）跑一遍全链路，验证两者的接口契约。"""

    class FakeResponse:
        status_code = 200
        text = ""

        @staticmethod
        def json():
            return {
                "choices": [{"message": {"content": json.dumps(
                    {"formulas": [{"target": "H2", "formula": "=AVERAGE(B2:F2)",
                                   "explanation": "五人平均分"}]}, ensure_ascii=False)}}],
                "usage": {"prompt_tokens": 380, "completion_tokens": 52},
            }

    class FakeSession:
        def post(self, url, headers=None, data=None, timeout=None, stream=False):
            assert headers["Authorization"].startswith("Bearer ")
            return FakeResponse()

    client = DeepSeekClient(settings)
    client._session = FakeSession()
    service = FormulaService(settings, logger=None, client=client)

    proposal = service.propose(sample_xlsx, "算每科平均分", target="H2")
    assert proposal.ok
    assert proposal.formula == "=AVERAGE(B2:F2)"
    assert proposal.predicted == "87.6"
    assert proposal.usage.prompt_tokens == 380


# ------------------------------------------------------------------ 流式接缝
def test_propose_passes_on_delta_to_client(sample_xlsx, settings):
    """交互入口传了流式回调时，pipeline 必须把它透传给客户端（走 SSE）。"""

    class RecordingClient:
        def __init__(self):
            self.seen_on_delta = None

        def chat(self, messages, *, json_mode=True, max_tokens=800, on_delta=None):
            self.seen_on_delta = on_delta
            if on_delta is not None:
                on_delta("先看看表结构", "reasoning")
                on_delta('{"formulas"', "content")
            return ChatResult(
                content=json.dumps(
                    {"formulas": [
                        {"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "总分"}]},
                    ensure_ascii=False,
                ),
                usage=Usage(prompt_tokens=10, completion_tokens=5, calls=1),
            )

    client = RecordingClient()
    service = FormulaService(settings, logger=None, client=client)
    seen: list[tuple[str, str]] = []
    proposal = service.propose(
        sample_xlsx, "算总分", on_delta=lambda piece, kind: seen.append((kind, piece))
    )

    assert proposal.ok
    assert callable(client.seen_on_delta)
    assert seen == [("reasoning", "先看看表结构"), ("content", '{"formulas"')]


# ------------------------------------------------------------------ auto 思考档
def test_auto_thinking_disables_for_simple_request(sample_xlsx, settings):
    """auto 档：短需求 + 小表直接跳过思考，把等待时间降下来。"""

    class ThinkingRecorder:
        def __init__(self):
            self.thinking_values: list[str | None] = []

        def chat(self, messages, *, json_mode=True, max_tokens=800, on_delta=None, thinking=None):
            self.thinking_values.append(thinking)
            return ChatResult(
                content=json.dumps(
                    {"formulas": [
                        {"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "总分"}]},
                    ensure_ascii=False,
                ),
                usage=Usage(prompt_tokens=10, completion_tokens=5, calls=1),
            )

    client = ThinkingRecorder()
    service = FormulaService(replace(settings, thinking="auto"), logger=None, client=client)
    proposal = service.propose(sample_xlsx, "算总分")

    assert proposal.ok
    assert client.thinking_values == ["disabled"]


def test_auto_thinking_upgrades_on_repair_round(sample_xlsx, settings):
    """auto 档：首轮没过校验后，修复轮次自动升级为思考模式兜底。"""

    class RepairRecorder:
        def __init__(self):
            self.thinking_values: list[str | None] = []

        def chat(self, messages, *, json_mode=True, max_tokens=800, on_delta=None, thinking=None):
            self.thinking_values.append(thinking)
            formula = "=SUM(B2:G2)" if len(self.thinking_values) == 1 else "=SUM(B2:F2)"
            return ChatResult(
                content=json.dumps(
                    {"formulas": [
                        {"target": "G2", "formula": formula, "explanation": "总分"}]},
                    ensure_ascii=False,
                ),
                usage=Usage(prompt_tokens=10, completion_tokens=5, calls=1),
            )

    client = RepairRecorder()
    service = FormulaService(replace(settings, thinking="auto"), logger=None, client=client)
    proposal = service.propose(sample_xlsx, "算总分")

    assert proposal.ok
    assert client.thinking_values == ["disabled", "enabled"]


def test_auto_thinking_heuristics(sample_xlsx, settings):
    """启发式各分支：简单直算关闭；信号词/长需求/大表/修复轮次保留思考。"""
    service = FormulaService(replace(settings, thinking="auto"), logger=None, client=FakeClient([]))
    with WorkbookView(sample_xlsx) as view:
        digest = view.digest("Sheet1")

        assert service._auto_thinking("算总分", digest, 0) == "disabled"
        assert service._auto_thinking("跨表汇总销售额", digest, 0) == "enabled"
        assert service._auto_thinking("很长的需求描述" * 5, digest, 0) == "enabled"
        assert service._auto_thinking("算总分", digest, 1) == "enabled"  # 修复轮次


# ------------------------------------------------------------------ 新建表格格式
_NEW_TABLE_REPLY = {
    "formulas": [
        {"target": "H2", "value": "科目", "explanation": "表头：科目"},
        {"target": "I2", "value": "总分", "explanation": "表头：总分"},
        {"target": "H3", "value": "Math", "explanation": "科目名"},
        {"target": "I3", "formula": "=SUM(B2:F2)", "explanation": "各科总分", "fill_to": "I5"},
    ],
    "assumptions": [],
    "clarification": None,
}


def test_model_table_spec_drives_new_table_format(sample_xlsx, settings):
    """模型给了 table 几何：预览标注表格格式，写入后表头浅蓝底、整表细边框。"""
    reply = dict(_NEW_TABLE_REPLY, table={"header": "H2:I2", "range": "H2:I5"})
    service = _service(settings, [reply])
    proposal = service.propose(sample_xlsx, "新建一个表格统计各科总分", new_table=True)

    assert proposal.ok
    expected = TableFormat(sheet="Sheet1", header_range="H2:I2", table_range="H2:I5")
    assert proposal.table == expected
    assert "表格格式: Sheet1!H2:I5（表头 H2:I2 浅蓝底、范围加细边框）" in proposal.render()
    assert not [w for w in (proposal.validation.warnings or []) if "表格格式" in w]

    applied = service.apply(proposal)
    assert applied["tables"] == [expected.label()]

    workbook = load_workbook(sample_xlsx)
    sheet = workbook["Sheet1"]
    for coord in ("H2", "I2"):
        assert sheet[coord].fill.fgColor.rgb == "FFD9E1F2"
    for coord in ("H2", "I2", "H3", "I3", "I4", "I5"):
        assert sheet[coord].border.left.style == "thin"
        assert sheet[coord].border.bottom.style == "thin"
    assert sheet["H3"].fill.fgColor.rgb != "FFD9E1F2"  # 数据行不铺表头底色
    workbook.close()


def test_new_table_format_derived_from_writes_without_model_geometry(sample_xlsx, settings):
    """模型没给 table 但需求侧识别出新建表格：按写入包围盒兜底（表头=首行）。"""
    service = _service(settings, [_NEW_TABLE_REPLY])
    proposal = service.propose(sample_xlsx, "新建一个表格统计各科总分", new_table=True)

    assert proposal.ok
    assert proposal.table == TableFormat(
        sheet="Sheet1", header_range="H2:I2", table_range="H2:I5"
    )
    # 提示词确实提醒了模型给出 table（6/8=0.75 的写满度满足兜底条件）
    prompt = service.client.calls[-1][-1]["content"]
    assert "这条需求是在新建一张表格" in prompt


def test_new_table_format_skipped_when_writes_scattered(sample_xlsx, settings):
    """写入分散在老表与新表两处（包围盒 6/49 稀疏）：不框大表，给提示而不是乱套格式。"""
    reply = {
        "formulas": [
            {"target": "G2", "formula": "=SUM(B2:F2)", "explanation": "总分", "fill_to": "G6"},
            {"target": "A8", "value": "汇总", "explanation": "标题"},
        ],
        "table": True,  # 只标记新建表格、没给几何
        "assumptions": [],
        "clarification": None,
    }
    service = _service(settings, [reply])
    proposal = service.propose(sample_xlsx, "新建一个表格统计各科总分", new_table=True)

    assert proposal.ok
    assert proposal.table is None
    assert any("未自动套用表格格式" in w for w in proposal.validation.warnings)


def test_new_table_format_rejects_unrelated_model_range(sample_xlsx, settings):
    """模型抄了示例坐标（A24:D28）却没盖住任何写入：拒绝并提示，不误框文件空白区。"""
    reply = dict(_NEW_TABLE_REPLY, table={"header": "A24:D24", "range": "A24:D28"})
    service = _service(settings, [reply])
    proposal = service.propose(sample_xlsx, "在H2:I5统计各科总分")  # 未识别为新建表格

    assert proposal.ok
    assert proposal.table is None
    assert any("已跳过表格格式" in w for w in proposal.validation.warnings)


# ------------------------------------------------------------------ 表格加框（本地意图）
def test_frame_table_builds_border_only_proposal(sample_xlsx, settings):
    """“将 A14 到 B18 框起来”：本地生成仅加框提案（0 写入），确认后给整个范围套细边框。"""
    service = _service(settings, [])  # 不备任何模型回复：误调模型会立刻 IndexError
    proposal = service.frame_table(sample_xlsx, cell_range="A14:B18")

    assert proposal.ok
    assert proposal.writes == []
    assert proposal.table == TableFormat(
        sheet="Sheet1", table_range="A14:B18", force_border=True
    )
    assert "表格格式: Sheet1!A14:B18（范围加细边框）" in proposal.render()

    applied = service.apply(proposal)
    assert applied["count"] == 0
    assert applied["tables"] == ["Sheet1!A14:B18（范围加细边框）"]

    workbook = load_workbook(sample_xlsx)
    sheet = workbook["Sheet1"]
    assert sheet["A14"].border.left.style == "thin"
    assert sheet["B18"].border.bottom.style == "thin"
    assert sheet["A14"].value is None  # 加框不改单元格内容
    workbook.close()


def test_frame_table_without_range_asks_for_it(sample_xlsx, settings):
    """没说范围：给出追问而不是报错，下一条输入作为补充自动续接。"""
    service = _service(settings, [])
    proposal = service.frame_table(sample_xlsx)

    assert not proposal.ok
    assert "范围" in proposal.clarification


def test_frame_table_rejects_oversized_range(sample_xlsx, settings):
    """一次加框超过 MAX_WRITE_CELLS 个单元格时拒绝并给出上限提示。"""
    service = _service(settings, [])
    proposal = service.frame_table(sample_xlsx, cell_range="A1:L100")  # 1200 格

    assert not proposal.ok
    assert "200" in proposal.error
