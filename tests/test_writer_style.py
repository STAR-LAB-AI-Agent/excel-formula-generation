"""写入时的样式沿用：新增单元格继承同行相邻格式，已有样式保持不变。"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from excel_formula.writer import CellWrite, TableFormat, apply_writes


def _book_with_header(tmp_path: Path, name: str = "style.xlsx") -> Path:
    """复刻真实场景：首行表头带填充色与加粗，数据区无样式。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["班级", "班主任", "教室"])
    sheet.append(["一班", "刘老师", "101"])
    sheet.append(["二班", "陈老师", "102"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFD9E1F2")
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return path


def _snapshot(path: Path, sheet_name: str) -> dict[str, tuple]:
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    data = {c.coordinate: (c.value, tuple(c._style) if c._style else None) for c in sheet._cells.values()}
    workbook.close()
    return data


def test_header_cell_inherits_left_neighbor_style(tmp_path, settings):
    """新表头沿用同行最近格式：D1 ← C1（表头颜色与加粗）。"""
    path = _book_with_header(tmp_path)
    result = apply_writes(
        path,
        [CellWrite(sheet="Sheet1", cell="D1", formula="", value="班级学生平均成绩")],
        settings,
    )
    assert result["styled"] == ["D1 ← C1"]
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    assert sheet["D1"].value == "班级学生平均成绩"
    assert sheet["D1"].font.bold is True
    assert sheet["D1"].fill.fgColor.rgb == "FFD9E1F2"
    workbook.close()


def test_existing_custom_style_is_kept(tmp_path, settings):
    """目标格已有自定义格式时不覆盖，也不登记为沿用。"""
    path = _book_with_header(tmp_path)
    workbook = load_workbook(path)
    cell = workbook["Sheet1"]["D1"]
    cell.value = "旧标题"
    cell.fill = PatternFill("solid", fgColor="FFFFF2CC")
    workbook.save(path)
    workbook.close()

    result = apply_writes(
        path, [CellWrite(sheet="Sheet1", cell="D1", formula="", value="新标题")], settings
    )
    assert result["styled"] == []
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    assert sheet["D1"].value == "新标题"
    assert sheet["D1"].fill.fgColor.rgb == "FFFFF2CC"
    assert sheet["D1"].font.bold in (False, None)
    workbook.close()


def test_data_rows_inherit_column_style(tmp_path, settings):
    """数据区逐行沿用：D2 ← C2、D3 ← C3、D4 ← C4（边框等）。"""
    path = _book_with_header(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    border = Border(left=Side(style="thin"), bottom=Side(style="thin"))
    for row in range(2, 5):
        sheet.cell(row=row, column=3).border = border
    workbook.save(path)
    workbook.close()

    writes = [
        CellWrite(sheet="Sheet1", cell=f"D{row}", formula=f"=C{row}*1") for row in range(2, 5)
    ]
    result = apply_writes(path, writes, settings)
    assert result["styled"] == ["D2 ← C2", "D3 ← C3", "D4 ← C4"]
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    for row in range(2, 5):
        assert sheet[f"D{row}"].border.left.style == "thin"
    workbook.close()


def test_no_styled_neighbors_keeps_plain(sample_xlsx, settings):
    """整行没有任何格式时保持默认样式，不凭空造格式。"""
    result = apply_writes(
        sample_xlsx, [CellWrite(sheet="Sheet1", cell="C2", formula="=B2*2")], settings
    )
    assert result["styled"] == []
    workbook = load_workbook(sample_xlsx)
    cell = workbook["Sheet1"]["C2"]
    assert cell.value == "=B2*2"
    assert not cell.has_style
    workbook.close()


def test_right_neighbor_style_used_when_left_empty(tmp_path, settings):
    """左侧无格式、右侧有格式时（往左插新列）沿用右侧格式。"""
    path = _book_with_header(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    sheet["B2"].fill = PatternFill("solid", fgColor="FFC6E0B4")
    workbook.save(path)
    workbook.close()

    result = apply_writes(
        path, [CellWrite(sheet="Sheet1", cell="A2", formula="", value="姓名")], settings
    )
    assert result["styled"] == ["A2 ← B2"]
    workbook = load_workbook(path)
    assert workbook["Sheet1"]["A2"].fill.fgColor.rgb == "FFC6E0B4"
    workbook.close()


def test_mixed_batch_only_plain_cells_get_style(tmp_path, settings):
    """同批“表头常量 + 数据公式”：只有无格式的新格被套样式。"""
    path = _book_with_header(tmp_path)
    writes = [
        CellWrite(sheet="Sheet1", cell="D1", formula="", value="班级学生平均成绩"),
        CellWrite(sheet="Sheet1", cell="D2", formula="=C2"),
    ]
    result = apply_writes(path, writes, settings)
    assert result["styled"] == ["D1 ← C1"]
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    assert sheet["D1"].font.bold is True
    assert sheet["D1"].fill.fgColor.rgb == "FFD9E1F2"
    assert not sheet["D2"].has_style
    workbook.close()


def test_new_cells_chain_style_in_same_batch(tmp_path, settings):
    """同批连续新列：后一格可沿用前一格刚套好的格式。"""
    path = _book_with_header(tmp_path)
    writes = [
        CellWrite(sheet="Sheet1", cell="E1", formula="", value="甲"),
        CellWrite(sheet="Sheet1", cell="F1", formula="", value="乙"),
    ]
    result = apply_writes(path, writes, settings)
    assert result["styled"] == ["E1 ← C1", "F1 ← E1"]
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    assert sheet["E1"].fill.fgColor.rgb == "FFD9E1F2"
    assert sheet["F1"].fill.fgColor.rgb == "FFD9E1F2"
    workbook.close()


def test_existing_styles_untouched(tmp_path, settings):
    """套样式只作用于写入目标，表内其他单元格的格式保持不变。"""
    path = _book_with_header(tmp_path)
    before = _snapshot(path, "Sheet1")
    apply_writes(
        path, [CellWrite(sheet="Sheet1", cell="D1", formula="", value="新增列")], settings
    )
    after = _snapshot(path, "Sheet1")
    changed = [coord for coord, state in before.items() if coord != "D1" and after.get(coord) != state]
    assert changed == []


# ------------------------------------------------------------------ 新建表格格式
def _table_header_and_rows() -> list[CellWrite]:
    return [
        CellWrite(sheet="Sheet1", cell="A8", formula="", value="班级"),
        CellWrite(sheet="Sheet1", cell="B8", formula="", value="平均分"),
        CellWrite(sheet="Sheet1", cell="A9", formula="", value="一班"),
        CellWrite(sheet="Sheet1", cell="B9", formula="=C2"),
        CellWrite(sheet="Sheet1", cell="A10", formula="", value="二班"),
        CellWrite(sheet="Sheet1", cell="B10", formula="=C3"),
    ]


def test_new_table_applies_header_fill_and_border(tmp_path, settings):
    """新建表格：表头铺浅蓝底、整表范围加细边框，数据行不铺底纹。"""
    path = _book_with_header(tmp_path)
    table = TableFormat(sheet="Sheet1", header_range="A8:B8", table_range="A8:B10")
    result = apply_writes(path, _table_header_and_rows(), settings, tables=[table])

    assert result["tables"] == ["Sheet1!A8:B10（表头 A8:B8 浅蓝底、范围加细边框）"]
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    for coord in ("A8", "B8"):
        assert sheet[coord].fill.fgColor.rgb == "FFD9E1F2"
    for coord in ("A8", "B8", "A9", "B9", "A10", "B10"):
        cell = sheet[coord]
        assert cell.border.left.style == "thin"
        assert cell.border.right.style == "thin"
        assert cell.border.top.style == "thin"
        assert cell.border.bottom.style == "thin"
    assert sheet["A9"].fill.fgColor.rgb != "FFD9E1F2"  # 数据行不铺表头底色
    workbook.close()


def test_new_table_borders_unwritten_range_cells(tmp_path, settings):
    """未写入的边框格也会套上格式：范围是矩形，不能中间缺框。"""
    path = _book_with_header(tmp_path)
    table = TableFormat(sheet="Sheet1", header_range="A8:B8", table_range="A8:B10")
    writes = [w for w in _table_header_and_rows() if w.cell not in {"A10", "B10"}]
    apply_writes(path, writes, settings, tables=[table])

    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    assert sheet["A10"].border.left.style == "thin"
    assert sheet["B10"].border.bottom.style == "thin"
    workbook.close()


def test_table_cells_skip_neighbor_inheritance(tmp_path, settings):
    """表内单元格不再沿用相邻格式：以表格格式为准，不登记“D ← C”沿用。"""
    path = _book_with_header(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    sheet["C8"].fill = PatternFill("solid", fgColor="FFFFF2CC")  # 表外的黄格，诱惑沿用
    workbook.save(path)
    workbook.close()

    table = TableFormat(sheet="Sheet1", header_range="A8:B8", table_range="A8:B8")
    writes = [
        CellWrite(sheet="Sheet1", cell="A8", formula="", value="班级"),
        CellWrite(sheet="Sheet1", cell="B8", formula="", value="人数"),
    ]
    result = apply_writes(path, writes, settings, tables=[table])
    assert result["styled"] == []
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    assert sheet["A8"].fill.fgColor.rgb == "FFD9E1F2"
    assert sheet["B8"].fill.fgColor.rgb == "FFD9E1F2"
    workbook.close()


def test_table_format_keeps_user_style_in_range(tmp_path, settings):
    """表格范围内已有的自定义格式保持原样，只补缺的那部分。"""
    path = _book_with_header(tmp_path)
    workbook = load_workbook(path)
    cell = workbook["Sheet1"]["A8"]
    cell.value = "旧表头"
    cell.fill = PatternFill("solid", fgColor="FFFFF2CC")
    workbook.save(path)
    workbook.close()

    table = TableFormat(sheet="Sheet1", header_range="A8:B8", table_range="A8:B8")
    writes = [
        CellWrite(sheet="Sheet1", cell="A8", formula="", value="新表头"),
        CellWrite(sheet="Sheet1", cell="B8", formula="", value="人数"),
    ]
    apply_writes(path, writes, settings, tables=[table])

    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    assert sheet["A8"].fill.fgColor.rgb == "FFFFF2CC"  # 用户格式未被覆盖
    assert sheet["B8"].fill.fgColor.rgb == "FFD9E1F2"
    workbook.close()


# ------------------------------------------------------------------ 范围加框（仅格式）
def test_frame_range_borders_without_writing(tmp_path, settings):
    """“把范围框起来”：零写入只套边框，已有填充保留、范围外不动、内容不改。"""
    path = _book_with_header(tmp_path)
    table = TableFormat(sheet="Sheet1", table_range="A1:B3", force_border=True)
    result = apply_writes(path, [], settings, tables=[table])

    assert result["count"] == 0
    assert result["written"] == []
    assert result["styled"] == []
    assert result["tables"] == ["Sheet1!A1:B3（范围加细边框）"]

    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    for coord in ("A1", "B1", "A2", "B2", "A3", "B3"):
        assert sheet[coord].border.left.style == "thin"
        assert sheet[coord].border.bottom.style == "thin"
    assert sheet["A1"].fill.fgColor.rgb == "FFD9E1F2"  # 已有填充不被清掉
    assert sheet["A1"].value == "班级"  # 单元格内容不变
    assert sheet["C1"].border.left.style is None  # 范围外不动
    workbook.close()


def test_empty_batch_without_tables_is_rejected(tmp_path, settings):
    """既没有公式也没有表格格式的空批次仍然拒绝，避免空写。"""
    path = _book_with_header(tmp_path)
    with pytest.raises(ValueError):
        apply_writes(path, [], settings)
