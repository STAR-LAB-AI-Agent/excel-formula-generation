"""流水线编排：需求 -> 摘要 -> 模型 -> 校验(带重试) -> 预览 -> 写入。"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from openpyxl.utils import get_column_letter

from . import data_validation as dv
from .config import MAX_WRITE_CELLS, Settings
from .evaluator import ExcelError, UnsupportedFormula, evaluate_formula, format_value
from .excel_reader import SheetDigest, WorkbookView, cell_formula_text
from .external import ExternalBookLoader
from .formula_parser import FormulaSyntaxError, parse_range, parse_target
from .intent import align_sheet_name
from .llm_client import (
    DeepSeekClient,
    LLMError,
    Usage,
    build_dropdown_messages,
    build_explain_messages,
    build_generate_messages,
    build_repair_message,
    extract_candidates,
    extract_dropdown_spec,
    extract_table_spec,
    parse_json_payload,
)
from .validator import ValidationResult, validate_formula
from .writer import CellWrite, TableFormat, apply_writes, expand_fill


@dataclass
class Attempt:
    """一次生成尝试的记录，用于日志与"重试了几次"的展示。"""

    round: int
    formula: str
    ok: bool
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {"round": self.round, "formula": self.formula, "ok": self.ok, "errors": self.errors}


@dataclass
class Proposal:
    """待用户确认的公式方案。"""

    ok: bool
    file: Path
    sheet: str
    writes: list[CellWrite] = field(default_factory=list)
    explanation: str = ""
    validation: ValidationResult | None = None
    predicted: str | None = None
    predicted_note: str | None = None
    assumptions: list[str] = field(default_factory=list)
    clarification: str | None = None
    attempts: list[Attempt] = field(default_factory=list)
    usage: Usage = field(default_factory=Usage)
    error: str | None = None
    # 新建表格时自动套用的格式（表头浅蓝底、整表加框），None 表示不涉及
    table: TableFormat | None = None
    dropdown: dv.DropdownSpec | None = None
    file_version: str | None = None
    operation: str = "generate"

    @property
    def formula(self) -> str:
        return self.writes[0].formula if self.writes else ""

    @property
    def target(self) -> str:
        return self.writes[0].cell if self.writes else ""

    def to_dict(self) -> dict:
        return {
            "ok": self.ok,
            "file": str(self.file),
            "sheet": self.sheet,
            "target": self.target,
            "formula": self.formula,
            "explanation": self.explanation,
            "cells": [w.to_dict() for w in self.writes],
            "validation": self.validation.to_dict() if self.validation else None,
            "predicted_value": self.predicted,
            "predicted_note": self.predicted_note,
            "assumptions": self.assumptions,
            "clarification": self.clarification,
            "attempts": [a.to_dict() for a in self.attempts],
            "usage": self.usage.to_dict(),
            "error": self.error,
            "table": self.table.to_dict() if self.table else None,
            "dropdown": self.dropdown.to_dict() if self.dropdown else None,
            "file_version": self.file_version,
            "operation": self.operation,
        }

    def render(self) -> str:
        """给用户看的预览文本。"""
        if self.clarification:
            return f"需要补充信息：{self.clarification}"
        if not self.ok:
            label = "未能设置下拉列表" if self.operation == "dropdown" else "未能生成可用公式"
            return f"{label}：{self.error or '校验未通过'}"
        if self.dropdown:
            spec = self.dropdown
            if self.usage.calls:
                handling = (f"本地规则未命中，DeepSeek 兜底理解 {self.usage.calls} 次调用 / "
                            f"{self.usage.total_tokens} Token")
            else:
                handling = "本地处理 · 0 Token"
            return "\n".join([
                f"文件: {self.file.name}    工作表: {spec.sheet}",
                f"下拉列表范围: {spec.cell_range}",
                f"选项来源: {spec.formula1}",
                "选项: " + "、".join(dv.option_label(v) for v in spec.options[:20])
                + (f" … 共 {len(spec.options)} 项" if len(spec.options) > 20 else ""),
                f"{handling}；仅设置规则，不改动单元格现有内容。",
                *[f"提示: {warning}" for warning in spec.warnings],
            ])

        lines = [f"文件: {self.file.name}    工作表: {self.sheet}"]
        # 值写入恒列出内容；公式只列每个填充块的首格（其余落进下方的“填充”范围行）
        anchors = [
            w for w in self.writes
            if not w.formula or w.explanation or w.predicted or w.predicted_note
        ]
        if len(anchors) > 1:
            for item in anchors:
                if item.formula:
                    line = f"  {item.cell} = {item.formula}"
                    if item.predicted is not None:
                        line += f"    预期结果: {item.predicted}"
                    elif item.predicted_note:
                        line += f"    预期结果: 未验证（{item.predicted_note}）"
                else:
                    line = f"  {item.cell} = {item.value}（值）"
                lines.append(line)
                if item.explanation:
                    lines.append(f"      说明: {item.explanation}")
            lines.append("（预期结果为本地独立计算，未写入文件）")
        elif self.writes:
            first = self.writes[0]
            is_value = not first.formula
            if is_value:
                lines.append(f"目标: {self.target}    值: {first.value}")
            else:
                lines.append(f"目标: {self.target}    公式: {self.formula}")
            if self.explanation:
                lines.append(f"说明: {self.explanation}")
            if not is_value:
                if self.predicted is not None:
                    lines.append(f"预期结果: {self.predicted}（本地独立计算，未写入文件）")
                elif self.predicted_note:
                    lines.append(f"预期结果: 未验证（{self.predicted_note}）")
        # 自动填充的单元格按连续段报告范围（同时填充多列时不会合并成误导的大范围）
        for span in self._fill_spans(anchors):
            lines.append(f"填充: {span[0]} → {span[-1]}（共 {len(span)} 个单元格）")
        if self.table:
            lines.append(f"表格格式: {self.table.label()}")
        overwrite = [w for w in self.writes if w.old_formula or w.old_value is not None]
        if overwrite:
            preview = ", ".join(
                f"{w.cell}={w.old_formula or w.old_value}" for w in overwrite[:3]
            )
            lines.append(f"⚠ 将覆盖已有内容: {preview}{' ...' if len(overwrite) > 3 else ''}")
        if self.validation and self.validation.warnings:
            lines.extend(f"提示: {w}" for w in self.validation.warnings)
        if self.assumptions:
            lines.extend(f"假设: {a}" for a in self.assumptions)
        if len(self.attempts) > 1:
            lines.append(f"（经过 {len(self.attempts)} 次生成，前 {len(self.attempts) - 1} 次校验未通过）")
        return "\n".join(lines)

    def _fill_spans(self, anchors: list[CellWrite]) -> list[list[str]]:
        """自动填充格子按连续段分组；返回每段单元格列表（单格段与空段剔除）。"""
        anchor_ids = {id(w) for w in anchors}
        spans: list[list[str]] = []
        current: list[str] = []
        for item in self.writes:
            if id(item) in anchor_ids or not item.formula:
                # 锚点或值写入会打断填充段
                if current:
                    spans.append(current)
                current = []
            else:
                current.append(item.cell)
        if current:
            spans.append(current)
        return [span for span in spans if len(span) > 1]


# EXCELCR_THINKING=auto 时的"复杂需求"关键词：命中即保留思考模式
_COMPLEX_SIGNALS = (
    "跨表", "另一张", "多条件", "同时满足", "并且", "嵌套", "查找", "匹配",
    "排名", "占比", "百分比", "累计", "环比", "同比", "汇总", "分档", "评级",
)
# 值写入（分类名、标题等常量）的长度上限：防止模型把整段说明塞进单元格
_MAX_VALUE_LENGTH = 200
# 来源表共用的文本预算；超限明确告知模型，不把缺失内容当作空表。
_MAX_SOURCE_CONTEXT_CHARS = 40_000
# 兜底推断表格范围时要求的包围盒写满度：太稀疏说明写入分散在多处，不能框成一张表
_MIN_TABLE_DENSITY = 0.5
# 模型给出的表格范围至少要盖住这么大比例的写入才被采信
_MIN_WRITE_COVERAGE = 0.8


def _range_text(col1: int, row1: int, col2: int, row2: int) -> str:
    return f"{get_column_letter(col1)}{row1}:{get_column_letter(col2)}{row2}"


def _writes_coverage(writes: list[CellWrite], ref) -> float:
    """本次写入落在给定区域内的比例，用于判断模型给出的表格范围是否名副其实。"""
    if not writes:
        return 0.0
    inside = 0
    for item in writes:
        _, col, row = parse_target(item.cell)
        if ref.contains(col, row):
            inside += 1
    return inside / len(writes)


def _notify(callback: Callable[[str], None] | None, text: str) -> None:
    """阶段提示回调：提示展示失败（如 SSE 断连）绝不能影响主流程。"""
    if callback is None:
        return
    try:
        callback(text)
    except Exception:
        pass


def _target_position(target: str | None) -> tuple[int, int] | None:
    """把目标坐标（G2、Sheet1!G2）解析成求值用的 (行, 列)；拿不到返回 None。

    ROW()/COLUMN() 需要公式所在位置；解析失败时交给求值器按"未验证"保守处理。
    """
    if not target:
        return None
    try:
        _, col, row = parse_target(target)
    except FormulaSyntaxError:
        return None
    return (row, col)


class FormulaService:
    """意图的统一入口：生成公式、校验公式、解释公式、表格加框（外加表结构概览）。"""

    def __init__(self, settings: Settings, logger=None, client: object | None = None):
        self.settings = settings
        self.logger = logger
        self.client = client or DeepSeekClient(settings, logger)

    # ============================================================ 意图 0：概览
    def describe(self, file: str | Path, sheet: str | None = None) -> dict:
        path = self.settings.resolve_path(file)
        with WorkbookView(path) as view:
            digest = view.digest(sheet)
            result = {
                "file": str(path),
                "sheets": view.overview(),
                "digest": digest.to_dict(),
                "digest_text": digest.to_prompt(),
            }
        if self.logger:
            self.logger.info("读取结构 file=%s sheet=%s", path.name, result["digest"]["sheet"])
        return result

    # ============================================================ 意图 1：生成
    def propose(
        self,
        file: str | Path,
        request: str,
        *,
        sheet: str | None = None,
        target: str | None = None,
        new_table: bool = False,
        on_delta: Callable[[str, str], None] | None = None,
    ) -> Proposal:
        path = self.settings.resolve_path(file)
        if not request or not request.strip():
            raise ValueError("请描述你想要的计算，例如“在G列算每科总分”")

        with WorkbookView(path) as view:
            # 提取器可能把“班级信息表教室列…”整段抓成表名，也可能什么都没抓到；回到原文找回真实表名
            aligned = align_sheet_name(request, sheet, view.sheet_names)
            if aligned != sheet and self.logger:
                self.logger.info("工作表名对齐 %r -> %r", sheet, aligned)
            sheet = aligned
            sheet_name = view.resolve_sheet(sheet)
            digest = view.digest(sheet_name)
            if self.logger:
                self.logger.info(
                    "生成公式 file=%s sheet=%s target=%s 需求=%s",
                    path.name, sheet_name, target or "自动", request.strip()[:80],
                )
            proposal = Proposal(ok=False, file=path, sheet=sheet_name)
            sources, omitted = self._source_context(view, sheet_name, request)
            messages = build_generate_messages(
                digest.to_prompt(), request, target=target, sheet_names=view.sheet_names,
                source_digests=sources, omitted_sheets=omitted, new_table=new_table,
            )

            for round_index in range(self.settings.max_repair_rounds + 1):
                # 传了回调才走流式；不传时保持一次性返回，兼容程序化调用与测试替身
                chat_kwargs: dict = {"on_delta": on_delta} if on_delta is not None else {}
                if self.settings.thinking == "auto":
                    # auto 档：简单需求跳过思考提速；首轮失败后自动升级为思考再修
                    chat_kwargs["thinking"] = self._auto_thinking(request, digest, round_index)
                try:
                    result = self.client.chat(messages, **chat_kwargs)
                except LLMError as exc:
                    if exc.usage:  # 输出被截断等失败仍占用了 Token，不能漏算
                        proposal.usage.add(exc.usage)
                    proposal.error = str(exc)
                    if self.logger:
                        self.logger.error("模型调用失败：%s", exc)
                    return proposal
                proposal.usage.add(result.usage)

                try:
                    payload = parse_json_payload(result.content)
                except LLMError as exc:
                    proposal.error = str(exc)
                    return proposal

                candidates, assumptions, clarification = extract_candidates(payload)
                proposal.assumptions = assumptions
                if clarification and not candidates:
                    proposal.clarification = clarification
                    if self.logger:
                        self.logger.info("模型请求补充信息：%s", clarification)
                    return proposal
                if not candidates:
                    proposal.error = "模型没有给出公式"
                    if self.logger:
                        self.logger.warning(
                            "模型没有给出公式：返回里没有可写入内容 需求=%s",
                            request.strip()[:60],
                        )
                    return proposal

                # 模型可能一次给出多个目标单元格（例如“E2 F2 G2 分别求最小值”），
                # 逐个校验，全部通过才算成功；任何一个失败就把所有错误一起回传
                accepted: list[CellWrite] = []
                errors: list[str] = []
                validation: ValidationResult | None = None
                seen: set[str] = set()
                labelled = len(candidates) > 1
                anchors: set[str] = set()

                for index, candidate in enumerate(candidates):
                    formula = candidate["formula"]
                    cell = candidate["target"] or (target if index == 0 else None)
                    if not cell:
                        proposal.clarification = "请告诉我把内容写到哪个单元格，例如 G2。"
                        return proposal
                    cell = cell.upper()
                    if cell in seen:
                        errors.append(f"{cell}: 同一单元格被指定了多个写入")
                        continue
                    seen.add(cell)

                    if not formula:
                        # 值写入：分类名、标题等已知常量，无需公式校验
                        try:
                            accepted.append(
                                self._build_value_write(
                                    view, sheet_name, cell,
                                    candidate.get("value"), candidate["explanation"],
                                )
                            )
                        except ValueError as exc:
                            prefix = f"{cell}: " if labelled else ""
                            errors.append(prefix + str(exc))
                        continue

                    one, cell_errors = self._check_candidate(
                        view, sheet_name, digest, cell, formula, candidate.get("fill_to")
                    )
                    validation = _merge_validation(validation, one)
                    if cell_errors:
                        prefix = f"{cell}: " if labelled else ""
                        errors.extend(prefix + message for message in cell_errors)
                        continue

                    writes = self._build_writes(
                        view, sheet_name, cell, formula, candidate.get("fill_to")
                    )
                    writes[0].explanation = candidate["explanation"]
                    anchors.add(writes[0].cell)
                    accepted.extend(writes)

                if not errors:
                    prediction = self._check_predictions(view, accepted, anchors)
                    validation = _merge_validation(validation, prediction)
                    errors.extend(prediction.errors)

                formulas = [c["formula"] for c in candidates]
                proposal.attempts.append(
                    Attempt(
                        round=round_index + 1,
                        formula="; ".join(formulas),
                        ok=not errors,
                        errors=list(errors),
                    )
                )
                if errors:
                    if self.logger:
                        self.logger.warning(
                            "第 %d 次校验失败 formulas=%s errors=%s",
                            round_index + 1, formulas, errors,
                        )
                    # 命中“重试也修不好”的错误类别：直接收手，不再白烧一轮 Token
                    unrepairable = validation is not None and validation.unrepairable
                    if unrepairable or round_index >= self.settings.max_repair_rounds:
                        proposal.validation = validation
                        proposal.error = "；".join(errors)
                        if unrepairable:
                            proposal.error += "。该写法本地不支持，重试无法解决，已停止"
                        return proposal
                    messages = messages + [
                        {"role": "assistant", "content": result.content},
                        build_repair_message(formulas, errors),
                    ]
                    continue

                proposal.ok = True
                proposal.writes = accepted
                proposal.explanation = accepted[0].explanation
                proposal.validation = validation
                proposal.predicted = accepted[0].predicted
                proposal.predicted_note = accepted[0].predicted_note
                proposal.table, table_note = self._resolve_table(
                    payload, accepted, new_table=new_table, sheet=sheet_name
                )
                if table_note:
                    if proposal.validation is not None:
                        proposal.validation.warnings.append(table_note)
                    if self.logger:
                        self.logger.warning("表格格式未套用：%s", table_note)
                if self.logger:
                    self.logger.info(
                        "公式通过校验 共 %d 个单元格 targets=%s 预期值=%s",
                        len(accepted),
                        ", ".join(w.cell for w in accepted[:5]),
                        accepted[0].predicted or accepted[0].predicted_note,
                    )
                return proposal
        return proposal

    # ============================================================ 意图 2：校验
    def validate(
        self,
        file: str | Path,
        formula: str,
        *,
        sheet: str | None = None,
        target: str | None = None,
    ) -> dict:
        path = self.settings.resolve_path(file)
        with WorkbookView(path) as view:
            sheet_name = view.resolve_sheet(sheet)
            digest = view.digest(sheet_name)
            validation, _ = validate_formula(
                formula,
                target=target,
                sheet=sheet_name,
                digest=digest,
                sheet_names=view.sheet_names,
            )
            predicted, note = (
                self._predict(
                    view, sheet_name, formula, current_cell=_target_position(target)
                )
                if validation.ok
                else (None, "公式未通过校验，已跳过本地计算")
            )
        if self.logger:
            self.logger.info(
                "校验公式 formula=%s ok=%s errors=%s", formula, validation.ok, validation.errors
            )
        return {
            "file": str(path),
            "sheet": sheet_name,
            "formula": formula,
            "target": target,
            "validation": validation.to_dict(),
            "predicted_value": predicted,
            "predicted_note": note,
        }

    # ============================================================ 意图 3：解释
    def explain(
        self,
        formula: str | None = None,
        *,
        file: str | Path | None = None,
        sheet: str | None = None,
        cell: str | None = None,
        on_delta: Callable[[str, str], None] | None = None,
    ) -> dict:
        digest_text = None
        resolved_formula = (formula or "").strip()
        path = None

        if file:
            path = self.settings.resolve_path(file)
            with WorkbookView(path) as view:
                sheet_name = view.resolve_sheet(sheet)
                digest_text = view.digest(sheet_name).to_prompt()
                if cell and not resolved_formula:
                    found = view.cell_formula(sheet_name, cell.upper())
                    if not found:
                        raise ValueError(f"{sheet_name}!{cell.upper()} 中没有公式")
                    resolved_formula = found
        if not resolved_formula:
            raise ValueError("请提供要解释的公式，或指定包含公式的单元格")

        # 先本地校验一次，静态问题不必花 Token 就能告知用户
        validation, _ = validate_formula(resolved_formula)
        chat_kwargs = {"on_delta": on_delta} if on_delta is not None else {}
        if self.settings.thinking == "auto":
            # 短公式直译即可；嵌套较深的公式才值得开思考
            chat_kwargs["thinking"] = "enabled" if len(resolved_formula) >= 60 else "disabled"
        result = self.client.chat(
            build_explain_messages(resolved_formula, digest_text),
            json_mode=False,
            max_tokens=1200,
            **chat_kwargs,
        )
        if self.logger:
            self.logger.info("解释公式 formula=%s", resolved_formula)
        return {
            "file": str(path) if path else None,
            "formula": resolved_formula,
            "explanation": result.content.strip(),
            "validation": validation.to_dict(),
            "usage": result.usage.to_dict(),
        }

    # ============================================================ 意图 4：表格加框
    def frame_table(
        self,
        file: str | Path,
        *,
        sheet: str | None = None,
        cell_range: str | None = None,
    ) -> Proposal:
        """给指定范围套细边框（本地处理，0 Token）：不改任何单元格内容。

        承接“将 A14 到 B18 框起来”这类纯格式需求——模型只会生成公式、给不了
        格式，这类请求在本地完成，并复用同一套预览/确认/写入通道。
        """
        path = self.settings.resolve_path(file)
        with WorkbookView(path) as view:
            sheet_name = view.resolve_sheet(sheet)
            proposal = Proposal(ok=False, file=path, sheet=sheet_name)
            if not cell_range:
                proposal.clarification = "请告诉我给哪个范围加边框，例如“将 A14 到 B18 框起来”。"
                return proposal
            try:
                ref = parse_range(cell_range)
            except FormulaSyntaxError as exc:
                proposal.error = f"范围写法无法解析：{exc}"
                return proposal
            if (ref.col2 - ref.col1 + 1) * (ref.row2 - ref.row1 + 1) > MAX_WRITE_CELLS:
                proposal.error = f"一次最多给 {MAX_WRITE_CELLS} 个单元格加边框，请缩小范围"
                return proposal
        proposal.ok = True
        proposal.table = TableFormat(
            sheet=sheet_name,
            table_range=_range_text(ref.col1, ref.row1, ref.col2, ref.row2),
            force_border=True,
        )
        if self.logger:
            self.logger.info(
                "表格加框 file=%s sheet=%s range=%s",
                path.name, sheet_name, proposal.table.table_range,
            )
        return proposal

    # ============================================================ 列表数据验证
    def propose_dropdown(
        self,
        file: str | Path,
        request: str,
        *,
        sheet: str | None = None,
        on_stage: Callable[[str], None] | None = None,
    ) -> Proposal:
        """本地生成规则方案，保留文件版本，确认后才执行保存。

        本地规则未命中时（来源说不清、句式没见过）自动把原句交给模型理解一次，
        模型输出的表/列/区域仍经 dv.make_spec 严格校验，幻觉来源会被原样拒绝。
        """
        path = self.settings.resolve_path(file)
        proposal = Proposal(ok=False, file=path, sheet=sheet or "", operation="dropdown")
        if self.logger:
            self.logger.info("下拉列表需求 file=%s sheet=%s 需求=%s",
                             path.name, sheet or "自动", request[:120])
        local_hit = False
        try:
            with dv.WRITE_LOCK:
                version = dv.file_version(path)
                warning = dv.extension_warning(path)
                if warning:
                    raise ValueError(warning)
                with WorkbookView(path) as view:
                    default = view.resolve_sheet(sheet)
                    target_sheet, target, options, source = dv.parse_request(request, view.sheet_names, default)
                    proposal.sheet = target_sheet
                    proposal.dropdown = dv.make_spec(view, target_sheet, target, options, source)
                dv.require_version(path, version)
                proposal.file_version = version
                proposal.ok = True
                local_hit = True
        except dv.DropdownClarification as exc:
            # “没听懂/缺来源”交给模型理解一次；确定性拒绝（删除、多操作等）不走这里
            proposal.clarification = str(exc)
            self._dropdown_model_fallback(path, request, sheet, proposal, on_stage)
        except ValueError as exc:
            proposal.error = str(exc)
            if str(exc).startswith("无法识别下拉来源"):
                self._dropdown_model_fallback(path, request, sheet, proposal, on_stage)
        if local_hit:
            _notify(on_stage, "本地规则命中：不调用模型（0 Token）")
        return proposal

    def _dropdown_model_fallback(
        self,
        path: Path,
        request: str,
        sheet: str | None,
        proposal: Proposal,
        on_stage: Callable[[str], None] | None = None,
    ) -> None:
        """本地失败后用一次模型调用理解原句；任何失败都静默保留原追问。

        兜底成功的前提是模型给出的目标与来源能通过 dv.make_spec 的全套校验，
        因此幻觉表名/列标题只会退回本地提示，不会落到文件上。
        目标表沿用“未指定就用当前表”规则：模型点名的表必须与原句绑定，否则拉回当前表。
        """
        _notify(on_stage, "本地规则未命中，改由 DeepSeek 理解原句（1 次调用）…")
        try:
            version = dv.file_version(path)
            with WorkbookView(path) as view:
                default = view.resolve_sheet(sheet)
                sources, omitted = self._source_context(view, default, request)
                parts = [f"目标工作表「{default}」\n{view.digest(default).to_prompt()}"]
                parts.extend(f"来源工作表「{name}」\n{text}" for name, text in sources.items())
                if omitted:
                    parts.append("因上下文预算未提供内容的工作表: " + "、".join(omitted))
                messages = build_dropdown_messages(
                    "\n\n".join(parts), request, default=default, sheet_names=view.sheet_names,
                )
                try:
                    result = self.client.chat(messages, thinking="disabled", max_tokens=1500)
                except LLMError as exc:
                    if exc.usage:  # 失败也可能已计费（如输出截断），如实记账
                        proposal.usage.add(exc.usage)
                    raise
                proposal.usage.add(result.usage)
                payload = parse_json_payload(result.content)
                spec = extract_dropdown_spec(payload)
                if spec is None:
                    return
                # 未指定目标表时优先当前表：句里把表名连到目标格的才作数；
                # 模型若只是把来源表名塞进 sheet（无绑定），同样拉回当前表。
                hint = dv.target_sheet_hint(request, view.sheet_names, spec["target"])
                if hint:
                    chosen = hint
                elif (spec["sheet"] and spec["sheet"] != spec.get("source_sheet")
                      and spec["sheet"] in request):
                    chosen = spec["sheet"]
                else:
                    chosen = default
                target_sheet = view.resolve_sheet(chosen)
                proposal.sheet = target_sheet
                proposal.dropdown = dv.make_spec(
                    view, target_sheet, spec["target"], spec["options"], spec["source"]
                )
            dv.require_version(path, version)
            proposal.file_version = version
            proposal.clarification = None
            proposal.error = None
            proposal.ok = True
            if self.logger:
                self.logger.info("下拉兜底成功 file=%s sheet=%s", path.name, proposal.sheet)
        except Exception as exc:
            # 模型没帮上忙：保留本地提示；已发生的调用已记在 usage 里
            if self.logger:
                self.logger.info("下拉兜底未成功（%s），保留本地提示", exc)

    def select_dropdown(
        self, file: str | Path, *, sheet: str, cell: str, rule_id: str,
        option_index: int, file_version: str,
    ) -> dict:
        """选择已存在规则中的一项；不接受浏览器任意指定写入值。"""
        path = self.settings.resolve_path(file)
        if not sheet or not cell:
            raise ValueError("请选择工作表和单元格")
        with dv.WRITE_LOCK:
            dv.require_version(path, file_version)
            warning = dv.extension_warning(path)
            if warning:
                raise ValueError(warning)
            with WorkbookView(path) as view:
                sheet = view.resolve_sheet(sheet)
                coord, value = dv.selected_value(view, sheet, cell, rule_id, option_index)
            return apply_writes(
                path, [CellWrite(sheet=sheet, cell=coord, formula="", value=value)],
                self.settings, logger=self.logger, expected_version=file_version,
            )

    # ============================================================ 写入
    def apply(self, proposal: Proposal, *, output: str | Path | None = None) -> dict:
        if not proposal.ok or (not proposal.writes and proposal.table is None and proposal.dropdown is None):
            raise ValueError("方案未通过校验，拒绝写入")
        out_path = self.settings.resolve_path(output, must_exist=False) if output else None
        return apply_writes(
            proposal.file, proposal.writes, self.settings, logger=self.logger, output=out_path,
            tables=[proposal.table] if proposal.table else None,
            dropdown=proposal.dropdown, expected_version=proposal.file_version,
        )

    def write_formula(
        self,
        file: str | Path,
        target: str,
        formula: str,
        *,
        sheet: str | None = None,
        fill_to: str | None = None,
        output: str | Path | None = None,
    ) -> dict:
        """直接写入指定公式（仍然先做本地校验），供"我自己写好了公式"的场景使用。"""
        path = self.settings.resolve_path(file)
        with WorkbookView(path) as view:
            sheet_name = view.resolve_sheet(sheet)
            digest = view.digest(sheet_name)
            validation, errors = self._check_candidate(
                view, sheet_name, digest, target, formula, fill_to
            )
            if errors:
                raise ValueError("公式未通过校验：" + "；".join(errors))
            writes = self._build_writes(view, sheet_name, target, formula, fill_to)
        out_path = self.settings.resolve_path(output, must_exist=False) if output else None
        return apply_writes(path, writes, self.settings, logger=self.logger, output=out_path)

    # ============================================================ 内部方法
    def _source_context(
        self, view: WorkbookView, target_sheet: str, request: str
    ) -> tuple[dict[str, str], list[str]]:
        """目标表之外的同簿工作表也提供 TSV，优先保留需求明确提到的表。"""
        names = [name for name in view.sheet_names if name != target_sheet]
        names.sort(key=lambda name: name.casefold() not in request.casefold())
        sources: dict[str, str] = {}
        omitted: list[str] = []
        remaining = _MAX_SOURCE_CONTEXT_CHARS
        for name in names:
            text = view.digest(name).to_prompt()
            if len(text) > remaining:
                omitted.append(name)
                continue
            sources[name] = text
            remaining -= len(text)
        return sources, omitted

    # ------------------------------------------------ 新建表格格式
    def _resolve_table(
        self,
        payload: dict,
        writes: list[CellWrite],
        *,
        new_table: bool,
        sheet: str,
    ) -> tuple[TableFormat | None, str | None]:
        """确定新建表格的几何：优先采信模型给出的 table，其次按写入范围兜底。

        返回 (格式, 提示)。兜底只在需求侧识别出“新建表格”时启用，且写入范围必须规整；
        宁可不套格式，也不把零散写入框成一张误导的大表。
        """
        spec = extract_table_spec(payload)
        has_geometry = bool(spec and (spec.get("range") or spec.get("header")))
        if spec is not None:
            table = self._table_from_spec(spec, sheet, writes)
            if table is not None:
                return table, None
        if new_table:
            table = self._derive_table(sheet, writes)
            if table is not None:
                return table, None
            # 模型只标了“新建表格”没给几何、或几何不可用时，问题都回到写入本身是否规整
            return None, "写入范围不构成规整表格（存在跳格或过大），未自动套用表格格式"
        if has_geometry:
            return None, "模型给出的表格范围无效，已跳过表格格式"
        return None, None

    @staticmethod
    def _table_from_spec(spec: dict, sheet: str, writes: list[CellWrite]) -> TableFormat | None:
        """校验模型给出的表格范围：可解析、不超上限、盖住绝大多数写入。

        表头缺失或越界时退化为范围首行；范围不成立返回 None，由调用方决定是否兜底。
        """
        try:
            ref = parse_range(spec.get("range") or "")
        except FormulaSyntaxError:
            return None
        if ref.cell_count < 2 or ref.cell_count > MAX_WRITE_CELLS:
            return None
        if _writes_coverage(writes, ref) < _MIN_WRITE_COVERAGE:
            return None
        header_text = _range_text(ref.col1, ref.row1, ref.col2, ref.row1)  # 默认表头为范围首行
        if spec.get("header"):
            try:
                header = parse_range(spec["header"])
            except FormulaSyntaxError:
                header = None
            if header is not None:
                col1, col2 = max(header.col1, ref.col1), min(header.col2, ref.col2)
                row1, row2 = max(header.row1, ref.row1), min(header.row2, ref.row2)
                if col1 <= col2 and row1 <= row2:
                    header_text = _range_text(col1, row1, col2, row2)
        return TableFormat(
            sheet=sheet,
            header_range=header_text,
            table_range=_range_text(ref.col1, ref.row1, ref.col2, ref.row2),
        )

    @staticmethod
    def _derive_table(sheet: str, writes: list[CellWrite]) -> TableFormat | None:
        """按写入范围推断新建表格几何：表头=包围盒首行，范围=整个包围盒。

        只有包围盒写满度达 _MIN_TABLE_DENSITY 才认账：写入若分散在老表与新表两处，
        大包围盒会把老数据一起框住，此时宁可不套格式。
        """
        coords = {
            (col, row)
            for item in writes
            if item.sheet == sheet
            for _, col, row in [parse_target(item.cell)]
        }
        if not coords:
            return None
        cols = [col for col, _ in coords]
        rows = [row for _, row in coords]
        col1, col2, row1, row2 = min(cols), max(cols), min(rows), max(rows)
        total = (col2 - col1 + 1) * (row2 - row1 + 1)
        if total < 2 or total > MAX_WRITE_CELLS:
            return None
        if len(coords) / total < _MIN_TABLE_DENSITY:
            return None
        return TableFormat(
            sheet=sheet,
            header_range=_range_text(col1, row1, col2, row1),
            table_range=_range_text(col1, row1, col2, row2),
        )

    def _check_predictions(
        self, view: WorkbookView, writes: list[CellWrite], anchors: set[str]
    ) -> ValidationResult:
        """逐格试算整批（含填充），确定的 Excel 错误进入修复；能力不足仍标未验证。

        仅在内存值视图中暂存本批内容，避免用旧的分类名/分母误判。
        本批公式依赖没有缓存时保守标未验证，不假装已算出新结果。
        """
        result = ValidationResult(ok=True)
        saved = []
        loader = ExternalBookLoader()
        try:
            occupied: set[tuple[str, str]] = set()
            for item in writes:
                key = (item.sheet, item.cell)
                if key in occupied:
                    result.errors.append(f"{item.cell}: 填充区域与其他写入重叠")
                    continue
                occupied.add(key)
                cell = view.wb_values[item.sheet][item.cell]
                saved.append((cell, cell.value, cell.data_type))
                cell.value = item.content
            if result.errors:
                result.ok = False
                return result
            for item in writes:
                if not item.formula:
                    continue
                value, note = self._evaluate(
                    view, item.sheet, item.formula,
                    loader=loader, current_cell=_target_position(item.cell),
                )
                if item.cell in anchors:
                    item.predicted = format_value(value) if note is None else None
                    item.predicted_note = note
                elif note:
                    result.warnings.append(f"{item.cell}: 本地试算未验证（{note}）")
                if isinstance(value, ExcelError):
                    result.errors.append(
                        f"{item.sheet}!{item.cell}: 本地试算得到 {value.code}；"
                        "请核对来源表的真实列、匹配条件和数值数据区，"
                        "不要仅用 IFERROR 返回 0 掩盖引用错误"
                    )
        finally:
            for cell, value, data_type in reversed(saved):
                cell.value = value
                cell.data_type = data_type
            self._close_loader(loader)
        result.ok = not result.errors
        return result

    def _auto_thinking(self, request: str, digest: SheetDigest, round_index: int) -> str:
        """EXCELCR_THINKING=auto 的复杂度启发式：简单直算跳过思考省时间，复杂需求保留思考。

        硬信号优先：修复轮次（上一轮没过本地校验，说明问题不简单）、表格规模
        （大表定位数据区更容易出错）；需求侧看长度与复杂关键词。
        """
        if round_index > 0:
            return "enabled"
        if len(request) >= 24:
            return "enabled"
        if digest.max_row > 40 or digest.max_column > 12:
            return "enabled"
        if any(word in request for word in _COMPLEX_SIGNALS):
            return "enabled"
        return "disabled"

    def _check_candidate(
        self,
        view: WorkbookView,
        sheet_name: str,
        digest: SheetDigest,
        target: str,
        formula: str,
        fill_to: str | None,
    ) -> tuple[ValidationResult, list[str]]:
        """校验首个单元格；若需填充，额外校验最后一个单元格平移后的公式。"""
        validation, _ = validate_formula(
            formula,
            target=target,
            sheet=sheet_name,
            digest=digest,
            sheet_names=view.sheet_names,
        )
        errors = list(validation.errors)
        if not errors and fill_to:
            try:
                writes = expand_fill(sheet_name, target, formula, fill_to)
            except (ValueError, FormulaSyntaxError) as exc:
                errors.append(str(exc))
                return validation, errors
            tail = writes[-1]
            tail_validation, _ = validate_formula(
                tail.formula,
                target=tail.cell,
                sheet=sheet_name,
                digest=digest,
                sheet_names=view.sheet_names,
            )
            for message in tail_validation.errors:
                errors.append(f"填充到 {tail.cell} 时: {message}")
            validation.warnings.extend(
                f"填充到 {tail.cell} 时: {w}" for w in tail_validation.warnings
            )
        return validation, errors

    def _mark_old_content(self, view: WorkbookView, sheet_name: str, item: CellWrite) -> None:
        """记录目标单元格当前内容，供预览里的覆盖提醒使用。"""
        _, col, row = parse_target(item.cell)
        raw = cell_formula_text(view.formula_sheet(sheet_name).cell(row=row, column=col))
        if raw is not None and raw.startswith("="):
            item.old_formula = raw
        else:
            item.old_value = view.value_sheet(sheet_name).cell(row=row, column=col).value

    def _build_writes(
        self,
        view: WorkbookView,
        sheet_name: str,
        target: str,
        formula: str,
        fill_to: str | None,
    ) -> list[CellWrite]:
        writes = expand_fill(sheet_name, target, formula, fill_to)
        for item in writes:
            self._mark_old_content(view, sheet_name, item)
        return writes

    def _build_value_write(
        self,
        view: WorkbookView,
        sheet_name: str,
        cell: str,
        value: object,
        explanation: str,
    ) -> CellWrite:
        """值写入：把已知常量（区域名、标题等）直接写入单元格，不经过公式校验。"""
        if isinstance(value, bool) or not isinstance(value, (str, int, float)):
            raise ValueError("value 只能是文本或数字")
        if isinstance(value, str):
            if not value.strip():
                raise ValueError("常量值不能为空")
            if value.startswith("="):
                raise ValueError("常量值不能以 = 开头，公式请使用 formula 字段")
            if len(value) > _MAX_VALUE_LENGTH:
                raise ValueError(f"常量值超过 {_MAX_VALUE_LENGTH} 字，疑似把说明文字写进了单元格")
        write = CellWrite(sheet=sheet_name, cell=cell, formula="", value=value)
        write.explanation = explanation
        write.predicted = str(value)
        self._mark_old_content(view, sheet_name, write)
        return write

    def _predict(
        self, view: WorkbookView, sheet_name: str, formula: str,
        current_cell: tuple[int, int] | None = None,
    ) -> tuple[str | None, str | None]:
        """可选功能：在 Python 中独立算一遍公式，作为写入前的预期结果。

        跨工作簿引用会通过 ExternalBookLoader 按需打开外部工作簿试算；
        找不到或打不开时按"未验证（原因）"返回，不影响写入。
        current_cell 是公式将写入的位置（行、列），供 ROW()/COLUMN() 求值。
        """
        value, note = self._evaluate(view, sheet_name, formula, current_cell=current_cell)
        return (format_value(value), None) if note is None else (None, note)

    def _evaluate(
        self, view: WorkbookView, sheet_name: str, formula: str,
        *, loader: ExternalBookLoader | None = None,
        current_cell: tuple[int, int] | None = None,
    ) -> tuple[object, str | None]:
        """保留 ExcelError 类型，避免将同样内容的普通文本误判为计算错误。"""
        _, ast = validate_formula(formula, sheet=sheet_name)
        if ast is None:
            return None, "公式无法解析"
        sheets = {name: view.wb_values[name] for name in view.sheet_names}
        formulas = {name: view.wb_formulas[name] for name in view.sheet_names}
        owns_loader = loader is None
        if loader is None:
            loader = ExternalBookLoader()
        try:
            value = evaluate_formula(
                ast, sheets, sheet_name, load_external=loader, book_path=view.path,
                formula_sheets=formulas, current_cell=current_cell,
            )
        except UnsupportedFormula as exc:
            return None, str(exc)
        except (ZeroDivisionError, ValueError, TypeError, IndexError, KeyError) as exc:
            return None, f"本地计算出错：{exc}"
        finally:
            if owns_loader:
                self._close_loader(loader)
        return value, None

    def _close_loader(self, loader: ExternalBookLoader) -> None:
        try:
            if self.logger and loader.loaded:
                self.logger.info(
                    "本地试算加载外部工作簿：%s", "、".join(str(p) for p in loader.loaded)
                )
        finally:
            loader.close()


def _merge_validation(
    base: ValidationResult | None, extra: ValidationResult
) -> ValidationResult:
    """多个候选公式共用一份校验结果：警告累加，ok 取与。"""
    if base is None:
        return extra
    base.ok = base.ok and extra.ok
    base.errors.extend(extra.errors)
    base.warnings.extend(extra.warnings)
    base.unrepairable = base.unrepairable or extra.unrepairable
    base.functions.extend(f for f in extra.functions if f not in base.functions)
    base.refs.extend(r for r in extra.refs if r not in base.refs)
    return base
