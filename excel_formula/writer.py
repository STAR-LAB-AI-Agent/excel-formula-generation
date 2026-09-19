"""公式写入：备份 -> 写入 -> 保存，并把相对引用填充展开为多个单元格。"""
from __future__ import annotations

import os
import re
import shutil
import tempfile
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import PILImage
from openpyxl.formula.translate import Translator
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from .config import MAX_WRITE_CELLS, Settings
from .data_validation import (
    WRITE_LOCK, DropdownSpec, apply_spec, extension_warning, file_version, require_version,
)
from .evaluator import needs_array_formula
from .formula_parser import FormulaSyntaxError, parse_formula, parse_range, parse_target


@dataclass
class CellWrite:
    """一次待写入操作。formula 为空表示写常量值（value）：分类名、标题等非公式内容。

    old_formula/old_value/predicted 均仅用于预览与回溯。
    """

    sheet: str
    cell: str
    formula: str
    value: object = None
    explanation: str = ""
    old_formula: str | None = None
    old_value: object = None
    predicted: str | None = None
    predicted_note: str | None = None

    @property
    def content(self) -> object:
        """写入单元格的实际内容：公式或常量值（二选一）。"""
        return self.formula if self.formula else self.value

    def to_dict(self) -> dict:
        return {
            "sheet": self.sheet,
            "cell": self.cell,
            "formula": self.formula,
            "value": self.value.isoformat() if isinstance(self.value, (date, time)) else self.value,
            "explanation": self.explanation,
            "overwrites": self.old_formula or (
                None if self.old_value is None else str(self.old_value)
            ),
            "predicted_value": self.predicted,
            "predicted_note": self.predicted_note,
        }


def expand_fill(sheet: str, target: str, formula: str, fill_to: str | None) -> list[CellWrite]:
    """把 target + fill_to 展开成一列（或一行）单元格，相对引用自动平移。"""
    _, col, row = parse_target(target)
    base = f"{get_column_letter(col)}{row}"
    writes = [CellWrite(sheet=sheet, cell=base, formula=formula)]
    if not fill_to:
        return writes

    _, end_col, end_row = parse_target(fill_to)
    if end_col != col and end_row != row:
        raise ValueError(f"填充范围 {target}:{fill_to} 必须在同一行或同一列")

    cells: list[str] = []
    if end_row != row:
        step = 1 if end_row > row else -1
        cells = [f"{get_column_letter(col)}{r}" for r in range(row + step, end_row + step, step)]
    elif end_col != col:
        step = 1 if end_col > col else -1
        cells = [f"{get_column_letter(c)}{row}" for c in range(col + step, end_col + step, step)]

    if len(cells) + 1 > MAX_WRITE_CELLS:
        raise ValueError(
            f"一次最多写入 {MAX_WRITE_CELLS} 个单元格，当前请求 {len(cells) + 1} 个，请缩小范围"
        )
    translator = Translator(formula, origin=base)
    for cell in cells:
        writes.append(CellWrite(sheet=sheet, cell=cell, formula=translator.translate_formula(cell)))
    return writes


def _is_array_semantics(formula: str) -> bool:
    """公式是否含数组语义运算（区域参与一元/二元运算、IF 数组形式）。

    这类公式本地按数组语义求值；若以裸公式文本存进 .xlsx，Excel/WPS 打开时
    会按传统"隐式交叉"只取交叉单值，与本地结果及网页展示不一致。写入时改用
    数组公式（CSE）形态，各版本 Excel/WPS 都按数组语义还原。解析失败按非数组
    处理——公式语法错误由校验环节负责报告。
    """
    try:
        return needs_array_formula(parse_formula(formula))
    except FormulaSyntaxError:
        return False


# ---------------------------------------------------------------- 新建表格格式
# 表头浅蓝底 + 整表细边框：新建表格时由 pipeline 给出几何，apply_writes 负责套用。
_TABLE_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFD9E1F2")
_TABLE_SIDE = Side(style="thin")
_TABLE_BORDER = Border(left=_TABLE_SIDE, right=_TABLE_SIDE, top=_TABLE_SIDE, bottom=_TABLE_SIDE)


@dataclass
class TableFormat:
    """一张表格的格式：整表范围加细边框，表头行可选铺浅蓝底。

    header_range 为 None 表示只加边框（“把范围框起来”这类显式请求）；
    force_border 为 True 时已有样式的格子也补边框——用户点名要框，不能默默跳过。
    """

    sheet: str
    table_range: str
    header_range: str | None = None
    force_border: bool = False

    def label(self) -> str:
        if self.header_range:
            return f"{self.sheet}!{self.table_range}（表头 {self.header_range} 浅蓝底、范围加细边框）"
        return f"{self.sheet}!{self.table_range}（范围加细边框）"

    def to_dict(self) -> dict:
        return {
            "sheet": self.sheet,
            "header_range": self.header_range,
            "table_range": self.table_range,
            "label": self.label(),
        }


def _ref_cells(ref) -> list[tuple[int, int]]:
    """把区域引用展开为 (行, 列) 列表。"""
    return [
        (row, col)
        for row in range(ref.row1, ref.row2 + 1)
        for col in range(ref.col1, ref.col2 + 1)
    ]


def backup_file(path: Path, settings: Settings) -> Path | None:
    """写入前备份到 backups/ 目录，便于人工回退。"""
    if not settings.backup:
        return None
    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f") + "_" + uuid.uuid4().hex[:8]
    destination = backup_dir / f"{path.stem}_{stamp}{path.suffix}"
    shutil.copy2(path, destination)
    return destination


def _inherited_style(sheet_obj, row: int, col: int):
    """取同行最近单元格的样式：先向左，再向右。

    只认“已有格式”的格子（has_style），空白格不参与，避免凭空造样式；
    用 _cells 直取，不通过 ws.cell() 以免把扫描过的空单元格物化进文件。
    """
    columns = list(range(col - 1, 0, -1)) + list(range(col + 1, sheet_obj.max_column + 1))
    for column in columns:
        cell = sheet_obj._cells.get((row, column))
        if cell is not None and cell.has_style:
            return cell._style, cell.coordinate
    return None, None


# ---------------------------------------------------------------- 富对象保真
# openpyxl 是"解析-重建"式保存：它模型里没有的对象会被静默丢弃。
# 这里在写入前后各扫一次 zip 部件，用于提前拦截与事后报警。
_RICH_PART_PATTERNS = {
    "charts": re.compile(r"^xl/charts/chart\d+\.xml$"),
    "pivots": re.compile(r"^xl/pivotTables/pivotTable\d+\.xml$"),
    "drawings": re.compile(r"^xl/drawings/drawing\d+\.xml$"),
}
# 丢失/保留清单展示用的中文名与量词
_RICH_PART_LABELS = {
    "charts": ("图表", "个"),
    "images": ("图片", "张"),
    "pivots": ("数据透视表", "个"),
}


def inspect_rich_objects(path: Path) -> dict[str, int]:
    """统计文件内的富对象部件：图表 / 图片 / 数据透视表 / 绘图。

    WMF 图片单独计数：openpyxl 不支持该格式、注定丢失，不计入常规图片。
    """
    counts = {"charts": 0, "images": 0, "pivots": 0, "drawings": 0, "wmf_images": 0}
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return counts
    for name in names:
        if name.startswith("xl/media/"):
            if name.lower().endswith(".wmf"):
                counts["wmf_images"] += 1
            else:
                counts["images"] += 1
            continue
        for key, pattern in _RICH_PART_PATTERNS.items():
            if pattern.match(name):
                counts[key] += 1
                break
    return counts


def _fidelity_report(before: dict[str, int], after: dict[str, int]) -> dict:
    """对比保存前后的部件数量，生成丢失与保留清单。"""
    lost: list[str] = []
    preserved: list[str] = []
    for key, (label, unit) in _RICH_PART_LABELS.items():
        missing = before[key] - after[key]
        if missing > 0:
            lost.append(f"{label} {missing} {unit}")
        elif before[key]:
            preserved.append(f"{label} {after[key]} {unit}")
    drawing_missing = before["drawings"] - after["drawings"]
    if drawing_missing > 0 and not lost:
        # 图表/图片丢失时绘图容器会跟着少，这里避免重复报警
        lost.append(f"绘图对象 {drawing_missing} 个")
    return {"lost": lost, "preserved": preserved}


def apply_writes(
    path: Path,
    writes: list[CellWrite],
    settings: Settings,
    *,
    logger=None,
    output: Path | None = None,
    tables: list[TableFormat] | None = None,
    dropdown: DropdownSpec | None = None,
    expected_version: str | None = None,
) -> dict:
    """所有保存共用锁，版本检查与原子替换之间不能插入另一次写入。"""
    with WRITE_LOCK:
        path = settings.resolve_path(path)
        output = settings.resolve_path(output, must_exist=False) if output else None
        version = file_version(path)
        if expected_version is not None:
            require_version(path, expected_version)
        if dropdown is not None or expected_version is not None:
            warning = extension_warning(path)
            if warning:
                raise ValueError(warning)
        return _apply_writes_locked(
            path, writes, settings, logger=logger, output=output, tables=tables,
            dropdown=dropdown, original_version=version,
        )


def _apply_writes_locked(
    path: Path,
    writes: list[CellWrite],
    settings: Settings,
    *,
    logger=None,
    output: Path | None = None,
    tables: list[TableFormat] | None = None,
    dropdown: DropdownSpec | None = None,
    original_version: str,
) -> dict:
    """执行写入。output 为空表示原地保存（会先备份）。

    tables 为表格格式（表头浅蓝底、整表加框；“仅加框”时 header_range 为空），
    为空表示不套用。允许只给 tables 不给 writes——显式“框起来”不改单元格内容。
    """
    if not writes and not tables and dropdown is None:
        raise ValueError("没有需要写入的内容")
    if dropdown and dropdown.unchanged and not writes and not tables and output is None:
        return {"file": str(path), "backup": None, "written": [], "count": 0,
                "styled": [], "tables": [], "dropdown": dropdown.to_dict(),
                "unchanged": True, "fidelity": {"warnings": [], "preserved": []}}
    if len(writes) > MAX_WRITE_CELLS:
        raise ValueError(f"一次最多写入 {MAX_WRITE_CELLS} 个单元格")

    # 写前检测：注定保不住的对象提前拦住或提示
    before = inspect_rich_objects(path)
    warnings: list[str] = []
    if before["images"] and not PILImage:
        raise ValueError(
            f"{path.name} 内含 {before['images']} 张图片，当前环境未安装 Pillow，"
            "保存会导致图片全部丢失；请先执行 pip install Pillow 后重试"
        )
    if before["wmf_images"]:
        warnings.append(
            f"{path.name} 内含 {before['wmf_images']} 张 WMF 图片，"
            "openpyxl 不支持该格式，保存后会丢失"
        )

    target_path = output or path
    backup = None

    workbook = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    styled: list[str] = []
    # 新建表格的几何：表内单元格不参与相邻格式沿用，统一走表格格式
    table_plans: list[tuple[TableFormat, set, list[tuple[int, int]]]] = []
    table_cells: dict[str, set[str]] = {}
    for table in tables or []:
        if table.sheet not in workbook.sheetnames:
            raise KeyError(f"工作表 {table.sheet!r} 不存在")
        cells = _ref_cells(parse_range(table.table_range))
        header_ref = parse_range(table.header_range) if table.header_range else None
        header_cells = {
            (row, col) for row, col in cells if header_ref and header_ref.contains(col, row)
        }
        table_plans.append((table, header_cells, cells))
        table_cells.setdefault(table.sheet, set()).update(
            f"{get_column_letter(col)}{row}" for row, col in cells
        )
    try:
        for item in writes:
            if item.sheet not in workbook.sheetnames:
                raise KeyError(f"工作表 {item.sheet!r} 不存在")
            content = item.content
            if content is None or content == "":
                raise ValueError(f"{item.sheet}!{item.cell} 没有可写入的内容")
            sheet_obj = workbook[item.sheet]
            target = sheet_obj[item.cell]
            # 新增单元格沿用同行相邻单元格的格式（表头颜色、边框等）；
            # 已自带格式的格子保持原样，不覆盖用户自己设置过的样式；
            # 属于新建表格的格子跳过沿用，由下方的表格格式统一处理。
            if not target.has_style and target.coordinate not in table_cells.get(item.sheet, ()):
                style, source = _inherited_style(sheet_obj, target.row, target.column)
                if style is not None:
                    target._style = style
                    styled.append(f"{target.coordinate} ← {source}")
            if item.formula and _is_array_semantics(item.formula):
                # 数组语义公式以 CSE 形态写入，保证 Excel/WPS 按数组而非隐式交叉求值
                target.value = ArrayFormula(item.cell, item.formula)
            else:
                target.value = content
            if not item.formula and isinstance(content, str):
                target.data_type = "s"  # 下拉选项即使以 = 开头，也只能作为文本保存
        # 表格格式：表头铺浅蓝底（如有），整表（含未写入的边框格）加细边框；
        # 自动新建表格时已有样式的格子保持原样，避免覆盖用户手动格式；
        # 显式“框起来”请求（force_border）则点名要框，已有样式也补上边框。
        styled_tables: list[str] = []
        for table, header_cells, cells in table_plans:
            sheet_obj = workbook[table.sheet]
            touched = 0
            for row, col in cells:
                cell = sheet_obj.cell(row=row, column=col)
                if cell.has_style and not table.force_border:
                    continue
                if (row, col) in header_cells:
                    cell.fill = _TABLE_HEADER_FILL
                cell.border = _TABLE_BORDER
                touched += 1
            if touched:
                styled_tables.append(table.label())
        if dropdown is not None:
            apply_spec(workbook, dropdown)
        require_version(path, original_version)
        backup = backup_file(path, settings) if target_path == path else None
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=target_path.parent, prefix=".excelcr-", suffix=target_path.suffix, delete=False,
            ) as stream:
                temporary = Path(stream.name)
            workbook.save(temporary)
            require_version(path, original_version)
            os.replace(temporary, target_path)
        except PermissionError as exc:
            raise PermissionError(
                f"无法写入 {target_path.name}，文件可能正被 Excel 打开，请关闭后重试"
            ) from exc
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
    finally:
        workbook.close()

    # 写后校验：部件数量对比，防止 openpyxl 静默丢弃对象
    after = inspect_rich_objects(target_path)
    report = _fidelity_report(before, after)
    if report["lost"]:
        if backup:
            advice = "；原文件已备份，可人工回退"
        elif target_path != path:
            advice = "；原文件未改动，缺失仅发生在输出文件"
        else:
            advice = ""
        warnings.append(
            "保存后检测到对象丢失：" + "、".join(report["lost"]) + "（openpyxl 无法完整保留）" + advice
        )

    if logger:
        logger.info(
            "写入完成 file=%s cells=%s backup=%s",
            target_path.name,
            ",".join(f"{w.sheet}!{w.cell}" for w in writes) or (
                f"下拉列表 {dropdown.sheet}!{dropdown.cell_range}" if dropdown else "无（仅套用表格格式）"
            ),
            backup.name if backup else "无",
        )
        if styled:
            logger.info("样式沿用：%s", "，".join(styled))
        if styled_tables:
            logger.info("表格格式：%s", "，".join(styled_tables))
        if report["preserved"]:
            logger.info("富对象已保留：%s", "、".join(report["preserved"]))
        for warning in warnings:
            logger.warning("保真提示：%s", warning)

    return {
        "file": str(target_path),
        "backup": str(backup) if backup else None,
        "written": [w.to_dict() for w in writes],
        "count": len(writes),
        "styled": styled,
        "tables": styled_tables,
        "dropdown": dropdown.to_dict() if dropdown else None,
        "fidelity": {
            "before": before,
            "after": after,
            "preserved": report["preserved"],
            "warnings": warnings,
        },
    }
