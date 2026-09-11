"""Excel 公式的分词器与递归下降语法分析器（不依赖模型，纯确定性实现）。

产出的 AST 同时供 validator（静态检查）和 evaluator（独立求值）使用。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.utils import column_index_from_string, get_column_letter

from .config import MAX_EXCEL_COLUMN, MAX_EXCEL_ROW


class FormulaSyntaxError(ValueError):
    """公式语法错误，携带出错位置以便回传模型修复。"""

    def __init__(self, message: str, position: int = -1):
        super().__init__(message if position < 0 else f"{message}（位置 {position}）")
        self.message = message
        self.position = position


# ------------------------------------------------------------------ AST 节点
@dataclass
class Number:
    value: float


@dataclass
class Text:
    value: str


@dataclass
class Bool:
    value: bool


@dataclass
class ErrorLiteral:
    value: str


@dataclass
class Ref:
    """单元格或区域引用。列/行下标均为 1 起始，包含边界。"""

    text: str
    col1: int
    row1: int
    col2: int
    row2: int
    sheet: str | None = None
    is_range: bool = False
    whole_column: bool = False

    @property
    def cell_count(self) -> int:
        return (self.col2 - self.col1 + 1) * (self.row2 - self.row1 + 1)

    def contains(self, col: int, row: int) -> bool:
        return self.col1 <= col <= self.col2 and self.row1 <= row <= self.row2

    def normalized(self) -> str:
        head = f"{self.sheet}!" if self.sheet else ""
        start = f"{get_column_letter(self.col1)}{self.row1}"
        if not self.is_range:
            return head + start
        return f"{head}{start}:{get_column_letter(self.col2)}{self.row2}"


@dataclass
class DefinedName:
    """定义名称 / 表名等非单元格标识符。"""

    name: str
    position: int = -1


@dataclass
class FuncCall:
    name: str
    args: list = field(default_factory=list)
    position: int = -1


@dataclass
class Unary:
    op: str
    operand: object


@dataclass
class Binary:
    op: str
    left: object
    right: object


# ------------------------------------------------------------------ 分词
@dataclass
class Token:
    kind: str
    text: str
    position: int


_SHEET = r"(?:'(?:[^']|'')+'|[A-Za-z_\u4e00-\u9fff][A-Za-z0-9_.\u4e00-\u9fff]*)!"
_A1 = r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}"

TOKEN_RE = re.compile(
    rf"""
      (?P<ws>\s+)
    | (?P<string>"(?:[^"]|"")*")
    | (?P<unterminated>"(?:[^"]|"")*$)
    | (?P<error>\#(?:REF!|VALUE!|DIV/0!|NAME\?|N/A|NULL!|NUM!|GETTING_DATA))
    | (?P<sheet>{_SHEET})
    | (?P<colrange>\$?[A-Za-z]{{1,3}}:\$?[A-Za-z]{{1,3}}(?![A-Za-z0-9$]))
    | (?P<range>{_A1}\s*:\s*{_A1})
    | (?P<bool>\b(?:TRUE|FALSE)\b(?!\s*\())
    | (?P<func>[A-Za-z_\u4e00-\u9fff][A-Za-z0-9_.\u4e00-\u9fff]*)\s*(?=\()
    | (?P<cell>{_A1}(?![A-Za-z0-9_$]))
    | (?P<number>(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?)
    | (?P<op><=|>=|<>|[-+*/^&=<>])
    | (?P<percent>%)
    | (?P<lparen>\()
    | (?P<rparen>\))
    | (?P<sep>[,;])
    | (?P<name>[A-Za-z_\u4e00-\u9fff\[][A-Za-z0-9_.\u4e00-\u9fff\[\]]*)
    """,
    re.VERBOSE | re.IGNORECASE,
)


def tokenize(formula: str) -> list[Token]:
    """把公式体（不含开头的 =）切成 token 序列。"""
    tokens: list[Token] = []
    pos = 0
    length = len(formula)
    while pos < length:
        match = TOKEN_RE.match(formula, pos)
        if match is None:
            raise FormulaSyntaxError(f"无法识别的字符 {formula[pos]!r}", pos + 1)
        kind = match.lastgroup
        # func 分组带有 lookahead，lastgroup 可能落在内部分组上，统一用 groupdict 判断
        if kind is None or match.group(kind) is None:
            kind = next(k for k, v in match.groupdict().items() if v is not None)
        text = match.group(kind)
        pos = match.end()
        if kind == "ws":
            continue
        if kind == "unterminated":
            raise FormulaSyntaxError("字符串引号未闭合", match.start() + 1)
        tokens.append(Token(kind=kind, text=text, position=match.start() + 1))
    return tokens


# ------------------------------------------------------------------ 语法分析
_COMPARISON_OPS = {"=", "<>", "<", ">", "<=", ">="}


class Parser:
    def __init__(self, tokens: list[Token]):
        self.tokens = tokens
        self.index = 0
        self.pending_sheet: str | None = None

    # ------------------------------------------------------------ 基础操作
    def peek(self) -> Token | None:
        return self.tokens[self.index] if self.index < len(self.tokens) else None

    def next(self) -> Token:
        token = self.peek()
        if token is None:
            raise FormulaSyntaxError("公式在此处意外结束")
        self.index += 1
        return token

    def accept(self, kind: str, text: str | None = None) -> Token | None:
        token = self.peek()
        if token and token.kind == kind and (text is None or token.text.upper() == text.upper()):
            self.index += 1
            return token
        return None

    def expect(self, kind: str, label: str) -> Token:
        token = self.accept(kind)
        if token is None:
            current = self.peek()
            where = current.position if current else -1
            got = current.text if current else "公式结尾"
            raise FormulaSyntaxError(f"此处应为 {label}，实际是 {got!r}", where)
        return token

    # ------------------------------------------------------------ 表达式层级
    def parse(self) -> object:
        node = self.parse_comparison()
        remaining = self.peek()
        if remaining is not None:
            raise FormulaSyntaxError(f"多余内容 {remaining.text!r}", remaining.position)
        return node

    def parse_comparison(self) -> object:
        node = self.parse_concat()
        while True:
            token = self.peek()
            if token and token.kind == "op" and token.text in _COMPARISON_OPS:
                self.index += 1
                node = Binary(token.text, node, self.parse_concat())
            else:
                return node

    def parse_concat(self) -> object:
        node = self.parse_additive()
        while True:
            token = self.peek()
            if token and token.kind == "op" and token.text == "&":
                self.index += 1
                node = Binary("&", node, self.parse_additive())
            else:
                return node

    def parse_additive(self) -> object:
        node = self.parse_multiplicative()
        while True:
            token = self.peek()
            if token and token.kind == "op" and token.text in {"+", "-"}:
                self.index += 1
                node = Binary(token.text, node, self.parse_multiplicative())
            else:
                return node

    def parse_multiplicative(self) -> object:
        node = self.parse_power()
        while True:
            token = self.peek()
            if token and token.kind == "op" and token.text in {"*", "/"}:
                self.index += 1
                node = Binary(token.text, node, self.parse_power())
            else:
                return node

    def parse_power(self) -> object:
        node = self.parse_unary()
        token = self.peek()
        if token and token.kind == "op" and token.text == "^":
            self.index += 1
            return Binary("^", node, self.parse_power())
        return node

    def parse_unary(self) -> object:
        token = self.peek()
        if token and token.kind == "op" and token.text in {"-", "+"}:
            self.index += 1
            return Unary(token.text, self.parse_unary())
        return self.parse_postfix()

    def parse_postfix(self) -> object:
        node = self.parse_primary()
        while self.accept("percent"):
            node = Unary("%", node)
        return node

    def parse_primary(self) -> object:
        token = self.next()
        kind = token.kind

        if kind == "number":
            return Number(float(token.text))
        if kind == "string":
            return Text(token.text[1:-1].replace('""', '"'))
        if kind == "bool":
            return Bool(token.text.upper() == "TRUE")
        if kind == "error":
            return ErrorLiteral(token.text.upper())
        if kind == "lparen":
            node = self.parse_comparison()
            self.expect("rparen", "右括号 )")
            return node
        if kind == "sheet":
            sheet = _clean_sheet(token.text)
            nxt = self.peek()
            if nxt is None or nxt.kind not in {"cell", "range", "colrange"}:
                raise FormulaSyntaxError(f"工作表引用 {token.text!r} 后缺少单元格地址", token.position)
            self.index += 1
            return _make_ref(nxt, sheet)
        if kind in {"cell", "range", "colrange"}:
            return _make_ref(token, None)
        if kind == "func":
            self.expect("lparen", "左括号 (")
            args: list[object] = []
            if self.peek() and self.peek().kind == "rparen":
                self.index += 1
                return FuncCall(token.text.upper(), args, token.position)
            while True:
                args.append(self.parse_comparison())
                if self.accept("sep"):
                    continue
                break
            self.expect("rparen", "右括号 )")
            return FuncCall(token.text.upper(), args, token.position)
        if kind == "name":
            return DefinedName(token.text, token.position)

        raise FormulaSyntaxError(f"意外的符号 {token.text!r}", token.position)


def parse_formula(formula: str) -> object:
    """解析完整公式（必须以 = 开头），返回 AST 根节点。"""
    if not isinstance(formula, str):
        raise FormulaSyntaxError("公式必须是字符串")
    text = formula.strip()
    if not text:
        raise FormulaSyntaxError("公式为空")
    if not text.startswith("="):
        raise FormulaSyntaxError("公式必须以 = 开头")
    body = text[1:].strip()
    if not body:
        raise FormulaSyntaxError("= 后面没有内容")
    if text.count('"') % 2 == 1:
        raise FormulaSyntaxError("双引号数量不成对")
    return Parser(tokenize(body)).parse()


def walk(node: object):
    """深度优先遍历 AST。"""
    yield node
    if isinstance(node, FuncCall):
        for arg in node.args:
            yield from walk(arg)
    elif isinstance(node, Binary):
        yield from walk(node.left)
        yield from walk(node.right)
    elif isinstance(node, Unary):
        yield from walk(node.operand)


# ------------------------------------------------------------------ 引用工具
def _clean_sheet(text: str) -> str:
    name = text[:-1]  # 去掉结尾的 !
    if name.startswith("'") and name.endswith("'"):
        name = name[1:-1].replace("''", "'")
    return name


def _split_coord(coord: str) -> tuple[int, int]:
    match = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]{1,7})", coord.strip())
    if not match:
        raise FormulaSyntaxError(f"非法单元格地址 {coord!r}")
    col = column_index_from_string(match.group(1).upper())
    row = int(match.group(2))
    if col > MAX_EXCEL_COLUMN or row > MAX_EXCEL_ROW:
        raise FormulaSyntaxError(f"单元格地址超出 Excel 上限：{coord}")
    return col, row


def _make_ref(token: Token, sheet: str | None) -> Ref:
    text = token.text.replace(" ", "")
    if token.kind == "cell":
        col, row = _split_coord(text)
        return Ref(text=text, col1=col, row1=row, col2=col, row2=row, sheet=sheet)
    if token.kind == "range":
        left, right = text.split(":")
        col1, row1 = _split_coord(left)
        col2, row2 = _split_coord(right)
        return Ref(
            text=text,
            col1=min(col1, col2),
            row1=min(row1, row2),
            col2=max(col1, col2),
            row2=max(row1, row2),
            sheet=sheet,
            is_range=True,
        )
    # colrange，例如 A:A
    left, right = text.split(":")
    col1 = column_index_from_string(left.replace("$", "").upper())
    col2 = column_index_from_string(right.replace("$", "").upper())
    return Ref(
        text=text,
        col1=min(col1, col2),
        row1=1,
        col2=max(col1, col2),
        row2=MAX_EXCEL_ROW,
        sheet=sheet,
        is_range=True,
        whole_column=True,
    )


def parse_target(target: str) -> tuple[str | None, int, int]:
    """解析目标单元格，如 'G2' 或 'Sheet1!G2'，返回 (工作表, 列, 行)。"""
    text = (target or "").strip().replace(" ", "")
    if not text:
        raise FormulaSyntaxError("目标单元格不能为空")
    sheet = None
    if "!" in text:
        sheet_part, _, text = text.rpartition("!")
        sheet = _clean_sheet(sheet_part + "!")
    if ":" in text:
        raise FormulaSyntaxError(f"目标必须是单个单元格，不能是区域：{target}")
    col, row = _split_coord(text)
    return sheet, col, row
