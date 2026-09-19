"""公式独立求值器：在 Python 中重算一遍公式，用于写入前的"公式自动验证"。

只支持常用函数子集；遇到不支持的函数会抛出 UnsupportedFormula，
上层据此把结果标记为"未验证"，而不是阻断流程。
"""
from __future__ import annotations

import calendar
import datetime as _dt
import math
import re
import statistics
from dataclasses import dataclass
from pathlib import Path

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .excel_reader import cell_formula_text
from .formula_parser import (
    ArrayLiteral,
    Binary,
    Bool,
    DefinedName,
    ErrorLiteral,
    FormulaSyntaxError,
    FuncCall,
    is_external_sheet,
    Number,
    parse_formula,
    Ref,
    Text,
    Unary,
)


class UnsupportedFormula(RuntimeError):
    """本地求值器不支持该公式（函数未实现或引用无缓存值）。"""


@dataclass(frozen=True)
class ExcelError:
    code: str

    def __str__(self) -> str:
        return self.code


@dataclass(frozen=True)
class DateValue:
    """日期/时间值的只读包装。

    日期在 Excel 里本质是序列号（1900 日期系统，1899-12-30 为 0）。求值全程按
    序列号参与算术、比较与汇总；与单元格数字格式相关的行为（显示格式）无法本地
    还原，日期与无法解析的文本比较时仍抛 UnsupportedFormula 标"未验证"。
    """

    raw: object

    def __str__(self) -> str:
        return format_value(self)


_EXCEL_EPOCH = _dt.datetime(1899, 12, 30)  # 1900 日期系统的序列号 0
DATE_HINT = "日期与无法解析的文本无法比较"

# 文本形式的日期：DATEVALUE、日期条件（如 ">2025/1/5"）等场景按这些格式解析
_DATE_TEXT_FORMATS = (
    "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日",
    "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M",
)


def _to_serial(value) -> float:
    """日期/时间对象 → Excel 序列号：datetime 带时间小数，date 为整数天，time 为日内小数。"""
    if isinstance(value, _dt.datetime):
        return (value - _EXCEL_EPOCH).total_seconds() / 86400
    if isinstance(value, _dt.date):
        return float((value - _EXCEL_EPOCH.date()).days)
    return (value.hour * 3600 + value.minute * 60 + value.second) / 86400


def _from_serial(serial: float) -> DateValue:
    """序列号 → 日期/时间值：整数返回 date，含小数返回 datetime。"""
    days = math.floor(serial)
    fraction = serial - days
    base = _EXCEL_EPOCH.date() + _dt.timedelta(days=days)
    if fraction < 1e-9:
        return DateValue(base)
    clock = _dt.timedelta(seconds=round(fraction * 86400))
    return DateValue(_dt.datetime.combine(base, _dt.time()) + clock)


def _as_date(value) -> _dt.date:
    """把 date/datetime 统一成 date。"""
    return value.date() if isinstance(value, _dt.datetime) else value


def _parse_date_text(text: str):
    """按常见格式解析日期文本，失败返回 None。"""
    raw = text.strip()
    for fmt in _DATE_TEXT_FORMATS:
        try:
            return _dt.datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None

# 跨工作簿递归求值的最大嵌套深度：防御异常的超长引用链
_MAX_EXTERNAL_DEPTH = 24

DIV0 = ExcelError("#DIV/0!")
VALUE_ERR = ExcelError("#VALUE!")
NA_ERR = ExcelError("#N/A")
NUM_ERR = ExcelError("#NUM!")


class Evaluator:
    """按 AST 求值。所有取值来自 data_only 工作簿的缓存值。"""

    def __init__(
        self,
        sheets: dict[str, Worksheet],
        default_sheet: str,
        load_external=None,
        book_path: str | Path | None = None,
        current_cell: tuple[int, int] | None = None,
        _stack: set | None = None,
        _book=None,
        _formulas: dict[str, Worksheet] | None = None,
        _expand_uncached: bool = False,
        _cache: dict | None = None,
    ):
        self.sheets = sheets
        self.default_sheet = default_sheet
        # 跨工作簿支持：load_external(sheet_name, parent_dir) → (外部工作簿, 表名) | (None, 原因)
        self.load_external = load_external
        self.book_path = Path(book_path) if book_path else None
        # ROW()/COLUMN() 无参形式所需的公式所在单元格 (行, 列)。递归求值链里，
        # 被引用的公式用自己所在的位置；调用方给不出位置时为 None，这两个函数
        # 按"未验证"处理，而不是猜一个行号。
        self.current_cell = current_cell
        # 递归求值外部公式时的循环检测：整条求值链共享同一个集合
        self._stack: set = set() if _stack is None else _stack
        # 当前上下文来自外部工作簿时：_book 为其句柄，_formulas 为其公式视图（用于继续递归）
        self._book = _book
        self._formulas = _formulas
        # 主簿里没有缓存的公式是否递归求值（展示场景开启；校验场景保持"未验证"）
        self._expand_uncached = _expand_uncached
        # 主簿递归求值的结果缓存：同一次顶层计算里同一格只算一次
        self._cache: dict = {} if _cache is None else _cache

    # -------------------------------------------------------------- 对外入口
    def evaluate(self, node: object):
        value = self._eval(node)
        if isinstance(value, list):  # 顶层拿到区域时取首个值，与 Excel 隐式交叉行为近似
            flat = _flatten(value)
            return flat[0] if flat else None
        return value

    # -------------------------------------------------------------- 递归求值
    def _eval(self, node: object):
        if isinstance(node, Number):
            return node.value
        if isinstance(node, Text):
            return node.value
        if isinstance(node, Bool):
            return node.value
        if isinstance(node, ErrorLiteral):
            return ExcelError(node.value)
        if isinstance(node, ArrayLiteral):
            return [[self._eval(item) for item in row] for row in node.rows]
        if isinstance(node, Ref):
            return self._eval_ref(node)
        if isinstance(node, Unary):
            return self._eval_unary(node)
        if isinstance(node, Binary):
            return self._eval_binary(node)
        if isinstance(node, FuncCall):
            return self._eval_func(node)
        if isinstance(node, DefinedName):
            raise UnsupportedFormula(f"不支持定义名称 {node.name}")
        raise UnsupportedFormula(f"不支持的表达式节点 {type(node).__name__}")

    def _eval_ref(self, ref: Ref):
        sheet_name = ref.sheet or self.default_sheet
        ws = _lookup_sheet(self.sheets, sheet_name)
        if ws is None:
            if is_external_sheet(sheet_name):
                return self._eval_external_ref(sheet_name, ref)
            raise UnsupportedFormula(f"找不到工作表 {sheet_name}")
        row2 = max(ws.max_row or 1, 1) if ref.whole_column else ref.row2
        formulas_ws = _lookup_sheet(self._formulas, sheet_name) if self._formulas else None
        return self._collect(
            ref,
            row2,
            lambda r, c: self._cell_value(self._book, sheet_name, ws, formulas_ws, r, c),
        )

    def _collect(self, ref: Ref, row2: int, get_cell):
        """按引用取值：单元格返回标量，区域返回二维列表。"""
        if not ref.is_range:
            return get_cell(ref.row1, ref.col1)
        return [
            [get_cell(row, col) for col in range(ref.col1, ref.col2 + 1)]
            for row in range(ref.row1, row2 + 1)
        ]

    def _eval_external_ref(self, sheet_name: str, ref: Ref):
        """跨工作簿引用：经外部加载器读缓存值，公式无缓存时递归本地求值。"""
        if self.load_external is None:
            raise UnsupportedFormula(
                f"跨工作簿引用 {sheet_name} 无法本地验证（未启用外部工作簿加载）"
            )
        parent = self.book_path.parent if self.book_path else None
        book, inner = self.load_external(sheet_name, parent)
        if book is None:
            raise UnsupportedFormula(f"跨工作簿引用 {sheet_name} 无法本地验证：{inner}")
        ws_values = book.find_values(inner)
        if ws_values is None:
            raise UnsupportedFormula(
                f"外部工作簿 {book.path.name} 中找不到工作表 {inner!r}"
            )
        ws_formulas = book.find_formulas(inner)
        row2 = max(ws_values.max_row or 1, 1) if ref.whole_column else ref.row2
        return self._collect(
            ref,
            row2,
            lambda r, c: self._cell_value(book, inner, ws_values, ws_formulas, r, c),
        )

    def _cell_value(self, book, sheet_name: str, ws_values, ws_formulas, row: int, col: int):
        """读工作簿单元格：优先缓存值；公式无缓存时——外部簿或 expand_uncached
        模式递归本地求值，否则由 _coerce 抛出"未验证"（校验场景的保守语义）。
        """
        cell = ws_values.cell(row=row, column=col)
        value = cell.value
        if cell.data_type == "e":
            return ExcelError(str(value))
        if ws_formulas is not None and (
            value is None or (isinstance(value, str) and value.startswith("="))
        ):
            raw = cell_formula_text(ws_formulas.cell(row=row, column=col))
            if raw is not None and raw.startswith("="):
                if book is not None:
                    return self._eval_external_formula(book, sheet_name, row, col, raw)
                if self._expand_uncached:
                    return self._eval_local_formula(sheet_name, row, col, raw)
                return _coerce(raw)
        return _coerce(value)

    def _eval_local_formula(self, sheet_name: str, row: int, col: int, formula_text: str):
        """主簿内没有缓存值的公式递归本地求值（expand_uncached 模式）。

        展示场景（如网页表格）希望像 Excel 一样给出结果，因此对链式依赖逐层
        展开；校验场景保持默认关闭——引用未计算的公式即报"未验证"。带循环
        检测、深度上限与结果缓存，语义与跨工作簿递归保持一致。
        """
        marker = f"{sheet_name}!{get_column_letter(col)}{row}"
        key = ("<self>", sheet_name.casefold(), row, col)
        if key in self._cache:
            return self._cache[key]
        if key in self._stack:
            raise UnsupportedFormula(f"循环引用：{marker} 在求值链上重复出现")
        if len(self._stack) >= _MAX_EXTERNAL_DEPTH:
            raise UnsupportedFormula(f"公式引用链过深（超过 {_MAX_EXTERNAL_DEPTH} 层）：{marker}")
        self._stack.add(key)
        try:
            try:
                ast = parse_formula(formula_text)
            except FormulaSyntaxError as exc:
                raise UnsupportedFormula(f"公式 {marker} 无法解析：{exc}") from exc
            nested = Evaluator(
                self.sheets,
                sheet_name,
                load_external=self.load_external,
                book_path=self.book_path,
                current_cell=(row, col),
                _stack=self._stack,
                _formulas=self._formulas,
                _expand_uncached=True,
                _cache=self._cache,
            )
            value = nested.evaluate(ast)
        finally:
            self._stack.discard(key)
        self._cache[key] = value
        return value

    def _eval_external_formula(self, book, sheet_name: str, row: int, col: int, formula_text: str):
        """对外部工作簿里没有缓存值的公式做递归本地求值（带循环检测）。"""
        marker = f"{book.path.name}!{sheet_name}!{get_column_letter(col)}{row}"
        key = (str(book.path), sheet_name.casefold(), row, col)
        if key in self._stack:
            raise UnsupportedFormula(f"跨工作簿循环引用：{marker} 在求值链上重复出现")
        if len(self._stack) >= _MAX_EXTERNAL_DEPTH:
            raise UnsupportedFormula(f"跨工作簿引用链过深（超过 {_MAX_EXTERNAL_DEPTH} 层）：{marker}")
        self._stack.add(key)
        try:
            try:
                ast = parse_formula(formula_text)
            except FormulaSyntaxError as exc:
                raise UnsupportedFormula(f"外部公式 {marker} 无法解析：{exc}") from exc
            nested = Evaluator(
                book.values,
                sheet_name,
                load_external=self.load_external,
                book_path=book.path,
                current_cell=(row, col),
                _stack=self._stack,
                _book=book,
                _formulas=book.formulas,
            )
            return nested.evaluate(ast)
        finally:
            self._stack.discard(key)

    def _eval_unary(self, node: Unary):
        value = self._eval(node.operand)
        if isinstance(value, ExcelError):
            return value
        if isinstance(value, list):
            # 区域/数组的一元运算：逐元素展开（-A1:A9、A1:A9% 等）
            return _map_array(lambda item: _unary_scalar(node.op, item), value)
        return _unary_scalar(node.op, value)

    def _eval_binary(self, node: Binary):
        left, right = self._eval(node.left), self._eval(node.right)
        for value in (left, right):
            if isinstance(value, ExcelError):
                return value
        if isinstance(left, list) or isinstance(right, list):
            # 区域/数组参与运算按 Excel 数组语义逐元素展开。旧实现把区域塌缩
            # 成首元素再比较/计算，会给出静默算错的结果（例如 (D2:D9="华北")
            # 只比较了 D2），因此一度改为直接拒绝；现在按广播规则正确展开。
            return _array_binary(node.op, left, right)
        return _binary_scalar(node.op, left, right)

    # -------------------------------------------------------------- 函数实现
    def _eval_func(self, node: FuncCall):
        name = node.name
        handler = _FUNCTIONS.get(name)
        if handler is None:
            raise UnsupportedFormula(f"本地验证暂不支持函数 {name}")
        # IF / IFERROR 需要延迟求值分支，单独处理
        if name == "IF":
            args = [self._eval(arg) for arg in node.args]
            if any(isinstance(arg, list) for arg in args):
                return _if_array(args)
            condition = args[0]
            if isinstance(condition, ExcelError):
                return condition
            if _to_bool(condition):
                return _single(args[1])
            if len(args) > 2:
                return _single(args[2])
            return False
        if name in {"IFERROR", "IFNA"}:
            try:
                value = _single(self._eval(node.args[0]))
            except UnsupportedFormula:
                raise
            except Exception:  # noqa: BLE001 - 求值异常等价于 Excel 报错
                return _single(self._eval(node.args[1]))
            if isinstance(value, ExcelError):
                return _single(self._eval(node.args[1]))
            return value
        if name == "IFS":
            # 条件-结果成对，命中第一个为真的条件即返回，未命中的分支不求值
            for index in range(0, len(node.args) - 1, 2):
                condition = self._eval(node.args[index])
                if isinstance(condition, ExcelError):
                    return condition
                if _to_bool(_single(condition)):
                    return _single(self._eval(node.args[index + 1]))
            return NA_ERR
        if name == "CHOOSE":
            # 只对选中的分支求值：CHOOSE(2, 坏公式, 好结果) 在 Excel 中不会报错
            index_value = self._eval(node.args[0])
            if isinstance(index_value, ExcelError):
                return index_value
            picks = _single(index_value)
            number = _to_number(picks)
            if isinstance(number, ExcelError):
                return number
            choice = int(number)
            if choice < 1 or choice > len(node.args) - 1:
                return VALUE_ERR
            return _single(self._eval(node.args[choice]))
        if name == "SWITCH":
            # 命中分支才求值：SWITCH(x, 1, "a", 2, 1/0) 不该因未命中的 1/0 而报错
            target = _single(self._eval(node.args[0]))
            if isinstance(target, ExcelError):
                return target
            index = 1
            while index + 1 < len(node.args):
                candidate = self._eval(node.args[index])
                if isinstance(candidate, ExcelError):
                    return candidate
                if _compare("=", _single(candidate), target) is True:
                    return _single(self._eval(node.args[index + 1]))
                index += 2
            if index < len(node.args):  # 最后的落单参数是默认值
                return _single(self._eval(node.args[index]))
            return NA_ERR
        if name in {"ROW", "COLUMN"}:
            # 行列信息只存在于引用坐标里，已求值的参数会丢失，须直接看 AST
            return self._eval_position_func(name, node)
        args = [self._eval(arg) for arg in node.args]
        for arg in args:
            if isinstance(arg, ExcelError) and name not in {"ISERROR", "ISERR", "ISNA"}:
                return arg
        return handler(args)

    def _eval_position_func(self, name: str, node: FuncCall):
        """ROW()/COLUMN()：无参返回公式所在单元格的行/列号；带引用参数返回该引用的
        首行/首列，区域形式按数组返回整段行/列号（标量上下文经隐式交叉取首值，
        与 Excel 行为一致）。整列引用拿不到有效行范围，保守报未验证。"""
        if not node.args:
            if self.current_cell is None:
                raise UnsupportedFormula(f"无法确定公式所在单元格，暂不验证 {name}()")
            return float(self.current_cell[0] if name == "ROW" else self.current_cell[1])
        target = node.args[0]
        if not isinstance(target, Ref):
            raise UnsupportedFormula(f"{name} 的参数必须是单元格或区域引用")
        if name == "ROW":
            if target.whole_column:
                raise UnsupportedFormula("本地验证暂不支持对整列引用使用 ROW")
            if target.row2 > target.row1:
                return [[float(row)] for row in range(target.row1, target.row2 + 1)]
            return float(target.row1)
        if target.col2 > target.col1:
            return [[float(col) for col in range(target.col1, target.col2 + 1)]]
        return float(target.col1)


# ------------------------------------------------------------------ 值处理工具
def _lookup_sheet(sheets, name: str):
    """Excel 的工作表名不区分大小写：先精确命中，再大小写不敏感兜底。"""
    if not sheets:
        return None
    ws = sheets.get(name)
    if ws is not None:
        return ws
    lowered = name.casefold()
    for key, ws in sheets.items():
        if key.casefold() == lowered:
            return ws
    return None


def _coerce(value: object):
    if isinstance(value, str) and value.startswith("="):
        # data_only 视图里仍是公式，说明该单元格没有缓存值
        raise UnsupportedFormula("引用的单元格是尚未计算的公式，无法本地验证")
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        # 包装成 DateValue：算术与比较按序列号进行；纯取值场景（如 VLOOKUP
        # 按行命中日期列）原样返回，以 ISO 文本展示。
        return DateValue(value)
    return value


def _flatten(value) -> list:
    if isinstance(value, list):
        out: list = []
        for item in value:
            out.extend(_flatten(item))
        return out
    return [value]


def _single(value):
    if isinstance(value, list):
        flat = _flatten(value)
        return flat[0] if flat else None
    return value


def _if_array(args: list):
    """IF 的数组形式，用于 IF({1,0},E2:E9,D2:D9) 这类“反向查找”构造。

    按 Excel 的数组广播规则：先取所有参与值形状的最大值当作结果形状，
    长度不足的维度自动伸展（1x2 的条件配 8x1 的区域得到 8x2）。
    """
    yes = args[1]
    no = args[2] if len(args) > 2 else False
    rows = max(_layer_count(value) for value in args) or 1
    cols = max(_layer_width(value) for value in args) or 1
    return [
        [_if_at(args[0], yes, no, row, col) for col in range(cols)]
        for row in range(rows)
    ]


def _layer_count(value) -> int:
    if not isinstance(value, list):
        return 1
    return len(value) if value and isinstance(value[0], list) else 1


def _layer_width(value) -> int:
    if not isinstance(value, list):
        return 1
    if not value:
        return 0
    return len(value[0]) if isinstance(value[0], list) else len(value)


def _at(value, row: int, col: int = 0):
    """取广播后 (row, col) 处的值：标量不变，一维退化为列向量，越界时钳到首行/首列。"""
    if not isinstance(value, list):
        return value
    if not value:
        return None
    line = value[row] if row < len(value) else value[0]
    if not isinstance(line, list):
        return line
    if not line:
        return None
    return line[col] if col < len(line) else line[0]


def _if_at(condition, yes, no, row: int, col: int):
    chosen = _at(condition, row, col)
    if isinstance(chosen, ExcelError):
        return chosen
    return _at(yes, row, col) if _to_bool(chosen) else _at(no, row, col)


def _broadcastable(size_a: int, size_b: int) -> bool:
    """数组广播兼容性：两侧尺寸相等，或任一为 1（沿该维伸展）。"""
    return size_a == size_b or size_a == 1 or size_b == 1


def _array_binary(op: str, left, right):
    """区域/数组参与二元运算的数组语义：先定广播形状，再逐元素按标量规则计算。

    形状规则与 IF 的数组形式一致（共用 _layer_count/_layer_width/_at）。维度不
    兼容（两侧都大于 1 且不相等，如 8 行比 6 行）时返回 #N/A——Excel 对尺寸不
    匹配的数组运算用 #N/A 填充缺失元素；这里宁可给出明确错误，也不钳制出
    "看起来正常"的静默算错结果。
    """
    left_rows, left_cols = _layer_count(left), _layer_width(left)
    right_rows, right_cols = _layer_count(right), _layer_width(right)
    if not _broadcastable(left_rows, right_rows) or not _broadcastable(left_cols, right_cols):
        return NA_ERR
    rows = max(left_rows, right_rows)
    cols = max(left_cols, right_cols)
    return [
        [_binary_scalar(op, _at(left, row, col), _at(right, row, col)) for col in range(cols)]
        for row in range(rows)
    ]


def _map_array(func, value):
    """对整个数组逐元素应用标量函数（一元运算复用），形状规整为二维。"""
    return [
        [func(_at(value, row, col)) for col in range(_layer_width(value))]
        for row in range(_layer_count(value))
    ]


_ARRAY_NATIVE_FUNCS = {"SUMPRODUCT"}


def needs_array_formula(node) -> bool:
    """AST 判定：公式是否需要以数组公式（CSE）形态写入文件。

    区域参与一元/二元运算（含比较）的公式本地按数组语义求值；但若以普通公式
    文本写入 .xlsx，Excel/WPS 打开时会按传统"隐式交叉"只取与公式同行/列交叉
    的单值，与本地结果不一致。以数组公式形态写入即可在各版本还原数组语义。

    SUMPRODUCT 除外：其参数天然按数组处理，各版本都会正确求值，不需要 CSE。
    """
    if isinstance(node, FuncCall):
        name = node.name.upper()
        if name in _ARRAY_NATIVE_FUNCS:
            return False
        if name in {"ROW", "COLUMN"} and any(
            isinstance(arg, Ref) and arg.is_range for arg in node.args
        ):
            # ROW(A1:A5) 这类区域形式返回数组，裸公式会被旧版 Excel/WPS
            # 隐式交叉成单值，需要 CSE 才能还原数组语义
            return True
        if name == "IF" and any(_could_be_array(arg) for arg in node.args):
            return True
        return any(needs_array_formula(arg) for arg in node.args)
    if isinstance(node, Unary):
        if _could_be_array(node.operand):
            return True
        return needs_array_formula(node.operand)
    if isinstance(node, Binary):
        if _could_be_array(node.left) or _could_be_array(node.right):
            return True
        return needs_array_formula(node.left) or needs_array_formula(node.right)
    return False


def _could_be_array(node) -> bool:
    """子表达式求值是否可能得到数组：区域、数组字面量，或数组形式的 IF。"""
    if isinstance(node, Ref):
        return node.is_range
    if isinstance(node, ArrayLiteral):
        return True
    if isinstance(node, Unary):
        return _could_be_array(node.operand)
    if isinstance(node, Binary):
        return _could_be_array(node.left) or _could_be_array(node.right)
    if isinstance(node, FuncCall):
        name = node.name.upper()
        if name in {"ROW", "COLUMN"} and any(
            isinstance(arg, Ref) and arg.is_range for arg in node.args
        ):
            return True
        return name == "IF" and any(_could_be_array(arg) for arg in node.args)
    return False


def _numbers(values) -> list[float]:
    out: list[float] = []
    for item in _flatten(values):
        if isinstance(item, DateValue):
            out.append(_to_serial(item.raw))  # 日期按序列号参与汇总（与 Excel 一致）
            continue
        if isinstance(item, bool) or item is None or isinstance(item, ExcelError):
            continue
        if isinstance(item, (int, float)):
            out.append(float(item))
    return out


def _aggregate_number(value):
    """条件汇总的取值：日期 → 序列号，数字原样，其余返回 None（按忽略处理）。"""
    if isinstance(value, DateValue):
        return _to_serial(value.raw)
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value)


def _to_number(value):
    if isinstance(value, DateValue):
        return _to_serial(value.raw)  # 日期在 Excel 里就是数字：序列号
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, ExcelError):
        return value
    try:
        return float(str(value).strip())
    except ValueError:
        return VALUE_ERR


def _to_text(value) -> str:
    if isinstance(value, DateValue):
        # 单元格数字格式本地不可知，统一用 ISO 文本：既可直接阅读，也让
        # ">="&DATE(...) 这类拼接出的条件文本能被 _parse_date_text 还原成日期
        return _date_text(value.raw)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _to_bool(value) -> bool:
    if isinstance(value, DateValue):
        return _to_serial(value.raw) != 0
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return str(value).strip().upper() == "TRUE"


def _typed(value):
    """类型判断类函数的取值：日期在 Excel 里属于数字（序列号）。"""
    if isinstance(value, DateValue):
        return _to_serial(value.raw)
    return value


def _compare(op: str, left, right):
    if isinstance(left, DateValue) or isinstance(right, DateValue):
        a, b = _date_side(left), _date_side(right)
        if a is None or b is None:
            raise UnsupportedFormula(DATE_HINT)
        return _order(op, a, b)
    if isinstance(left, str) or isinstance(right, str):
        a, b = _to_text(left).upper(), _to_text(right).upper()
    else:
        a, b = _to_number(left), _to_number(right)
        if isinstance(a, ExcelError) or isinstance(b, ExcelError):
            return VALUE_ERR
    return _order(op, a, b)


def _order(op: str, a, b):
    if op == "=":
        return a == b
    if op == "<>":
        return a != b
    if op == "<":
        return a < b
    if op == ">":
        return a > b
    if op == "<=":
        return a <= b
    return a >= b


def _date_side(value):
    """日期参与比较时把两侧折算成序列号；无法折算返回 None（上层标"未验证"）。"""
    if isinstance(value, DateValue):
        return _to_serial(value.raw)
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    if value is None:
        return 0.0
    if isinstance(value, str):
        parsed = _parse_date_text(value)
        return None if parsed is None else _to_serial(parsed)
    return None


def _date_arithmetic(op: str, left, right):
    """日期参与 +- 的 Excel 语义：日期±天数仍是日期，日期-日期是天数差。

    其余组合（日期×数字等）返回 None，交给通用数值路径按序列号计算。
    """
    left_date = isinstance(left, DateValue)
    right_date = isinstance(right, DateValue)
    try:
        if left_date and isinstance(right, (int, float)) and not isinstance(right, bool):
            delta = _dt.timedelta(days=float(right))
            return DateValue(left.raw + delta if op == "+" else left.raw - delta)
        if right_date and isinstance(left, (int, float)) and not isinstance(left, bool):
            if op == "+":
                return DateValue(right.raw + _dt.timedelta(days=float(left)))
            return _to_serial(left) - _to_serial(right.raw)
        if left_date and right_date:
            a, b = _to_serial(left.raw), _to_serial(right.raw)
            return a + b if op == "+" else a - b
    except OverflowError:
        return NUM_ERR
    return None


def _unary_scalar(op: str, value):
    """一元运算的标量规则；数组路径对每个元素复用同一语义。"""
    if isinstance(value, ExcelError):
        return value
    if op == "%":
        number = _to_number(value)
        return number if isinstance(number, ExcelError) else number / 100
    number = _to_number(value)
    if isinstance(number, ExcelError):
        return number
    return -number if op == "-" else number


def _binary_scalar(op: str, left, right):
    """二元运算的标量规则；数组路径按元素对复用这里，保证语义一致。"""
    for value in (left, right):
        if isinstance(value, ExcelError):
            return value
    if op == "&":
        return _to_text(left) + _to_text(right)
    if op in {"=", "<>", "<", ">", "<=", ">="}:
        return _compare(op, left, right)
    if op in {"+", "-"}:
        dated = _date_arithmetic(op, left, right)
        if dated is not None:
            return dated
    a, b = _to_number(left), _to_number(right)
    if isinstance(a, ExcelError):
        return a
    if isinstance(b, ExcelError):
        return b
    if op == "+":
        return a + b
    if op == "-":
        return a - b
    if op == "*":
        return a * b
    if op == "/":
        return DIV0 if b == 0 else a / b
    if op == "^":
        try:
            return math.pow(a, b)
        except (ValueError, OverflowError):
            return NUM_ERR
    raise UnsupportedFormula(f"不支持的运算符 {op}")


def _date_text(raw) -> str:
    """日期值的文本形式：零点整的 datetime 等价于 date，其余带上时间部分。"""
    if isinstance(raw, _dt.datetime):
        if raw.time() == _dt.time(0, 0):
            return raw.date().isoformat()
        return f"{raw.date().isoformat()} {raw.time().isoformat()}"
    return raw.isoformat()


_CRITERIA_RE = re.compile(r"^\s*(<=|>=|<>|=|<|>)?\s*(.*)$", re.DOTALL)


def _match_criteria(value, criteria) -> bool:
    """实现 COUNTIF/SUMIF 的条件匹配，如 ">80"、"Math"、"<>"。"""
    if isinstance(criteria, (int, float)) and not isinstance(criteria, bool):
        return _compare("=", value, criteria) is True
    match = _CRITERIA_RE.match(_to_text(criteria))
    op = match.group(1) or "="
    operand_text = match.group(2)
    try:
        operand: object = float(operand_text)
    except ValueError:
        operand = operand_text
    if isinstance(operand, str) and ("*" in operand or "?" in operand):
        raise UnsupportedFormula("本地验证不支持通配符条件")
    if isinstance(operand, str) and isinstance(value, DateValue):
        parsed = _parse_date_text(operand)
        if parsed is not None:
            operand = _to_serial(parsed)  # ">2025-01-15" 这类日期文本条件
    if value is None and op in {"=", "<>"}:
        return (operand_text == "") if op == "=" else (operand_text != "")
    result = _compare(op, value, operand)
    return result is True


def _sum_if(args, *, mode: str):
    """SUMIF / AVERAGEIF / COUNTIF 的统一实现。"""
    if mode == "COUNTIF":
        rng, criteria = args[0], _single(args[1])
        picked = [v for v in _flatten(rng) if _match_criteria(v, criteria)]
        return float(len(picked))
    rng, criteria = args[0], _single(args[1])
    target = args[2] if len(args) > 2 else rng
    source = _flatten(rng)
    values = _flatten(target)
    if len(values) < len(source):
        raise UnsupportedFormula("条件区域与求和区域大小不一致")
    numbers: list[float] = []
    for item, value in zip(source, values):
        if not _match_criteria(item, criteria):
            continue
        number = _aggregate_number(value)
        if number is not None:
            numbers.append(number)
    if mode == "SUMIF":
        return math.fsum(numbers)
    return math.fsum(numbers) / len(numbers) if numbers else DIV0


def _multi_criteria(args, *, mode: str):
    """COUNTIFS / SUMIFS / AVERAGEIFS。"""
    if mode == "COUNTIFS":
        pairs = list(zip(args[0::2], args[1::2]))
        sum_range = None
    else:
        sum_range = _flatten(args[0])
        pairs = list(zip(args[1::2], args[2::2]))
    if not pairs:
        raise UnsupportedFormula("条件参数不成对")
    length = len(_flatten(pairs[0][0]))
    hits = []
    for index in range(length):
        ok = True
        for rng, criteria in pairs:
            flat = _flatten(rng)
            if len(flat) != length:
                raise UnsupportedFormula("多条件区域大小不一致")
            if not _match_criteria(flat[index], _single(criteria)):
                ok = False
                break
        if ok:
            hits.append(index)
    if mode == "COUNTIFS":
        return float(len(hits))
    numbers: list[float] = []
    for index in hits:
        if index < len(sum_range):
            number = _aggregate_number(sum_range[index])
            if number is not None:
                numbers.append(number)
    if mode == "SUMIFS":
        return math.fsum(numbers)
    return math.fsum(numbers) / len(numbers) if numbers else DIV0


def _safe_div(numbers: list[float], func):
    return func(numbers) if numbers else DIV0


def _round(args, mode: str):
    value = _to_number(_single(args[0]))
    if isinstance(value, ExcelError):
        return value
    digits = int(_to_number(_single(args[1]))) if len(args) > 1 else 0
    factor = 10 ** digits
    if mode == "ROUND":
        return math.floor(abs(value) * factor + 0.5) / factor * (1 if value >= 0 else -1)
    if mode == "UP":
        return math.ceil(abs(value) * factor) / factor * (1 if value >= 0 else -1)
    return math.floor(abs(value) * factor) / factor * (1 if value >= 0 else -1)


def _mod(args):
    """MOD：结果符号跟随除数（Excel 语义）。math.fmod 跟随被除数，负被除数时会给出相反符号。"""
    divisor = _to_number(_single(args[1]))
    if isinstance(divisor, ExcelError):
        return divisor
    dividend = _to_number(_single(args[0]))
    if isinstance(dividend, ExcelError):
        return dividend
    if divisor == 0:
        return DIV0
    return dividend % divisor


def _trunc(args):
    value = _to_number(_single(args[0]))
    if isinstance(value, ExcelError):
        return value
    digits = _require_int(args[1]) if len(args) > 1 else 0
    if isinstance(digits, ExcelError):
        return digits
    factor = 10 ** digits
    return math.trunc(value * factor) / factor


def _rank(args):
    value = _to_number(_single(args[0]))
    numbers = _numbers(args[1])
    order = int(_to_number(_single(args[2]))) if len(args) > 2 else 0
    if not numbers or isinstance(value, ExcelError):
        return NA_ERR
    ordered = sorted(numbers) if order else sorted(numbers, reverse=True)
    return float(ordered.index(value) + 1) if value in ordered else NA_ERR


def _vlookup(args):
    key = _single(args[0])
    table = args[1]
    if not isinstance(table, list):
        raise UnsupportedFormula("VLOOKUP 第二参数必须是区域或数组")
    col_index = int(_to_number(_single(args[2])))
    exact = len(args) > 3 and not _to_bool(_single(args[3]))
    if not exact:
        raise UnsupportedFormula("本地验证仅支持精确匹配的 VLOOKUP（第 4 参数为 FALSE）")
    for row in table:
        cells = row if isinstance(row, list) else [row]
        if cells and _compare("=", cells[0], key) is True:
            if col_index > len(cells):
                return ExcelError("#REF!")
            return cells[col_index - 1]
    return NA_ERR


def _sumproduct(args):
    """SUMPRODUCT：各参数逐位相乘再求和。

    参数长度不一致时 Excel 报 #VALUE!，这里必须显式检查——沿用 zip 会把长数组
    静默截断，算出一个偏小的"看起来正常"的数字。
    """
    columns = [[float(value) for value in _numbers(arg)] for arg in args]
    if len({len(column) for column in columns}) > 1:
        return VALUE_ERR
    return math.fsum(math.prod(pair) for pair in zip(*columns))


def _text_func(args, func):
    return func(_to_text(_single(args[0])))


# ------------------------------------------------------------------ 查找族
def _index(args):
    """INDEX(区域, 行号, [列号])：标量取值；单行区域只给一个下标时按列号解释。"""
    table = args[0]
    if not isinstance(table, list):
        raise UnsupportedFormula("INDEX 第一参数必须是区域或数组")
    grid = table if (table and isinstance(table[0], list)) else [table]
    height = len(grid)
    width = len(grid[0]) if grid else 0
    row_num = _require_int(args[1])
    if isinstance(row_num, ExcelError):
        return row_num
    if height == 1 and len(args) <= 2 and width > 1:
        # 单行区域 + 单下标：按列号解释，行号固定为 1
        col_num = row_num
        row_num = 1
    else:
        if len(args) > 2:
            col_num = _require_int(args[2])
            if isinstance(col_num, ExcelError):
                return col_num
        else:
            col_num = 1
    if row_num < 1 or col_num < 1:
        raise UnsupportedFormula("本地验证不支持 INDEX 的 0 下标（整行/整列返回）")
    if row_num > height or col_num > width:
        return ExcelError("#REF!")
    return grid[row_num - 1][col_num - 1]


def _match(args):
    """MATCH(查找值, 单行或单列区域, [匹配类型])：0 精确，1 升序近似（默认），-1 降序近似。"""
    key = _single(args[0])
    flat = _flatten(args[1])
    match_type = _require_int(args[2]) if len(args) > 2 else 1
    if isinstance(match_type, ExcelError):
        return match_type
    if match_type == 0:
        for index, value in enumerate(flat):
            if _compare("=", value, key) is True:
                return float(index + 1)
        return NA_ERR
    best = None
    for index, value in enumerate(flat):
        hit = _compare("<=", value, key) if match_type > 0 else _compare(">=", value, key)
        if hit is True:
            best = index + 1
    return float(best) if best else NA_ERR


def _hlookup(args):
    key = _single(args[0])
    table = args[1]
    if not isinstance(table, list):
        raise UnsupportedFormula("HLOOKUP 第二参数必须是区域或数组")
    row_index = _require_int(args[2])
    if isinstance(row_index, ExcelError):
        return row_index
    exact = len(args) > 3 and not _to_bool(_single(args[3]))
    if not exact:
        raise UnsupportedFormula("本地验证仅支持精确匹配的 HLOOKUP（第 4 参数为 FALSE）")
    grid = table if (table and isinstance(table[0], list)) else [table]
    header = grid[0] if grid else []
    for col, cell in enumerate(header):
        if _compare("=", cell, key) is True:
            if row_index > len(grid):
                return ExcelError("#REF!")
            row = grid[row_index - 1]
            return row[col] if col < len(row) else ExcelError("#REF!")
    return NA_ERR


def _xlookup(args):
    """XLOOKUP：仅支持默认的精确匹配、从头搜索。"""
    key = _single(args[0])
    lookup_values = _flatten(args[1])
    return_values = _flatten(args[2])
    if len(lookup_values) != len(return_values):
        return VALUE_ERR
    if len(args) > 4:
        mode = _to_number(_single(args[4]))
        if mode != 0:
            raise UnsupportedFormula("本地验证仅支持 XLOOKUP 精确匹配（第 5 参数为 0/省略）")
    for index, candidate in enumerate(lookup_values):
        if _compare("=", candidate, key) is True:
            return return_values[index]
    return _single(args[3]) if len(args) > 3 else NA_ERR


def _lookup(args):
    """LOOKUP 向量形式：在升序的查找列中做近似匹配，返回结果列同位置的值。"""
    key = _single(args[0])
    lookup_values = _flatten(args[1])
    result_values = _flatten(args[2]) if len(args) > 2 else lookup_values
    if len(result_values) < len(lookup_values):
        return VALUE_ERR
    best = None
    for index, value in enumerate(lookup_values):
        if _compare("<=", value, key) is True:
            best = index
    return result_values[best] if best is not None else NA_ERR


def _require_int(node_value):
    """把可能是错误值/文本的标量参数转成 int，错误值原样返回。"""
    number = _to_number(_single(node_value))
    if isinstance(number, ExcelError):
        return number
    return int(number)


# ------------------------------------------------------------------ 统计扩展
def _mode(args):
    counts: dict = {}
    ordered: list[float] = []
    for number in _numbers(args):
        if number not in counts:
            ordered.append(number)
        counts[number] = counts.get(number, 0) + 1
    if not counts:
        return NA_ERR
    best = max(counts.values())
    if best == 1:
        return NA_ERR  # 每个数都只出现一次，Excel 返回 #N/A
    return next(number for number in ordered if counts[number] == best)


def _conditional_extreme(args, mode: str):
    """MAXIFS / MINIFS：首参数为取值区域，其后是成对的条件区域与条件。"""
    target = _flatten(args[0])
    pairs = list(zip(args[1::2], args[2::2]))
    if not pairs:
        raise UnsupportedFormula("条件参数不成对")
    length = len(target)
    picked: list[float] = []
    for index in range(length):
        ok = True
        for rng, criteria in pairs:
            flat = _flatten(rng)
            if len(flat) != length:
                raise UnsupportedFormula("条件区域与取值区域大小不一致")
            if not _match_criteria(flat[index], _single(criteria)):
                ok = False
                break
        if ok:
            number = _aggregate_number(target[index])
            if number is not None:
                picked.append(number)
    if not picked:
        return 0.0  # 没有任何满足条件的数值时 Excel 返回 0
    return max(picked) if mode == "MAX" else min(picked)


# ------------------------------------------------------------------ 数学扩展
def _exp(args):
    try:
        return math.exp(_to_number(_single(args[0])))
    except OverflowError:
        return NUM_ERR


def _ln(args):
    value = _to_number(_single(args[0]))
    if isinstance(value, ExcelError):
        return value
    return math.log(value) if value > 0 else NUM_ERR


def _log(args):
    value = _to_number(_single(args[0]))
    if isinstance(value, ExcelError):
        return value
    base = _to_number(_single(args[1])) if len(args) > 1 else 10.0
    if isinstance(base, ExcelError):
        return base
    if value <= 0 or base <= 0 or base == 1:
        return NUM_ERR
    return math.log(value, base)


def _log10(args):
    value = _to_number(_single(args[0]))
    if isinstance(value, ExcelError):
        return value
    return math.log10(value) if value > 0 else NUM_ERR


def _ceiling_floor(args, mode: str):
    """CEILING / FLOOR（正显著性）：结果朝 +∞ / -∞ 方向凑到显著性的整数倍。"""
    value = _to_number(_single(args[0]))
    significance = _to_number(_single(args[1]))
    if isinstance(value, ExcelError):
        return value
    if isinstance(significance, ExcelError):
        return significance
    if significance == 0:
        return DIV0
    if significance < 0:
        raise UnsupportedFormula("本地验证仅支持正显著性参数的 CEILING/FLOOR")
    ratio = value / significance
    if abs(ratio - round(ratio)) < 1e-9:  # 消除 0.1+0.2 类浮点噪声
        return round(ratio) * significance
    return (math.ceil(ratio) if mode == "CEILING" else math.floor(ratio)) * significance


# ------------------------------------------------------------------ 文本扩展
def _find(args, *, wildcard: bool):
    """FIND（区分大小写）/ SEARCH（不区分大小写，支持 * ? 通配符）：返回 1 起始位置。"""
    needle = _to_text(_single(args[0]))
    haystack = _to_text(_single(args[1]))
    start = _require_int(args[2]) if len(args) > 2 else 1
    if isinstance(start, ExcelError):
        return start
    if start < 1 or start > len(haystack) + 1:
        return VALUE_ERR
    hay = haystack[start - 1 :]
    if wildcard:
        pattern = "".join(
            ".*" if ch == "*" else "." if ch == "?" else re.escape(ch) for ch in needle
        )
        found = re.search(pattern, hay, re.IGNORECASE)
        return float(start + found.start()) if found else VALUE_ERR
    position = hay.find(needle)
    return float(start + position) if position >= 0 else VALUE_ERR


def _substitute(args):
    text = _to_text(_single(args[0]))
    old = _to_text(_single(args[1]))
    new = _to_text(_single(args[2]))
    if not old:
        return text
    if len(args) > 3:
        instance = _require_int(args[3])
        if isinstance(instance, ExcelError):
            return instance
        if instance < 1:
            return VALUE_ERR
        count = 0
        for index in range(len(text)):
            if text.startswith(old, index):
                count += 1
                if count == instance:
                    return text[:index] + new + text[index + len(old) :]
        return text  # 出现次数不足时原样返回
    return text.replace(old, new)


def _replace(args):
    text = _to_text(_single(args[0]))
    start = _require_int(args[1])
    length = _require_int(args[2])
    new = _to_text(_single(args[3]))
    if isinstance(start, ExcelError) or isinstance(length, ExcelError):
        return VALUE_ERR
    if start < 1 or length < 0:
        return VALUE_ERR
    return text[: start - 1] + new + text[start - 1 + length :]


def _rept(args):
    text = _to_text(_single(args[0]))
    count = _require_int(args[1])
    if isinstance(count, ExcelError):
        return count
    if count < 0 or len(text) * count > 32767:  # Excel 单元格上限
        return VALUE_ERR
    return text * count


def _char(args):
    code = _require_int(args[0])
    if isinstance(code, ExcelError):
        return code
    return chr(code) if 1 <= code <= 255 else VALUE_ERR


def _code(args):
    text = _to_text(_single(args[0]))
    return float(ord(text[0])) if text else VALUE_ERR


def _value(args):
    """VALUE：把文本转成数值，容忍千分位逗号、货币符号与百分号。"""
    raw = _single(args[0])
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return float(raw)
    text = _to_text(raw).strip()
    percent = text.endswith("%")
    if percent:
        text = text[:-1]
    cleaned = text.replace(",", "").replace("￥", "").replace("$", "").strip()
    try:
        number = float(cleaned)
    except ValueError:
        return VALUE_ERR
    return number / 100 if percent else number


# ------------------------------------------------------------------ 逻辑扩展
def _parity(args, *, want_even: bool):
    number = _to_number(_single(args[0]))
    if isinstance(number, ExcelError):
        return number
    remainder = abs(int(number)) % 2
    return remainder == 0 if want_even else remainder == 1


# ------------------------------------------------------------------ 日期时间
# 全部按 1900 日期系统（1899-12-30 = 序列号 0）实现，与 Excel 的日期序列号一致。
def _require_serial(node_value):
    """日期函数的参数 → 序列号：日期对象折算，数字原样，日期文本尝试解析。"""
    value = _single(node_value)
    if isinstance(value, str):
        parsed = _parse_date_text(value)
        if parsed is not None:
            return _to_serial(parsed)
    return _to_number(value)


def _date_func(args):
    """DATE：年月日 → 日期。月份/日期溢出按 Excel 规则进位，0~1899 的年自动加 1900。"""
    year = _require_int(args[0])
    month = _require_int(args[1])
    day = _require_int(args[2])
    for number in (year, month, day):
        if isinstance(number, ExcelError):
            return number
    if year < 0:
        return NUM_ERR
    if year < 1900:
        year += 1900
    if year > 9999:
        return NUM_ERR
    year += (month - 1) // 12
    month = (month - 1) % 12 + 1
    try:
        return DateValue(_dt.date(year, month, 1) + _dt.timedelta(days=day - 1))
    except (ValueError, OverflowError):
        return NUM_ERR


def _date_part(args, mode: str):
    serial = _require_serial(args[0])
    if isinstance(serial, ExcelError):
        return serial
    raw = _from_serial(serial).raw
    return float({"Y": raw.year, "M": raw.month, "D": raw.day}[mode])


def _time_part(args, mode: str):
    serial = _require_serial(args[0])
    if isinstance(serial, ExcelError):
        return serial
    total = round((serial - math.floor(serial)) * 86400)
    hours, remaining = divmod(total, 3600)
    minutes, seconds = divmod(remaining, 60)
    return float({"H": hours, "M": minutes, "S": seconds}[mode])


def _datevalue(args):
    text = _single(args[0])
    if isinstance(text, str):
        parsed = _parse_date_text(text)
        if parsed is not None:
            return DateValue(parsed)
    return VALUE_ERR


def _add_months(day: _dt.date, months: int) -> _dt.date:
    total = day.year * 12 + (day.month - 1) + months
    year, month = divmod(total, 12)
    month += 1
    return _dt.date(year, month, min(day.day, calendar.monthrange(year, month)[1]))


def _full_months(start: _dt.date, end: _dt.date) -> int:
    months = (end.year - start.year) * 12 + end.month - start.month
    if end.day < start.day:
        months -= 1
    return max(months, 0)


def _full_years(start: _dt.date, end: _dt.date) -> int:
    years = end.year - start.year
    if (end.month, end.day) < (start.month, start.day):
        years -= 1
    return max(years, 0)


def _shift_years(day: _dt.date, years: int) -> _dt.date:
    try:
        return day.replace(year=day.year + years)
    except ValueError:  # 2 月 29 日在目标年不存在
        return _dt.date(day.year + years, 2, 28)


def _edate(args):
    """EDATE：按月份偏移得到同一天（目标月天数不足时取月末）。"""
    serial = _require_serial(args[0])
    months = _require_int(args[1])
    for number in (serial, months):
        if isinstance(number, ExcelError):
            return number
    try:
        return DateValue(_add_months(_as_date(_from_serial(serial).raw), months))
    except (ValueError, OverflowError):
        return NUM_ERR


def _eomonth(args):
    """EOMONTH：月份偏移后取该月最后一天。"""
    serial = _require_serial(args[0])
    months = _require_int(args[1])
    for number in (serial, months):
        if isinstance(number, ExcelError):
            return number
    try:
        day = _as_date(_from_serial(serial).raw)
        shifted = _add_months(_dt.date(day.year, day.month, 1), months)
        last = calendar.monthrange(shifted.year, shifted.month)[1]
        return DateValue(_dt.date(shifted.year, shifted.month, last))
    except (ValueError, OverflowError):
        return NUM_ERR


def _days(args):
    end = _require_serial(args[0])
    start = _require_serial(args[1])
    for value in (end, start):
        if isinstance(value, ExcelError):
            return value
    return end - start


def _datedif(args):
    """DATEDIF：两日期的整年/整月/天数差，支持 Y/M/D/YM/YD/MD 单位。"""
    unit = _to_text(_single(args[2])).strip().upper()
    start = _require_serial(args[0])
    end = _require_serial(args[1])
    for value in (start, end):
        if isinstance(value, ExcelError):
            return value
    if end < start:
        return NUM_ERR
    start_day = _as_date(_from_serial(start).raw)
    end_day = _as_date(_from_serial(end).raw)
    if unit == "D":
        return float((end_day - start_day).days)
    if unit == "Y":
        return float(_full_years(start_day, end_day))
    if unit == "M":
        return float(_full_months(start_day, end_day))
    if unit == "YM":
        return float(_full_months(start_day, end_day) % 12)
    if unit == "YD":
        anchor = _shift_years(start_day, _full_years(start_day, end_day))
        return float((end_day - anchor).days)
    if unit == "MD":
        if end_day.day >= start_day.day:
            return float(end_day.day - start_day.day)
        borrowed = _add_months(_dt.date(end_day.year, end_day.month, 1), -1)
        return float(
            end_day.day + calendar.monthrange(borrowed.year, borrowed.month)[1] - start_day.day
        )
    return NUM_ERR


def _weekday(args):
    """WEEKDAY：1=周日..7=周六（默认）；2=周一..7=周日；3=周一=0；11~17 为周一起算的变体。"""
    serial = _require_serial(args[0])
    if isinstance(serial, ExcelError):
        return serial
    kind = _require_int(args[1]) if len(args) > 1 else 1
    if isinstance(kind, ExcelError):
        return kind
    weekday = _as_date(_from_serial(serial).raw).weekday()  # 周一=0 … 周日=6
    if kind == 1:
        return float((weekday + 1) % 7 + 1)
    if kind == 2:
        return float(weekday + 1)
    if kind == 3:
        return float(weekday)
    if 11 <= kind <= 17:
        return float((weekday - (kind - 11)) % 7 + 1)
    return NUM_ERR


def _weeknum(args):
    """WEEKNUM：1=周日为一周起点（默认），2=周一为一周起点。"""
    serial = _require_serial(args[0])
    if isinstance(serial, ExcelError):
        return serial
    kind = _require_int(args[1]) if len(args) > 1 else 1
    if isinstance(kind, ExcelError):
        return kind
    if kind not in (1, 2):
        return NUM_ERR
    day = _as_date(_from_serial(serial).raw)
    january_first = _dt.date(day.year, 1, 1)
    offset = (january_first.weekday() + 1) % 7 if kind == 1 else january_first.weekday()
    return float(((day - january_first).days + offset) // 7 + 1)


def _holiday_dates(node_value) -> set:
    """节假日参数（区域/数组）→ date 集合；无法识别的项忽略。"""
    days: set = set()
    for item in _flatten(node_value):
        if isinstance(item, DateValue):
            days.add(_as_date(item.raw))
        elif isinstance(item, (int, float)) and not isinstance(item, bool):
            days.add(_as_date(_from_serial(float(item)).raw))
        elif isinstance(item, str):
            parsed = _parse_date_text(item)
            if parsed is not None:
                days.add(parsed.date())
    return days


def _networkdays(args):
    start = _require_serial(args[0])
    end = _require_serial(args[1])
    for value in (start, end):
        if isinstance(value, ExcelError):
            return value
    holidays = _holiday_dates(args[2]) if len(args) > 2 else set()
    first, last = int(start), int(end)
    if first > last:
        first, last = last, first
    count = 0
    for serial in range(first, last + 1):
        day = _as_date(_from_serial(float(serial)).raw)
        if day.weekday() < 5 and day not in holidays:
            count += 1
    return float(count)


def _workday(args):
    serial = _require_serial(args[0])
    offset = _require_int(args[1])
    for value in (serial, offset):
        if isinstance(value, ExcelError):
            return value
    holidays = _holiday_dates(args[2]) if len(args) > 2 else set()
    day = _as_date(_from_serial(serial).raw)
    step = 1 if offset >= 0 else -1
    remaining = abs(offset)
    while remaining:
        day += _dt.timedelta(days=step)
        if day.weekday() < 5 and day not in holidays:
            remaining -= 1
    return DateValue(day)


_FUNCTIONS: dict[str, object] = {
    "SUM": lambda a: math.fsum(_numbers(a)),
    "PRODUCT": lambda a: math.prod(_numbers(a)) if _numbers(a) else 0.0,
    "AVERAGE": lambda a: _safe_div(_numbers(a), lambda n: math.fsum(n) / len(n)),
    "MEDIAN": lambda a: _safe_div(_numbers(a), statistics.median),
    "MAX": lambda a: max(_numbers(a)) if _numbers(a) else 0.0,
    "MIN": lambda a: min(_numbers(a)) if _numbers(a) else 0.0,
    "COUNT": lambda a: float(len(_numbers(a))),
    "COUNTA": lambda a: float(sum(1 for v in _flatten(a) if v is not None and v != "")),
    "COUNTBLANK": lambda a: float(sum(1 for v in _flatten(a) if v is None or v == "")),
    "STDEV": lambda a: _safe_div(_numbers(a), lambda n: statistics.stdev(n) if len(n) > 1 else DIV0),
    "STDEV.S": lambda a: _safe_div(_numbers(a), lambda n: statistics.stdev(n) if len(n) > 1 else DIV0),
    "STDEV.P": lambda a: _safe_div(_numbers(a), statistics.pstdev),
    "COUNTIF": lambda a: _sum_if(a, mode="COUNTIF"),
    "SUMIF": lambda a: _sum_if(a, mode="SUMIF"),
    "AVERAGEIF": lambda a: _sum_if(a, mode="AVERAGEIF"),
    "COUNTIFS": lambda a: _multi_criteria(a, mode="COUNTIFS"),
    "SUMIFS": lambda a: _multi_criteria(a, mode="SUMIFS"),
    "AVERAGEIFS": lambda a: _multi_criteria(a, mode="AVERAGEIFS"),
    "IF": lambda a: None,  # 由 Evaluator 特殊处理
    "IFERROR": lambda a: None,
    "IFNA": lambda a: None,
    "AND": lambda a: all(_to_bool(v) for v in _flatten(a) if v is not None),
    "OR": lambda a: any(_to_bool(v) for v in _flatten(a) if v is not None),
    "NOT": lambda a: not _to_bool(_single(a[0])),
    "ABS": lambda a: abs(_to_number(_single(a[0]))),
    "SQRT": lambda a: (
        math.sqrt(_to_number(_single(a[0]))) if _to_number(_single(a[0])) >= 0 else NUM_ERR
    ),
    "POWER": lambda a: math.pow(_to_number(_single(a[0])), _to_number(_single(a[1]))),
    "MOD": _mod,
    "INT": lambda a: float(math.floor(_to_number(_single(a[0])))),
    "TRUNC": _trunc,
    "SIGN": lambda a: float((_to_number(_single(a[0])) > 0) - (_to_number(_single(a[0])) < 0)),
    "ROUND": lambda a: _round(a, "ROUND"),
    "ROUNDUP": lambda a: _round(a, "UP"),
    "ROUNDDOWN": lambda a: _round(a, "DOWN"),
    "LARGE": lambda a: (
        sorted(_numbers(a[0]), reverse=True)[int(_to_number(_single(a[1]))) - 1]
        if 0 < int(_to_number(_single(a[1]))) <= len(_numbers(a[0])) else NUM_ERR
    ),
    "SMALL": lambda a: (
        sorted(_numbers(a[0]))[int(_to_number(_single(a[1]))) - 1]
        if 0 < int(_to_number(_single(a[1]))) <= len(_numbers(a[0])) else NUM_ERR
    ),
    "RANK": _rank,
    "RANK.EQ": _rank,
    "ROWS": lambda a: float(len(a[0]) if isinstance(a[0], list) else 1),
    "COLUMNS": lambda a: float(
        len(a[0][0]) if isinstance(a[0], list) and a[0] and isinstance(a[0][0], list) else 1
    ),
    "ROW": lambda a: None,  # 由 Evaluator 特殊处理（无参形式需要公式所在位置）
    "COLUMN": lambda a: None,
    "LEN": lambda a: float(len(_to_text(_single(a[0])))),
    "TRIM": lambda a: _text_func(a, lambda s: s.strip()),
    "UPPER": lambda a: _text_func(a, str.upper),
    "LOWER": lambda a: _text_func(a, str.lower),
    "PROPER": lambda a: _text_func(a, str.title),
    "LEFT": lambda a: _to_text(_single(a[0]))[: int(_to_number(_single(a[1]))) if len(a) > 1 else 1],
    "RIGHT": lambda a: _to_text(_single(a[0]))[
        -(int(_to_number(_single(a[1]))) if len(a) > 1 else 1):
    ],
    "MID": lambda a: _to_text(_single(a[0]))[
        int(_to_number(_single(a[1]))) - 1: int(_to_number(_single(a[1]))) - 1 + int(_to_number(_single(a[2])))
    ],
    "CONCAT": lambda a: "".join(_to_text(v) for v in _flatten(a)),
    "CONCATENATE": lambda a: "".join(_to_text(v) for v in _flatten(a)),
    "TEXTJOIN": lambda a: _to_text(_single(a[0])).join(
        _to_text(v) for v in _flatten(a[2:]) if not (_to_bool(_single(a[1])) and (v is None or v == ""))
    ),
    "ISBLANK": lambda a: _single(a[0]) is None,
    "ISNUMBER": lambda a: isinstance(_typed(_single(a[0])), (int, float))
    and not isinstance(_single(a[0]), bool),
    "ISTEXT": lambda a: isinstance(_typed(_single(a[0])), str),
    "ISERROR": lambda a: isinstance(_single(a[0]), ExcelError),
    "ISERR": lambda a: isinstance(_single(a[0]), ExcelError) and _single(a[0]) != NA_ERR,
    "ISNA": lambda a: _single(a[0]) == NA_ERR,
    "ISEVEN": lambda a: _parity(a, want_even=True),
    "ISODD": lambda a: _parity(a, want_even=False),
    "NA": lambda a: NA_ERR,
    "TRUE": lambda a: True,
    "FALSE": lambda a: False,
    "VLOOKUP": _vlookup,
    "HLOOKUP": _hlookup,
    "XLOOKUP": _xlookup,
    "LOOKUP": _lookup,
    "INDEX": _index,
    "MATCH": _match,
    "IFS": lambda a: None,  # 由 Evaluator 特殊处理（惰性求值）
    "CHOOSE": lambda a: None,  # 由 Evaluator 特殊处理（惰性求值）
    "SWITCH": lambda a: None,  # 由 Evaluator 特殊处理（惰性求值）
    "MODE": _mode,
    "MODE.SNGL": _mode,
    "MAXIFS": lambda a: _conditional_extreme(a, "MAX"),
    "MINIFS": lambda a: _conditional_extreme(a, "MIN"),
    "EXP": _exp,
    "LN": _ln,
    "LOG": _log,
    "LOG10": _log10,
    "CEILING": lambda a: _ceiling_floor(a, "CEILING"),
    "FLOOR": lambda a: _ceiling_floor(a, "FLOOR"),
    "FIND": lambda a: _find(a, wildcard=False),
    "SEARCH": lambda a: _find(a, wildcard=True),
    "SUBSTITUTE": _substitute,
    "REPLACE": _replace,
    "REPT": _rept,
    "CHAR": _char,
    "CODE": _code,
    "VALUE": _value,
    "XOR": lambda a: sum(1 for v in _flatten(a) if v is not None and _to_bool(v)) % 2 == 1,
    "VAR": lambda a: _safe_div(_numbers(a), lambda n: statistics.variance(n) if len(n) > 1 else DIV0),
    "VAR.S": lambda a: _safe_div(_numbers(a), lambda n: statistics.variance(n) if len(n) > 1 else DIV0),
    "VAR.P": lambda a: _safe_div(_numbers(a), statistics.pvariance),
    "STDEVP": lambda a: _safe_div(_numbers(a), statistics.pstdev),
    "SUMPRODUCT": _sumproduct,
    # 日期时间族
    "TODAY": lambda a: DateValue(_dt.date.today()),
    "NOW": lambda a: DateValue(_dt.datetime.now().replace(microsecond=0)),
    "DATE": _date_func,
    "YEAR": lambda a: _date_part(a, "Y"),
    "MONTH": lambda a: _date_part(a, "M"),
    "DAY": lambda a: _date_part(a, "D"),
    "HOUR": lambda a: _time_part(a, "H"),
    "MINUTE": lambda a: _time_part(a, "M"),
    "SECOND": lambda a: _time_part(a, "S"),
    "WEEKDAY": _weekday,
    "WEEKNUM": _weeknum,
    "EOMONTH": _eomonth,
    "EDATE": _edate,
    "DATEDIF": _datedif,
    "DAYS": _days,
    "DATEVALUE": _datevalue,
    "NETWORKDAYS": _networkdays,
    "WORKDAY": _workday,
}


def evaluate_formula(
    ast: object,
    sheets: dict[str, Worksheet],
    default_sheet: str,
    *,
    load_external=None,
    book_path: str | Path | None = None,
    formula_sheets: dict[str, Worksheet] | None = None,
    expand_uncached: bool = False,
    current_cell: tuple[int, int] | None = None,
):
    """对外接口：求值成功返回结果值，不支持时抛出 UnsupportedFormula。

    传入 load_external（ExternalBookLoader 实例或等价可调用对象）后，跨工作簿
    引用会尝试打开外部工作簿试算；book_path 是当前工作簿路径，用于解析公式里
    省略目录的外部引用（按同目录查找）。

    expand_uncached=True 时，主簿里引用到"没有缓存值的公式单元格"也会递归本地
    求值（展示场景用，让链式公式也能显示结果）；默认 False——校验场景把未计算
    的公式视为不可信输入，直接报"未验证"。

    current_cell=(行, 列) 是公式所在单元格位置，供 ROW()/COLUMN() 无参形式求值；
    调用方给不出位置时传 None，这两个函数按"未验证"处理。
    """
    return Evaluator(
        sheets,
        default_sheet,
        load_external=load_external,
        book_path=book_path,
        current_cell=current_cell,
        _formulas=formula_sheets,
        _expand_uncached=expand_uncached,
    ).evaluate(ast)


def format_value(value) -> str:
    if value is None:
        return "(空)"
    if isinstance(value, ExcelError):
        return value.code
    if isinstance(value, DateValue):
        raw = value.raw
        return raw.isoformat()[:10] if isinstance(raw, (_dt.datetime, _dt.date)) else raw.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6g}"
    return str(value)
