"""标准列表数据验证：规则解析、类型保留及网页选项元数据（纯本地）。"""
from __future__ import annotations

import datetime as dt
import difflib
import hashlib
import math
import re
import threading
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.parsers import expat

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .config import MAX_WRITE_CELLS
from .evaluator import DateValue, ExcelError, UnsupportedFormula, evaluate_formula, format_value
from .excel_reader import cell_formula_text
from .formula_parser import parse_formula, parse_range

MAX_OPTIONS = 1000
WRITE_LOCK = threading.RLock()
_X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
_COORD = r"\$?[A-Za-z]{1,3}\$?[1-9][0-9]{0,6}"
_RANGE = re.compile(r"(?<![A-Za-z0-9_$])" + _COORD + r"(?:\s*(?:[:：]|到|至)\s*" + _COORD + r")?(?![A-Za-z0-9_])")
_STATIC_REF = re.compile(r"(?:(?:'((?:[^']|'')*)'|([^!]+))!)?(" + _COORD + r"(?::" + _COORD + r")?)$")
_AREA = re.compile(_COORD + r"(?:\s*(?:[:：]|到|至)\s*" + _COORD + r")?")
# “数据来自于/来源于/取自于”这类书写：把多余的口语虚词一并吃掉；
# “引用自”单独列出，避免把“引用”后面的“自”留给表名。
_SOURCE = re.compile(r"(?:选项\s*)?(?:来源(?:为|是)?|来自|取自|源自|引用自|引用|指向)(?:于)?\s*[:：]?\s*")
# “选项是来自…”：来源引导词跟在值标记之后，单独识别
_SOURCE_LEAD = re.compile(r"(?:来源|来自|取自|源自|引用|指向)\s*(?:为|是|于|自)?\s*[:：]?\s*")
# “列表内容/内容为/值为”等值标记：长词在前，短词在后，避免从词中截断
_VALUES = re.compile(r"(?:选项|列表内容|可选值|内容|值)\s*(?:为|是|有|包括)?\s*[:：]?\s*")
# 无“!”的口语引用：“销售订单A2:A25”“销售订单的A2到A25”
_LOOSE_AREA = re.compile(r"[^!]+?" + _COORD + r"(?:\s*(?:[:：]|到|至)\s*" + _COORD + r")?")
# 句子尾注（“，以便B4查询”“便于下拉”）：只是说明，不属于来源本身
_TRAILING_NOTE = re.compile(r"[\s,，;；。]*(?:以便|便于|方便|用于|用来|以供|使其|从而|这样)[^、,，;；。]*$")
# 整列/整行写法：命中时给出更直接的提示
_COLUMN_ONLY = re.compile(r"(?:(?:'[^']*'|[^!]+)!)?\$?[A-Za-z]{1,3}(?:列|:\$?[A-Za-z]{1,3})$")
_ROW_ONLY = re.compile(r"(?:(?:'[^']*'|[^!]+)!)?[1-9][0-9]{0,6}:[1-9][0-9]{0,6}$")
_ACTION = r"(?:删除|取消|清除|移除|写入|写到|填入|计算|汇总|求和|设置|设为|添加|生成|加边框|框起来)"
_SECOND_ACTION = re.compile(
    r"(?:并且|并|同时|然后|另外|接着|再)\s*(?:(?:把|将|给|在)\s*)?(?:" + _COORD + r"\s*)?" + _ACTION
    + r"|(?:^|[，,；;。\n])\s*(?:" + _ACTION + r"\s*" + _COORD
    + r"|(?:把|将|给|在)\s*" + _COORD + r"\s*" + _ACTION + r")"
)


class DropdownClarification(ValueError):
    """缺少可确定的目标或来源，交由现有追问流程处理。"""


def file_version(path: Path) -> str:
    """对实际文件内容取摘要，避免同长度或同时间戳的修改漏检。"""
    with Path(path).open("rb") as stream:
        digest = hashlib.sha256()
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def require_version(path: Path, expected: str) -> None:
    if not expected or file_version(path) != expected:
        raise ValueError("文件已改变或版本已失效，请刷新工作表后重新操作")


def extension_warning(path: Path) -> str | None:
    """按命名空间检测扩展验证，不能依赖 XML 前缀恰好叫 x14。"""
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                continue
            found = False
            parser = expat.ParserCreate(namespace_separator="}")

            def start(tag, attrs):
                nonlocal found
                if tag == _X14 + "}dataValidations":
                    found = True

            def reject_dtd(*args):
                raise ValueError("工作表 XML 含不允许的 DTD，无法安全处理")

            parser.StartElementHandler = start
            parser.StartDoctypeDeclHandler = reject_dtd
            if archive.getinfo(name).file_size > 64 * 1024 * 1024:
                raise ValueError("工作表 XML 超出安全读取上限")
            with archive.open(name) as stream:
                for block in iter(lambda: stream.read(65536), b""):
                    parser.Parse(block, False)
                parser.Parse(b"", True)
            if found:
                return "文件包含 x14 扩展数据验证，openpyxl 保存可能丢失规则；暂不支持下拉创建或网页回写，请在 Excel 中操作"
    return None


def range_text(ref, absolute: bool = False) -> str:
    def coord(col, row):
        return f"${get_column_letter(col)}${row}" if absolute else f"{get_column_letter(col)}{row}"
    first, last = coord(ref.col1, ref.row1), coord(ref.col2, ref.row2)
    return first if first == last else f"{first}:{last}"


def range_cells(ref):
    for row in range(ref.row1, ref.row2 + 1):
        for col in range(ref.col1, ref.col2 + 1):
            yield f"{get_column_letter(col)}{row}"


def _overlaps(a, b) -> bool:
    return not (a.col2 < b.min_col or a.col1 > b.max_col or a.row2 < b.min_row or a.row1 > b.max_row)


def check_editable(ws, ref) -> None:
    if any(_overlaps(ref, merged) for merged in ws.merged_cells.ranges):
        raise ValueError("目标涉及合并单元格，请先在 Excel 中取消合并")
    if ws.protection.sheet:
        if any(ws[cell].protection.locked for cell in range_cells(ref)):
            raise ValueError("目标单元格受工作表保护，不能修改")


# 表名与目标格之间只剩这些虚词时，才算“明确指定目标表”：
# “订单查询的B3”“订单查询!B3”“在订单查询表里把B3”；来源短语里的表名（“用销售订单的订单号做B3”）不算。
_TARGET_LINK = re.compile(
    r"[」”』'’)\]]?\s*(?:工作表|表格|表)?\s*"
    r"(?:(?:里面|里头|里边|里|中|内|上|下)\s*)?(?:的|之)?\s*"
    r"(?:把|将|给|对)?\s*[:：!！]?\s*(?:把|将|给|对)?\s*"
)


def _linked_sheets(before: str, names: list[str]) -> list[str]:
    """找出与目标格“紧贴”的表名；只被顺带提到的表名不进入结果。"""
    names = [name for name in names
             if not any(name != other and name in other for other in names)]
    linked = []
    for name in names:
        start = 0
        while True:
            index = before.find(name, start)
            if index < 0:
                break
            if _TARGET_LINK.fullmatch(before[index + len(name):]):
                linked.append(name)
                break
            start = index + 1
    return linked


def _sheet_before(prefix: str, names: list[str], default: str) -> str:
    """未明确指定目标表时返回 default（当前表）；来源表名不会顶替目标表。"""
    before = prefix.rstrip()
    linked = _linked_sheets(before, names)
    if len(linked) > 1:
        raise DropdownClarification("请明确目标工作表，例如 订单查询!B3")
    if linked:
        return linked[0]
    if before.endswith(("!", "！")):
        # “某某表!B3”写死了目标表却对不上任何真实表名：不能默默落到当前表
        raise ValueError("指定的目标工作表不存在，请使用真实表名，例如 订单查询!B3")
    return default


def target_sheet_hint(text: str, names: list[str], cell: str | None = None) -> str | None:
    """从句子里找与目标格绑定的表名；只出现在来源处或找不到时返回 None。

    兜底流程用它把“未指定就用当前表”的同一套规则施加到模型输出上。
    """
    def norm(value: str) -> str:
        return re.sub(r"[\s$]", "", re.sub(r"[到至：]", ":", value)).upper()

    want = norm(str(cell or ""))
    for match in _RANGE.finditer(text):
        if want and norm(match.group()) != want:
            continue
        linked = _linked_sheets(text[:match.start()], names)
        return linked[0] if len(linked) == 1 else None
    return None


def parse_request(text: str, names: list[str], default: str) -> tuple[str, str, list[str] | None, str | None]:
    """只接受明确的设置指令；补充段单独解析，避免把选项中的地址误当目标。"""
    target = None
    sheet = default
    options = None
    source = None
    commands = []
    for segment in re.split(r"[（(]补充[:：]", text):
        segment = segment.rstrip("）) \n")
        source_mark = _SOURCE.search(segment)
        value_mark = _VALUES.search(segment)
        markers = [mark for mark in (source_mark, value_mark) if mark is not None]
        marker = min(markers, key=lambda mark: mark.start()) if markers else None
        command = segment[:marker.start()] if marker else segment
        commands.append(command)
        if marker:
            tail = segment[marker.end():].strip().rstrip("。；;")
            if marker is source_mark:
                source, options = tail.lstrip("="), None
            else:
                lead = _SOURCE_LEAD.match(tail)
                rest = tail[lead.end():].strip().lstrip("=") if lead else ""
                if lead and _source_like(rest):  # “选项是来自销售订单!A2:A25”
                    source, options = rest, None
                elif not lead and _source_like(tail):  # “选项是销售订单!A2:A25”
                    source, options = tail.lstrip("="), None
                elif _SECOND_ACTION.search(tail):
                    raise ValueError("选项后含有其他操作，请拆分指令；选项本身含操作语句时请改用区域来源")
                else:
                    options = [item.strip() for item in re.split(r"[、,，;；]", tail)]
                    source = None
        # 文件名本身可能带类似 A1 的字样；它不属于单元格指令。
        command = re.sub(r"[^\s，,；;]+\.xls[xm]", "", command, flags=re.I)
        matches = list(_RANGE.finditer(command))
        if len(matches) > 1:
            raise ValueError("一次只设置一个连续目标范围，请拆分指令")
        if matches:
            match = matches[0]
            target = re.sub(r"\s*(?:到|至|：)\s*", ":", match.group())
            sheet = _sheet_before(command[:match.start()], names, default)
    command_text = " ".join(commands)
    if re.search(r"取消|删除|去掉|清除|移除|不要|解释|什么意思|如何|怎么|能否|可以吗|是否", command_text):
        raise ValueError("这里只支持创建或设置下拉列表；解释、删除和询问不会修改文件")
    if re.search(r"公式|计算|汇总|求和|加边框|框起来", command_text):
        raise ValueError("请将下拉列表设置与公式、格式操作拆成两条指令")
    if not re.search(r"设置|设为|设成|添加|创建|新增|建立|生成|做成|改为|改成|加.*下拉|下拉列表|下拉框|下拉选项|下拉菜单|数据验证|数据有效性", command_text):
        raise DropdownClarification("请明确设置操作，例如：把 B3 设置为下拉列表，选项为：甲、乙")
    if not target:
        raise DropdownClarification("请补充目标单元格或范围，例如 B3 或 C2:C13")
    if not source and not options:
        raise DropdownClarification("请补充选项，例如“选项为：华东、华南”，或“内容为销售订单的订单号”")
    return sheet, range_text(parse_range(target)), options, source


def _source_like(text: str) -> bool:
    """“选项是 <引用>”的形态粗筛：区域引用或“表名+的+列标题”写法。最终仍由 resolve_source 严格校验。"""
    text = _TRAILING_NOTE.sub("", (text or "").strip()).strip()
    text = re.sub(r"[\s　]+", "", text.lstrip("="))
    if not text or re.search(r"[、,，;；]", text):
        return False
    if _STATIC_REF.fullmatch(text) or "!" in text:
        return True
    if _LOOSE_AREA.fullmatch(text):
        return True
    # “销售订单的订单编号”：表名+的+列标题（列是否存在、下方是否有数据由 resolve_source 校验）
    return bool(re.fullmatch(r"[^!、,，;；]+的[^!、,，;；]+", text))


def resolve_source(wb, sheet: str, expression: str) -> tuple[str, object]:
    """仅解析静态区域或直接指向静态区域的名称，不执行动态名称公式。

    容忍常见口语写法：到/至、缺“!”、“表的”虚词、表名带“表/工作表”后缀、句式尾注。
    """
    text = expression.strip().lstrip("=")
    if "[" in text or "]" in text:
        raise ValueError("下拉来源暂不支持外部工作簿")
    name = wb[sheet].defined_names.get(text) or wb.defined_names.get(text)
    if name is not None:
        text = name.attr_text or ""
    text = _TRAILING_NOTE.sub("", text).strip()
    located = _locate_static(wb, sheet, text)
    if located is None:
        located = _locate_header_column(wb, text)  # “销售订单的订单编号”：按列标题找数据区
    if located is None:
        raise ValueError(_source_hint(text))
    source_sheet, address = located
    ref = parse_range(address)
    if ref.col1 != ref.col2 and ref.row1 != ref.row2:
        raise ValueError("下拉来源必须为单行或单列区域")
    if ref.cell_count > MAX_OPTIONS:
        raise ValueError(f"下拉来源最多读取 {MAX_OPTIONS} 格，请缩小范围")
    return source_sheet, ref


def _source_candidates(text: str) -> list[str]:
    """同一来源的口语变体：去空白、到/至→冒号、去“的”。"""
    out: list[str] = []

    def push(item: str) -> None:
        item = item.strip()
        if item and item not in out:
            out.append(item)

    def swap(item: str) -> str:
        return re.sub(r"(?<=[A-Za-z0-9$])\s*(?:到|至)\s*(?=\$?[A-Za-z])", ":", item)

    def drop_de(item: str) -> str:
        return re.sub(r"(?<=[^'])\s*的\s*(?=\$?[A-Za-z])", "", item)

    for base in (text, re.sub(r"[\s　]+", "", text)):
        push(base)
        push(swap(base))
        push(drop_de(base))
        push(drop_de(swap(base)))
    return out


def _canonical_sheet(wb, raw: str) -> str:
    """表名回正：容忍“销售订单表/销售订单工作表”这类多写的后缀。"""
    if raw in wb.sheetnames:
        return raw
    for suffix in ("工作表", "表格", "表"):
        base = raw[: -len(suffix)].strip()
        if raw.endswith(suffix) and base in wb.sheetnames:
            return base
    raise ValueError(f"来源工作表 {raw!r} 不存在")


def _locate_static(wb, sheet: str, text: str):
    """把（可能带口语变体的）来源文本定位到 (表名, 区域)；无法识别返回 None。"""
    candidates = _source_candidates(text)
    for item in candidates:
        match = _STATIC_REF.fullmatch(item)
        if not match:
            continue
        quoted, plain, address = match.groups()
        if quoted is None and plain is None:
            return sheet, address
        raw = quoted.replace("''", "'") if quoted is not None else plain
        return _canonical_sheet(wb, raw), address
    # 无“!”的口语写法：“销售订单的A2到A25”“销售订单表A2:A25”
    for item in candidates:
        if "!" in item:
            continue
        for name in sorted(wb.sheetnames, key=len, reverse=True):
            if not item.startswith(name):
                continue
            rest = item[len(name):].lstrip("的表")
            area = _AREA.fullmatch(rest)
            if area:
                return name, re.sub(r"\s*(?:[:：]|到|至)\s*", ":", area.group())
    return None


# 列标题近似匹配的下限：覆盖“订单编号→订单号”这类口语说法，
# 再低容易错认无关列（演示表里“数量/单价/金额”等短列名彼此相似度均低于该值）。
_COLUMN_SIMILARITY = 0.65


def _header_column_span(ws, title: str):
    """在表头行找列标题：精确优先，再容忍口语近似写法；返回 (列号, 标题行号)。"""
    normalized = re.sub(r"[\s　]+", "", title)
    if not normalized:
        return None
    best = None
    best_score = 0.0
    for row in range(1, min(ws.max_row, 5) + 1):
        for cell in ws[row]:
            if cell.value is None:
                continue
            raw = cell_formula_text(cell) if cell.data_type == "f" else cell.value
            text = re.sub(r"[\s　]+", "", str(raw))
            if text == normalized:
                return cell.column, row
            score = difflib.SequenceMatcher(None, normalized, text).ratio()
            if score >= _COLUMN_SIMILARITY and score > best_score:
                best, best_score = (cell.column, row), score
    return best


def _locate_header_column(wb, text: str):
    """来源写成列标题：“销售订单的订单编号” → 销售订单!A2:A25（标题下到该列最后一个非空行）。"""
    compact = re.sub(r"[\s　]+", "", text)
    for name in sorted(wb.sheetnames, key=len, reverse=True):
        for suffix in ("工作表", "表格", "表", ""):
            prefix = name + suffix
            if not compact.startswith(prefix):
                continue
            rest = re.sub(r"^(?:里的|中的|的|里|中|[!！:：])", "", compact[len(prefix):])
            if not rest:
                continue
            found = _header_column_span(wb[name], rest)
            if found is None:
                continue
            column, header = found
            last = header
            for row in range(header + 1, wb[name].max_row + 1):
                if wb[name].cell(row=row, column=column).value not in (None, ""):
                    last = row
            if last <= header:
                continue
            letter = get_column_letter(column)
            return name, f"{letter}{header + 1}:{letter}{last}"
    return None


def _source_hint(text: str) -> str:
    """来源无法识别时的提示：整列/整行写法单独给出更直接的引导。"""
    compact = re.sub(r"[\s　]+", "", text)
    if _COLUMN_ONLY.fullmatch(compact) or _ROW_ONLY.fullmatch(compact):
        return "不支持整列或整行引用，请写明具体区域，例如 销售订单!A2:A25"
    return (f"无法识别下拉来源“{text}”：请写成“工作表名!A2:A25”或“A2:A25”这样的静态单行/单列区域，"
            "也可以引用列标题（如“销售订单的订单号”）；不支持公式、动态或级联来源")


def _value(value):
    if isinstance(value, DateValue):
        value = value.raw
    if isinstance(value, ExcelError):
        raise ValueError(f"来源公式结果为 {value.code}，不能作为下拉选项")
    if value is None or value == "":
        return None
    if not isinstance(value, (str, bool, int, float, dt.date, dt.time)):
        raise ValueError("来源包含不支持的选项类型")
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("来源包含无效数字")
    if isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
        raise ValueError("选项过长或包含 Excel 不允许的控制字符")
    return value


def option_label(value) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return format_value(value)


def option_kind(value) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (dt.date, dt.time)):
        return "date"
    return "text"


def option_index(options: list, value) -> int | None:
    for index, option in enumerate(options):
        if option_kind(option) == option_kind(value) and option == value:
            return index
    return None


def read_options(view, sheet: str, expression: str) -> tuple[list, str]:
    text = (expression or "").strip().lstrip("=")
    if text.startswith('"') and text.endswith('"'):
        options = [_value(item) for item in text[1:-1].split(",")]
        canonical = text
    else:
        source_sheet, ref = resolve_source(view.wb_formulas, sheet, text)
        canonical = f"{quote_sheetname(source_sheet)}!{range_text(ref, True)}"
        ws = view.wb_formulas[source_sheet]
        sheets = {name: view.wb_values[name] for name in view.sheet_names}
        formulas = {name: view.wb_formulas[name] for name in view.sheet_names}
        options = []
        for coord in range_cells(ref):
            cell = ws[coord]
            value = cell.value
            if cell.data_type == "e":
                raise ValueError(f"来源 {source_sheet}!{coord} 含 Excel 错误")
            if cell.data_type == "f":
                cached_cell = view.wb_values[source_sheet][coord]
                if cached_cell.data_type == "e":
                    raise ValueError(f"来源 {source_sheet}!{coord} 的缓存结果为 Excel 错误")
                value = cached_cell.value
                if value is None:
                    formula_text = cell_formula_text(cell)
                    if formula_text is None:
                        raise ValueError(f"来源 {source_sheet}!{coord} 无法可靠计算，请先在 Excel 中重算保存")
                    try:
                        value = evaluate_formula(parse_formula(formula_text), sheets, source_sheet,
                                                 formula_sheets=formulas, expand_uncached=True)
                    except (UnsupportedFormula, ValueError, TypeError, IndexError, KeyError, ZeroDivisionError) as exc:
                        raise ValueError(f"来源 {source_sheet}!{coord} 无法可靠计算，请先在 Excel 中重算保存") from exc
            options.append(_value(value))
    options = [value for value in options if value is not None]
    if not options:
        raise ValueError("下拉来源为空，没有可选择的数据")
    if len(options) > MAX_OPTIONS:
        raise ValueError(f"下拉选项不能超过 {MAX_OPTIONS} 项")
    return options, canonical


@dataclass
class DropdownSpec:
    sheet: str
    cell_range: str
    formula1: str
    options: list
    warnings: list[str] = field(default_factory=list)
    unchanged: bool = False

    def to_dict(self) -> dict:
        return {"sheet": self.sheet, "range": self.cell_range, "source": self.formula1,
                "options": [option_label(value) for value in self.options],
                "warnings": self.warnings, "unchanged": self.unchanged}


def conflict_rule(ws, ref):
    hits = [dv for dv in ws.data_validations.dataValidation
            if any(_overlaps(ref, part) for part in dv.sqref.ranges)]
    if not hits:
        return None
    if len(hits) != 1 or hits[0].type != "list" or str(hits[0].sqref) != range_text(ref):
        raise ValueError("目标与已有数据验证部分重叠或存在规则冲突，请在 Excel 中整理后重试")
    return hits[0]


def make_spec(view, sheet: str, target: str, options: list[str] | None, source: str | None) -> DropdownSpec:
    ref = parse_range(target)
    if ref.cell_count > MAX_WRITE_CELLS:
        raise ValueError(f"一次最多给 {MAX_WRITE_CELLS} 个单元格设置下拉列表")
    ws = view.wb_formulas[sheet]
    check_editable(ws, ref)
    old = conflict_rule(ws, ref)
    if options is not None:
        if not options or any(not isinstance(v, str) or not v.strip() for v in options):
            raise DropdownClarification("选项不能为空，请补充完整选项列表")
        if any(re.search(r'[,"\n\r、，;；]', v) for v in options):
            raise ValueError("选项含分隔符或引号，请改用单元格区域作为来源")
        formula1 = '"' + ",".join(options) + '"'
        if len(formula1) > 255:
            raise ValueError("固定列表超过 255 字符，请把选项放在单元格中并使用区域来源")
    else:
        formula1 = source or ""
    values, canonical = read_options(view, sheet, formula1)
    spec = DropdownSpec(sheet, range_text(ref), canonical, values)
    if old:
        try:
            _, old_source = read_options(view, sheet, old.formula1)
        except ValueError:
            old_source = None
        spec.unchanged = (old_source == canonical and not old.showDropDown and old.allowBlank
                          and old.showErrorMessage and old.errorStyle == "stop")
        spec.warnings.append("相同规则已存在，无需重复添加" if spec.unchanged else "将替换同范围已有的列表验证规则")
    invalid = [coord for coord in range_cells(ref)
               if ws[coord].value is not None and option_index(values, ws[coord].value) is None]
    if invalid:
        spec.warnings.append("现有内容不在选项中（保持不变）：" + "、".join(invalid[:8]))
    return spec


def apply_spec(wb, spec: DropdownSpec) -> bool:
    ws = wb[spec.sheet]
    ref = parse_range(spec.cell_range)
    if ref.cell_count > MAX_WRITE_CELLS:
        raise ValueError("下拉目标范围过大")
    check_editable(ws, ref)
    old = conflict_rule(ws, ref)
    if spec.unchanged:
        return False
    formula1 = spec.formula1
    if not formula1.startswith('"'):
        sheet, source = resolve_source(wb, spec.sheet, formula1)
        canonical = f"{quote_sheetname(sheet)}!{range_text(source, True)}"
        base = "_ExcelCR_DV_" + hashlib.sha256(canonical.encode()).hexdigest()[:12]
        name, count = base, 0
        # 同时避开不同来源的全局名称与可能遮蔽它的局部名称。
        while (name in ws.defined_names or
               (name in wb.defined_names and wb.defined_names[name].attr_text != canonical)):
            count += 1
            name = f"{base}_{count}"
        if name not in wb.defined_names:
            wb.defined_names.add(DefinedName(name, attr_text=canonical))
        formula1 = "=" + name
    if old is not None:
        ws.data_validations.dataValidation.remove(old)
    dv = DataValidation(type="list", formula1=formula1, showDropDown=False,
                        showErrorMessage=True, errorStyle="stop", allowBlank=True,
                        errorTitle="请选择列表中的数据", error="输入值不在下拉列表中，请重新选择。")
    dv.add(spec.cell_range)
    ws.add_data_validation(dv)
    return True


def rule_id(dv) -> str:
    return hashlib.sha256(ET.tostring(dv.to_tree())).hexdigest()[:24]


def dropdown_metadata(view, sheet: str, warning: str | None = None) -> list[dict]:
    ws = view.wb_formulas[sheet]
    result = []
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        item = {"id": rule_id(dv), "ranges": [str(r) for r in sorted(dv.sqref.ranges, key=str)],
                "options": [], "cells": {}, "error": warning, "hidden": bool(dv.showDropDown)}
        options = []
        try:
            options, _ = read_options(view, sheet, dv.formula1)
            item["options"] = [{"label": option_label(v), "type": option_kind(v)} for v in options]
        except ValueError as exc:
            item["error"] = str(exc)
        for part in dv.sqref.ranges:
            # 只为网页可显示的区域附加当前值，绝不展开整列验证。
            for row in range(part.min_row, min(part.max_row, 300) + 1):
                for col in range(part.min_col, min(part.max_col, 40) + 1):
                    cell = ws.cell(row, col)
                    raw = cell_formula_text(cell) if cell.data_type == "f" else cell.value
                    error = item["error"]
                    try:
                        check_editable(ws, parse_range(cell.coordinate))
                        hits = [rule for rule in ws.data_validations.dataValidation if cell.coordinate in rule]
                        if len(hits) != 1:
                            raise ValueError("此单元格存在重叠验证规则")
                    except ValueError as exc:
                        error = str(exc)
                    item["cells"][cell.coordinate] = {
                        "selected": option_index(options, raw),
                        "current": option_label(raw) if raw is not None else "",
                        "error": error,
                    }
        result.append(item)
    return result


def selected_value(view, sheet: str, cell: str, identity: str, index: int):
    ref = parse_range(cell)
    if ref.cell_count != 1:
        raise ValueError("一次只能选择一个单元格")
    ws = view.wb_formulas[sheet]
    check_editable(ws, ref)
    hits = [dv for dv in ws.data_validations.dataValidation if range_text(ref) in dv]
    if len(hits) != 1 or hits[0].type != "list" or rule_id(hits[0]) != identity:
        raise ValueError("下拉规则不存在、已变化或存在重叠，请刷新后重试")
    if hits[0].showDropDown:
        raise ValueError("此规则已隐藏下拉箭头，请在 Excel 中操作")
    options, _ = read_options(view, sheet, hits[0].formula1)
    if type(index) is not int or not 0 <= index < len(options):
        raise ValueError("下拉选项索引无效")
    return range_text(ref), options[index]
