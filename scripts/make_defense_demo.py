"""生成答辩演示用的多表 Excel 文件（与 main.py 同目录）。

运行：python scripts/make_defense_demo.py
生成：答辩演示.xlsx —— 5 张工作表，一张表对应一组演示能力：

  成绩单      12 名学生 × 3 个班，三科成绩已填；总分/平均分/评级列留空
              → 演示 生成公式并填充（SUM / ROUND+AVERAGE / 嵌套 IF 评级）
  班级信息    3 个班 + 班主任 + 教室；人数/平均分列留空
              → 演示 跨表条件聚合（COUNTIF / AVERAGEIF 引用成绩单）
  销售订单    24 笔订单（含日期列）；金额列留空
              → 演示 乘法列填充；SUMIFS/COUNTIFS 的区域/产品天然素材
  订单查询    预置 1 条 IFERROR+VLOOKUP 公式（B4），其余留空、无边框
              → 演示 解释公式（解释 B4）；把 B5:B8 查询补齐；"把 A3 到 B8 框起来"加框
  区域汇总    空白表
              → 演示 新建表格自动套格式（表头浅蓝底 + 整表细边框）

约定：所有"要写公式"的目标单元格一律留空；预置公式仅订单查询 B4 一处。
样式：表头浅蓝底加粗（与程序自动套用的一致），数据区细边框；
      目标列的表头有样式、数据区留裸——写入时会自动沿用左侧相邻列格式。

参考数值（演示时对照）：
  总分   265 260 275 160 / 236 253 266 259 / 208 280 245 248
  评级   良好 良好 优秀 待提升 / 待提升 良好 良好 良好 / 待提升 优秀 良好 良好
  班级人数 4 / 4 / 4        班级平均分 80.0 / 84.5 / 81.8
  订单金额合计（写完后可核）= 222,142
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "答辩演示.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # 浅黄：可改写的输入位
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
CENTER = Alignment(horizontal="center")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header(ws: Worksheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def _write_rows(ws: Worksheet, start_row: int, rows) -> None:
    for offset, row in enumerate(rows):
        for col, value in enumerate(row, start=1):
            if value is not None:
                ws.cell(row=start_row + offset, column=col, value=value)


def _border_area(ws: Worksheet, min_row: int, max_row: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = BORDER


def _widths(ws: Worksheet, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


# ------------------------------------------------------------------ 1. 成绩单
_SCORES = [
    ("S01", "张伟", "高一1班", 88, 92, 85),
    ("S02", "李娜", "高一1班", 82, 88, 90),
    ("S03", "王强", "高一1班", 95, 88, 92),
    ("S04", "赵敏", "高一1班", 52, 61, 47),
    ("S05", "刘洋", "高一2班", 79, 83, 74),
    ("S06", "陈静", "高一2班", 91, 78, 84),
    ("S07", "杨帆", "高一2班", 88, 92, 86),
    ("S08", "周雪", "高一2班", 85, 93, 81),
    ("S09", "吴磊", "高一3班", 73, 66, 69),
    ("S10", "郑好", "高一3班", 95, 94, 91),
    ("S11", "孙悦", "高一3班", 84, 79, 82),
    ("S12", "林浩", "高一3班", 82, 90, 76),
]


def make_scores(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "成绩单"
    header = ["学号", "姓名", "班级", "语文", "数学", "英语", "总分", "平均分", "评级"]
    ws.append(header)
    _write_rows(ws, 2, _SCORES)
    _header(ws, 1, len(header))
    _border_area(ws, 2, 13, 6)  # 数据区只框到英语列：总分/平均分/评级留裸，写入时自动沿用格式
    _widths(ws, {"A": 7, "B": 8, "C": 10, "D": 6, "E": 6, "F": 6, "G": 7, "H": 8, "I": 8})


# ------------------------------------------------------------------ 2. 班级信息
_CLASSES = [
    ("高一1班", "刘老师", "101"),
    ("高一2班", "陈老师", "102"),
    ("高一3班", "王老师", "103"),
]


def make_classes(wb: Workbook) -> None:
    ws = wb.create_sheet("班级信息")
    header = ["班级", "班主任", "教室", "班级人数", "班级平均分"]
    ws.append(header)
    _write_rows(ws, 2, _CLASSES)
    _header(ws, 1, len(header))
    _border_area(ws, 2, 4, 3)  # 人数/平均分列留裸
    _widths(ws, {"A": 10, "B": 9, "C": 7, "D": 9, "E": 12})


# ------------------------------------------------------------------ 3. 销售订单
_ORDERS = [
    ("SO-2601", dt.date(2026, 1, 5), "深圳鹏城数码", "华南", "显示器27寸", 12, 1099),
    ("SO-2602", dt.date(2026, 1, 12), "广州天河电子", "华南", "机械键盘", 30, 399),
    ("SO-2603", dt.date(2026, 1, 20), "上海睿智科技", "华东", "无线鼠标", 40, 129),
    ("SO-2604", dt.date(2026, 2, 3), "杭州西湖智能", "华东", "笔记本支架", 25, 89),
    ("SO-2605", dt.date(2026, 2, 11), "北京中关村商贸", "华北", "显示器27寸", 18, 1099),
    ("SO-2606", dt.date(2026, 2, 18), "成都天府软件", "西南", "扩展坞", 22, 249),
    ("SO-2607", dt.date(2026, 2, 25), "武汉光谷信息", "华中", "机械键盘", 35, 399),
    ("SO-2608", dt.date(2026, 3, 4), "西安高新电子", "西北", "无线鼠标", 50, 129),
    ("SO-2609", dt.date(2026, 3, 10), "深圳鹏城数码", "华南", "扩展坞", 16, 249),
    ("SO-2610", dt.date(2026, 3, 17), "上海睿智科技", "华东", "显示器27寸", 9, 1099),
    ("SO-2611", dt.date(2026, 3, 24), "北京中关村商贸", "华北", "机械键盘", 28, 399),
    ("SO-2612", dt.date(2026, 4, 2), "广州天河电子", "华南", "笔记本支架", 36, 89),
    ("SO-2613", dt.date(2026, 4, 9), "杭州西湖智能", "华东", "无线鼠标", 45, 129),
    ("SO-2614", dt.date(2026, 4, 16), "成都天府软件", "西南", "显示器27寸", 14, 1099),
    ("SO-2615", dt.date(2026, 4, 23), "武汉光谷信息", "华中", "扩展坞", 19, 249),
    ("SO-2616", dt.date(2026, 5, 6), "深圳鹏城数码", "华南", "机械键盘", 42, 399),
    ("SO-2617", dt.date(2026, 5, 13), "上海睿智科技", "华东", "笔记本支架", 31, 89),
    ("SO-2618", dt.date(2026, 5, 20), "北京中关村商贸", "华北", "无线鼠标", 55, 129),
    ("SO-2619", dt.date(2026, 5, 27), "西安高新电子", "西北", "显示器27寸", 11, 1099),
    ("SO-2620", dt.date(2026, 6, 3), "广州天河电子", "华南", "扩展坞", 24, 249),
    ("SO-2621", dt.date(2026, 6, 10), "杭州西湖智能", "华东", "机械键盘", 33, 399),
    ("SO-2622", dt.date(2026, 6, 17), "成都天府软件", "西南", "无线鼠标", 48, 129),
    ("SO-2623", dt.date(2026, 6, 24), "深圳鹏城数码", "华南", "显示器27寸", 20, 1099),
    ("SO-2624", dt.date(2026, 7, 1), "上海睿智科技", "华东", "扩展坞", 15, 249),
]


def make_orders(wb: Workbook) -> None:
    ws = wb.create_sheet("销售订单")
    header = ["订单号", "日期", "客户", "区域", "产品", "数量", "单价", "金额"]
    ws.append(header)
    _write_rows(ws, 2, _ORDERS)
    for row in ws.iter_rows(min_row=2, max_row=1 + len(_ORDERS), min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"
    for row in ws.iter_rows(min_row=2, max_row=1 + len(_ORDERS), min_col=7, max_col=7):
        for cell in row:
            cell.number_format = "#,##0"
    _header(ws, 1, len(header))
    _border_area(ws, 2, 1 + len(_ORDERS), 7)  # 金额列留裸
    _widths(ws, {"A": 10, "B": 12, "C": 17, "D": 8, "E": 14, "F": 7, "G": 8, "H": 11})


# ------------------------------------------------------------------ 4. 订单查询
def make_lookup(wb: Workbook) -> None:
    """预置 B4 一条 IFERROR+VLOOKUP（解释公式素材）；B5:B8 留空供现场生成；整表不加边框。"""
    ws = wb.create_sheet("订单查询")
    ws["A1"] = "订单查询面板"
    ws["A1"].font = TITLE
    labels = ["订单编号：", "客户名称", "区域", "产品", "数量", "金额"]
    for offset, label in enumerate(labels, start=3):
        ws.cell(row=offset, column=1, value=label)
    ws["B3"] = "SO-2603"
    ws["B3"].fill = INPUT_FILL  # 浅黄提示：这里是可改写的订单号
    ws["B4"] = '=IFERROR(VLOOKUP($B$3,销售订单!$A$2:$H$25,3,0),"")'
    _widths(ws, {"A": 14, "B": 22})


# ------------------------------------------------------------------ 5. 区域汇总（空白，供"新建表格"演示）
def make_blank(wb: Workbook) -> None:
    wb.create_sheet("区域汇总")


def main() -> None:
    wb = Workbook()
    make_scores(wb)
    make_classes(wb)
    make_orders(wb)
    make_lookup(wb)
    make_blank(wb)
    wb.save(OUT)
    wb.close()
    print(f"已生成 {OUT.name}")
    print("工作表：成绩单 / 班级信息 / 销售订单 / 订单查询 / 区域汇总")


if __name__ == "__main__":
    main()
