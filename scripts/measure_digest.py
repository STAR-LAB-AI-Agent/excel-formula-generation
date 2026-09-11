"""提示词规模检查：看一张表文本化后有多大、是否触发行数截断。

用法：
    python scripts/measure_digest.py                    # 用 测试数据.xlsx + 合成大表各测一次
    python scripts/measure_digest.py 某个文件.xlsx       # 只测指定文件
    python scripts/measure_digest.py 某个文件.xlsx 工作表  # 指定工作表

表格现在以无损 TSV 全量交给模型（表头行与数据起始行由模型自行判断），
因此提示规模随表变大而增长，直到 DIGEST_MAX_ROWS 触发头尾截断后趋于恒定。
字符数用作 Token 的近似（ASCII 约 4 字符/Token，中文约 1.5 字符/Token），
真实 Token 数以 DeepSeek 返回的 usage 为准，交互界面每次都会打印。
"""
from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402

from excel_formula.config import DIGEST_MAX_ROWS  # noqa: E402
from excel_formula.excel_reader import WorkbookView  # noqa: E402


def measure(path: Path, sheet: str | None = None) -> dict:
    with WorkbookView(path) as view:
        digest = view.digest(sheet)
        prompt = digest.to_prompt()
    return {
        "file": path.name,
        "sheet": digest.name,
        "size": f"{digest.max_row}行 × {digest.max_column}列",
        "prompt_chars": len(prompt),
        "rows_sent": len(digest.rows),
        "rows_omitted": digest.truncated_rows,
    }


def make_large_table(directory: Path, rows: int = 500) -> Path:
    """合成一张大表，用来验证行数闸门确实生效。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单"
    sheet.append(["订单号", "客户", "地区", "产品", "单价", "数量", "折扣", "下单日期"])
    for index in range(1, rows):
        sheet.append(
            [
                f"SO{index:05d}",
                f"客户{index % 37}",
                ["华东", "华北", "华南", "西南"][index % 4],
                f"产品{index % 19}",
                round(50 + index % 500 * 1.7, 2),
                index % 25 + 1,
                round((index % 5) / 20, 2),
                f"2026-0{index % 9 + 1}-15",
            ]
        )
    path = directory / "large.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def report(rows: list[dict]) -> None:
    header = f"{'文件':<16}{'规模':<16}{'提示字符':>10}{'发送行数':>10}{'省略行数':>10}"
    print(header)
    print("-" * len(header))
    for row in rows:
        print(
            f"{row['file']:<16}{row['size']:<16}{row['prompt_chars']:>10}"
            f"{row['rows_sent']:>10}{row['rows_omitted']:>10}"
        )


def main() -> int:
    if len(sys.argv) > 1:
        sheet = sys.argv[2] if len(sys.argv) > 2 else None
        report([measure(Path(sys.argv[1]), sheet)])
        return 0

    results = []
    default = Path("测试数据.xlsx")
    if default.is_file():
        results.append(measure(default))
    with TemporaryDirectory() as tmp:
        results.append(measure(make_large_table(Path(tmp))))
    report(results)
    print(f"\n说明：单元格全部原样发送，不做表头/类型推断；超过 {DIGEST_MAX_ROWS} 行才启用")
    print("      头尾截断（头 50 行 + 尾 10 行），截断位置会在提示里显式标注。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
